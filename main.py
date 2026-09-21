import sys
import os
import re
import csv
import threading
import base64
import io
from datetime import datetime
import tkinter as tk
from tkinter import messagebox, filedialog
from typing import Optional, Dict, Tuple, Any, List
from PIL import Image

import customtkinter as ctk

from license_manager import LicenseManager, get_hardware_id
from database import MinimarketDB
from smtp_manager import SMTPManager
from whatsapp_manager import GestorWhatsApp
from automator import Automator

# ==========================================
# CÓDIGOS DE PAÍSES PARA WHATSAPP
# ==========================================
PAISES_WHATSAPP: Dict[str, str] = {
    "🇧🇴 Bolivia (+591)": "+591",
    "🇵🇪 Perú (+51)": "+51",
    "🇦🇷 Argentina (+54)": "+54",
    "🇨🇱 Chile (+56)": "+56",
    "🇨🇴 Colombia (+57)": "+57",
    "🇪🇨 Ecuador (+593)": "+593",
    "🇵🇾 Paraguay (+595)": "+595",
    "🇺🇾 Uruguay (+598)": "+598",
    "🇧🇷 Brasil (+55)": "+55",
    "🇲🇽 México (+52)": "+52",
    "🇺🇸 USA (+1)": "+1",
    "🇪🇸 España (+34)": "+34",
    "🌐 Directo / Sin prefijo": ""
}

# ==========================================
# PLANTILLAS DE MENSAJERÍA PARA MINIMARKET
# ==========================================
# Plantillas optimizadas con Copywriting persuasivo para WhatsApp
PLANTILLAS_MINIMARKET: Dict[str, str] = {
    "🛒 Ofertas del Día en Abarrotes y Despensa": (
        "¡Hola {nombre}! 🛒 En *Minimarket* cuidamos tu economía familiar hoy {fecha}:\n\n"
        "🔥 *Arroz, azúcar y aceites* con 15% de descuento\n"
        "🥛 *Lácteos y embutidos frescos* a precios especiales de rebaja\n"
        "🥫 *3x2* en conservas, fideos y salsas seleccionadas\n\n"
        "📍 ¡Visítanos hoy o haz tu pedido por aquí y te lo llevamos volando a casa! 🛵💨"
    ),
    "🥑 Día de Mercado: Frutas y Verduras Fresquitas": (
        "¡Hola {nombre}! 🥑🍎 ¡Llegó mercadería fresquita a *Minimarket* hoy {fecha}!\n\n"
        "🥗 *Paltas cremosas, tomates, lechugas y verduras del día*\n"
        "🍌 *Plátanos dulces, manzanas, naranjas y frutas de temporada*\n"
        "🥚 *Huevos frescos de granja* por docena y maple al mejor precio\n\n"
        "✨ ¡Calidad y frescura garantizada para tu familia! Pasa por tu compra o pídelo a domicilio. 🛵💨"
    ),
    "🥖 Pan Caliente, Empanadas y Desayunos": (
        "¡Buenos días {nombre}! 🥖☀️\n\n"
        "En *Minimarket* ya salió la primera tanda de *pan caliente y crujiente* recién horneado:\n\n"
        "☕ Café selecto, leche fresca, huevos y mermeladas\n"
        "🥐 Empanaditas calientes y bocadillos para tu mañana\n"
        "🧀 Jamones y quesos frescos rebanados al instante\n\n"
        "¡Empieza tu día con la mejor energía! Pasa antes de que se termine. 🥐✨"
    ),
    "🥩 Carnicería y Embutidos de Calidad": (
        "¡Hola {nombre}! 🥩🍗 Hoy en *Minimarket* tenemos los mejores cortes y embutidos para tu almuerzo:\n\n"
        "🍗 *Pollo fresco y carnes seleccionadas*\n"
        "🥓 *Embutidos, chorizos parrilleros y salchichas*\n"
        "🧈 *Mantequilla y quesos artesanales*\n\n"
        "¡Ahorra sin sacrificar calidad en tu mesa! Haz tu encargo por aquí. 🛵🛒"
    ),
    "🥤 Combo Fin de Semana: Cervezas, Snacks y Bebidas": (
        "¡Hola {nombre}! 🍿🍻 ¡Llegó el fin de semana a *Minimarket*!\n\n"
        "❄️ *Cervezas, gaseosas y jugos al polo* (bien heladas)\n"
        "🍟 *Snacks, papitas, chocolates, galletas y piqueos en oferta*\n"
        "🧊 *Bolsas de hielo y carbón* listos para tu reunión o parrilla\n\n"
        "🎉 ¡No salgas de casa! Pide tu combo por WhatsApp y te lo llevamos en minutos. 🛵💨"
    ),
    "🧼 Pack Ahorro Hogar: Limpieza y Aseo": (
        "¡Hola {nombre}! 🧼✨ En *Minimarket* armamos packs de ahorro para que tu hogar quede impecable:\n\n"
        "🧺 *Detergentes y suavizantes* en tamaño económico\n"
        "🧽 *Lavavajillas, desinfectantes y lavandina multiusos*\n"
        "🧻 *Papel higiénico y toallas de cocina* en paquetes familiares\n\n"
        "¡Ven por tu pack de ahorro o solicítalo a domicilio hoy mismo! 🏪👍"
    ),
    "📦 Confirmación de Pedido / Delivery en Camino": (
        "¡Hola {nombre}! 📦 Tu pedido en *Minimarket* ya fue preparado con el máximo cuidado.\n\n"
        "🛵 Nuestro repartidor ya va en camino hacia tu dirección.\n"
        "💵 Recuerda tener listo tu medio de pago (Efectivo / QR / Transferencia / Tarjeta).\n\n"
        "¡Muchas gracias por tu preferencia y confianza de siempre! 😊🙏"
    ),
    "🎁 Promoción 2x1 y Liquidación de Stock": (
        "¡Hola {nombre}! 🚨 ¡ALERTA DE OFERTAS en *Minimarket*! 🚨\n\n"
        "Solo por hoy {fecha} tenemos promociones relámpago:\n"
        "💥 *2x1 en galletas, snacks y golosinas seleccionadas*\n"
        "💥 *20% de descuento* en jugos y bebidas refrescantes\n"
        "💥 Precios de costo en abarrotes seleccionados\n\n"
        "⏰ ¡Stock limitado! No te quedes sin tus favoritos. ¡Visítanos ya! 🏃‍♀️💨"
    ),
    "🎉 Promoción de Cumpleaños / Fidelización": (
        "¡Feliz Cumpleaños {nombre}! 🎂🎉🎈\n\n"
        "En *Minimarket* estamos felices de celebrar contigo en este día especial.\n"
        "🎁 Hoy tienes un *15% de descuento en toda tu compra* presentando este mensaje.\n\n"
        "¡Que pases un día maravilloso junto a tu familia! 🥳✨"
    ),
    "⭐ Cliente VIP: Descuento Exclusivo": (
        "¡Estimado(a) {nombre}! ⭐ En *Minimarket* premiamos tu preferencia constante:\n\n"
        "👑 Como cliente especial de la casa, hoy tienes un *descuento exclusivo del 10%* en tu próxima compra o delivery.\n"
        "¡Gracias por ser parte de nuestra gran familia! Estamos a tu orden. 🏪🤝"
    ),
    "💰 Recordatorio Amable de Saldo / Fiado": (
        "Estimado(a) {nombre}, le saludamos cordialmente de *Minimarket*. 🏪\n\n"
        "Le recordamos de manera amable que mantiene un saldo pendiente por compras recientes en nuestra tienda.\n"
        "Agradeceremos pueda pasar a cancelar o realizar su pago vía QR / transferencia a su comodidad.\n\n"
        "¡Agradecemos mucho su confianza y comprensión! Cualquier duda estamos a su servicio. 👍"
    ),
    "⏰ Horario de Atención y Nuevos Servicios": (
        "Estimado(a) {nombre}, en *Minimarket* seguimos a tu lado siempre cerca: 🏪\n\n"
        "⏰ *Horario continuo:* Lunes a Domingo de 7:00 AM a 10:00 PM\n"
        "🛵 *Delivery rápido* en toda la zona\n"
        "💳 *Aceptamos:* Efectivo, Transferencias, Pagos con QR y Tarjetas\n\n"
        "¡Siempre listos para atenderte con una sonrisa! 😊"
    ),
    "✍️ Mensaje Personalizado (En Blanco)": (
        "¡Hola {nombre}! Le escribimos de *Minimarket* para informarle que..."
    )
}

# Plantillas profesionales con Copywriting de Email Marketing para Minimarket (Asunto + Cuerpo)
PLANTILLAS_EMAIL_MINIMARKET: Dict[str, Dict[str, str]] = {
    "🛒 Super Ofertas de la Semana en Abarrotes y Despensa": {
        "asunto": "🛒 ¡Ahorra en grande esta semana en {minimarket}! Ofertas exclusivas para ti",
        "cuerpo": (
            "Estimado(a) {nombre},\n\n"
            "Esperamos que se encuentre muy bien. En {minimarket} queremos cuidar su economía familiar, "
            "por lo que esta semana traemos descuentos especiales en productos esenciales de la canasta básica:\n\n"
            "✅ Arroz, azúcar y aceite de las mejores marcas con hasta 15% de descuento.\n"
            "✅ Leche, yogures y quesos frescos seleccionados para su familia.\n"
            "✅ Promoción 3x2 en fideos y salsas de tomate.\n"
            "✅ Descuentos imperdibles en artículos de limpieza y aseo personal.\n\n"
            "📅 Promoción válida del {fecha} hasta agotar stock.\n"
            "📍 Visítenos hoy mismo o realice su pedido a domicilio y se lo llevamos directamente a su puerta.\n\n"
            "¡Agradecemos su preferencia!\n"
            "Atentamente,\n"
            "El equipo de {minimarket}"
        )
    },
    "🥑 Llegada de Frutas y Verduras Frescas del Día": {
        "asunto": "🥑 ¡Mercadería Fresca recién llegada a {minimarket}! Frutas y verduras del día",
        "cuerpo": (
            "¡Hola {nombre}!\n\n"
            "Le informamos que hoy {fecha} acabamos de recibir nuestro camión con las frutas y verduras más frescas del mercado:\n\n"
            "🥦 Paltas cremosas, tomates maduros, lechugas frescas y verduras del día.\n"
            "🍎 Manzanas crujientes, plátanos dulces, naranjas para jugo y frutas de temporada.\n"
            "🥚 Huevos frescos de granja por docena y maple al mejor precio.\n\n"
            "Disfrute de la máxima frescura y calidad que su mesa merece.\n\n"
            "¡Pase por su compra o escríbanos para apartar sus productos favoritos antes de que se agoten!\n\n"
            "Saludos cordiales,\n"
            "{minimarket}"
        )
    },
    "🥩 Carnicería, Embutidos y Lácteos a Precios Especiales": {
        "asunto": "🥩 Calidad y Frescura: Carnes, embutidos y lácteos en {minimarket}",
        "cuerpo": (
            "Estimado(a) {nombre},\n\n"
            "Para sus almuerzos y preparaciones familiares, en {minimarket} le ofrecemos los cortes y embutidos más frescos:\n\n"
            "🍗 Pollo fresco, carnes seleccionadas y cortes listos para la semana.\n"
            "🥓 Jamones, salchichas, chorizos y embutidos de primera calidad.\n"
            "🧈 Mantequilla, queso criollo y lácteos artesanales.\n\n"
            "Precios competitivos y la mejor higiene y conservación garantizada.\n\n"
            "¡Le esperamos con la cordialidad de siempre!\n\n"
            "{minimarket}"
        )
    },
    "🥐 Desayunos y Pan Caliente Recién Horneado": {
        "asunto": "🥐🥖 ¡Empieza tu día con pan calientito y desayunos en {minimarket}!",
        "cuerpo": (
            "¡Buenos días {nombre}!\n\n"
            "En {minimarket} ya tenemos el pan caliente, crujiente y recién salido del horno listo para su desayuno.\n\n"
            "☕ Acompañe su mañana con café selecto, leche fresca, té, mermeladas y quesos.\n"
            "🥪 Empanadas y bocadillos recién preparados.\n\n"
            "Abiertos desde muy temprano para que empiece su jornada con la mejor energía.\n\n"
            "¡Que tenga un excelente día!\n"
            "{minimarket}"
        )
    },
    "🥤 Combo Fin de Semana: Snacks, Bebidas y Cervezas Heladas": {
        "asunto": "🎉 ¡Llegó el Fin de Semana! Snacks, piqueos y bebidas heladas en {minimarket}",
        "cuerpo": (
            "¡Hola {nombre}!\n\n"
            "¿Planes para descansar o compartir con familia y amigos este fin de semana? En {minimarket} tenemos todo listo:\n\n"
            "❄️ Cervezas, gaseosas, aguas y jugos al polo (bien heladas).\n"
            "🍿 Papitas, piqueos, chocolates, galletas y snacks surtidos.\n"
            "🧊 Bolsas de hielo, vasos descartables y carbón para su parrilla.\n\n"
            "No se preocupe por salir: solicite su pedido a domicilio y se lo entregamos en minutos.\n\n"
            "¡A disfrutar el fin de semana!\n"
            "{minimarket}"
        )
    },
    "🧼 Pack Ahorro Hogar: Limpieza y Aseo Familiar": {
        "asunto": "🧼 ¡Tu Hogar Impecable! Ofertas en productos de limpieza en {minimarket}",
        "cuerpo": (
            "Estimado(a) {nombre},\n\n"
            "Mantener su hogar limpio y desinfectado ahora cuesta menos en {minimarket}.\n\n"
            "🧺 Detergentes en polvo y líquidos en formatos económicos.\n"
            "✨ Lavavajillas, desinfectantes multiusos y lavandina.\n"
            "🧻 Papel higiénico, servilletas y toallas de cocina en paquetes familiares.\n\n"
            "Calidad y rendimiento comprobado para su bolsillo.\n\n"
            "¡Visítenos hoy o solicite su despacho a domicilio!\n\n"
            "Atentamente,\n"
            "{minimarket}"
        )
    },
    "🛵 Delivery Express a Domicilio (Pedido Fácil)": {
        "asunto": "🛵 ¿Te falta algo en la cocina? Pide por Delivery en {minimarket}",
        "cuerpo": (
            "¡Hola {nombre}!\n\n"
            "¿Necesita algo urgente y no puede salir de casa o del trabajo? ¡En {minimarket} se lo solucionamos!\n\n"
            "🛵 Contamos con servicio de Delivery Express rápido y seguro en toda la zona.\n"
            "📲 Solo díganos qué necesita y nuestro repartidor llegará en pocos minutos a su dirección.\n"
            "💳 Aceptamos efectivo, transferencias bancarias, QR y tarjetas.\n\n"
            "Guarde nuestro contacto y realice su pedido cuando lo necesite.\n\n"
            "¡Estamos para servirle!\n"
            "{minimarket}"
        )
    },
    "🎂 Descuento de Cumpleaños para Clientes Frecuentes": {
        "asunto": "🎂🎁 ¡Feliz Cumpleaños {nombre}! Tenemos un regalo especial para ti en {minimarket}",
        "cuerpo": (
            "¡Feliz Cumpleaños {nombre}! 🎉🎂🎈\n\n"
            "De parte de toda la familia de {minimarket}, queremos desearle un día lleno de bendiciones, salud y alegría en compañía de sus seres queridos.\n\n"
            "🎁 Para celebrarlo, le obsequiamos un 15% de descuento en toda su compra durante la semana de su cumpleaños presentando este correo.\n\n"
            "¡Gracias por ser un cliente tan valioso para nuestro negocio!\n\n"
            "Un abrazo cordial,\n"
            "{minimarket}"
        )
    },
    "⭐ Programa Cliente VIP: Beneficio Exclusivo del Mes": {
        "asunto": "⭐ Eres un Cliente VIP en {minimarket}: Accede a beneficios exclusivos",
        "cuerpo": (
            "Estimado(a) {nombre},\n\n"
            "Queremos agradecerle especialmente por su lealtad y confianza constante en {minimarket}.\n\n"
            "Como cliente preferencial, este mes cuenta con beneficios únicos en nuestra tienda:\n"
            "🌟 Precios preferenciales en compras por mayor y canasta familiar.\n"
            "🌟 Atención prioritaria en pedidos a domicilio.\n"
            "🌟 Acceso anticipado a nuestras ofertas y promociones de temporada.\n\n"
            "¡Será un gusto atenderle siempre con el trato y dedicación que usted merece!\n\n"
            "Cordialmente,\n"
            "{minimarket}"
        )
    },
    "💰 Recordatorio Cordial y Amable de Saldo Pendiente": {
        "asunto": "🏪 Estado de Cuenta / Saldo Pendiente en {minimarket}",
        "cuerpo": (
            "Estimado(a) {nombre},\n\n"
            "Le enviamos un saludo cordial de parte de la administración de {minimarket}.\n\n"
            "Nos comunicamos respetuosamente para informarle que registra un saldo pendiente de pago correspondiente a sus compras recientes en nuestra tienda.\n\n"
            "Agradeceremos pueda pasar por el local para regularizar su cuenta a la brevedad posible, o realizar su abono a través de transferencia bancaria o código QR.\n\n"
            "Si ya realizó su pago recientemente, por favor desestime este mensaje. Para cualquier consulta o detalle del saldo, estamos a su total disposición.\n\n"
            "Muchas gracias por su atención y comprensión.\n\n"
            "Atentamente,\n"
            "Administración de {minimarket}"
        )
    },
    "⏰ Horarios de Atención, Medios de Pago y Novedades": {
        "asunto": "🏪 Novedades y Horarios de Atención en {minimarket}",
        "cuerpo": (
            "Estimado(a) cliente {nombre},\n\n"
            "En {minimarket} renovamos nuestro compromiso de brindarle el mejor servicio, variedad y comodidad:\n\n"
            "⏰ Horario continuo de atención: Lunes a Domingo de 7:00 AM a 10:00 PM.\n"
            "💳 Medios de pago disponibles: Efectivo, Tarjetas de Débito/Crédito, Transferencia bancaria y pagos con QR.\n"
            "🛍️ Nuevos productos agregados a nuestras estanterías esta semana.\n\n"
            "¡Siempre a la vuelta de su hogar para atenderle con una sonrisa!\n\n"
            "Atentamente,\n"
            "{minimarket}"
        )
    },
    "✍️ Mensaje Personalizado en Blanco": {
        "asunto": "Aviso Importante de {minimarket}",
        "cuerpo": (
            "Estimado(a) {nombre},\n\n"
            "Le escribimos de {minimarket} para comunicarle que...\n\n"
            "Quedamos a su disposición ante cualquier duda o consulta.\n\n"
            "Atentamente,\n"
            "{minimarket}"
        )
    }
}

# ==========================================
# CONSTANTES Y SEGMENTACIÓN DE CLIENTES
# ==========================================
CATEGORIAS_CLIENTE = [
    "🥉 Bronce",
    "🥈 Plata",
    "🥇 Oro",
    "💎 Platino",
    "✨ Diamante / VIP"
]

def normalizar_categoria(cat_raw: Any) -> str:
    """Normaliza cualquier texto de categoría a una de las 5 categorías oficiales."""
    if not cat_raw:
        return "🥉 Bronce"
    c = str(cat_raw).strip().lower()
    if "diamante" in c or "vip" in c:
        return "✨ Diamante / VIP"
    elif "platino" in c or "muy" in c:
        return "💎 Platino"
    elif "oro" in c or ("importante" in c and "poco" not in c and "no" not in c and "muy" not in c):
        return "🥇 Oro"
    elif "plata" in c or "poco" in c:
        return "🥈 Plata"
    elif "bronce" in c or "no tan" in c or "baja" in c:
        return "🥉 Bronce"
    return "🥉 Bronce"

def normalizar_estado(est_raw: Any) -> str:
    """Normaliza el estado estrictamente a 'Activo' o 'Inactivo'."""
    if not est_raw:
        return "Activo"
    e = str(est_raw).strip().lower()
    if "inactiv" in e or "baja" in e or "bloque" in e or "desactiv" in e:
        return "Inactivo"
    return "Activo"

def normalizar_pais(pais_raw: Any, tel_raw: Any = "") -> str:
    """
    Normaliza el nombre o código de país a la clave oficial de PAISES_WHATSAPP.
    Acepta nombres en minúsculas (ej. 'bolivia', 'peru', 'argentina') o prefijos ('+591', '591', etc.).
    Si el teléfono ya trae prefijo y país viene vacío, deduce el país por el teléfono.
    Por defecto retorna '🇧🇴 Bolivia (+591)'.
    """
    import unicodedata
    p_raw = str(pais_raw or "").strip().lower()
    # Eliminar acentos y diacríticos para comparación infalible
    p = ''.join(c for c in unicodedata.normalize('NFD', p_raw) if unicodedata.category(c) != 'Mn')
    t = re.sub(r'\D', '', str(tel_raw or ""))

    if p:
        if "boliv" in p or "591" in p:
            return "🇧🇴 Bolivia (+591)"
        elif "per" in p or p == "51" or "+51" in p:
            return "🇵🇪 Perú (+51)"
        elif "arg" in p or p == "54" or "+54" in p:
            return "🇦🇷 Argentina (+54)"
        elif "chil" in p or p == "56" or "+56" in p:
            return "🇨🇱 Chile (+56)"
        elif "colomb" in p or p == "57" or "+57" in p:
            return "🇨🇴 Colombia (+57)"
        elif "ecuad" in p or "593" in p:
            return "🇪🇨 Ecuador (+593)"
        elif "parag" in p or "595" in p:
            return "🇵🇾 Paraguay (+595)"
        elif "urug" in p or "598" in p:
            return "🇺🇾 Uruguay (+598)"
        elif "bras" in p or "braz" in p or p == "55" or "+55" in p:
            return "🇧🇷 Brasil (+55)"
        elif "mex" in p or p == "52" or "+52" in p:
            return "🇲🇽 México (+52)"
        elif "usa" in p or "esta" in p or "eeuu" in p or p == "1" or "+1" in p:
            return "🇺🇸 USA (+1)"
        elif "esp" in p or p == "34" or "+34" in p:
            return "🇪🇸 España (+34)"
        elif "direc" in p or "sin" in p or p == "0":
            return "🌐 Directo / Sin prefijo"

    # Si no se especificó país, deducir según los primeros dígitos del teléfono
    if t:
        if t.startswith("591") and len(t) > 8:
            return "🇧🇴 Bolivia (+591)"
        elif t.startswith("51") and len(t) > 8:
            return "🇵🇪 Perú (+51)"
        elif t.startswith("54") and len(t) > 9:
            return "🇦🇷 Argentina (+54)"
        elif t.startswith("56") and len(t) > 8:
            return "🇨🇱 Chile (+56)"
        elif t.startswith("57") and len(t) > 9:
            return "🇨🇴 Colombia (+57)"
        elif t.startswith("593") and len(t) > 8:
            return "🇪🇨 Ecuador (+593)"
        elif t.startswith("52") and len(t) > 9:
            return "🇲🇽 México (+52)"
        elif t.startswith("34") and len(t) > 8:
            return "🇪🇸 España (+34)"

    return "🇧🇴 Bolivia (+591)"


# ==========================================
# VALIDACIÓN DE LICENCIA PREVIA AL ARRANQUE
# ==========================================
def validar_licencia_o_salir():
    """Valida la licencia por hardware antes de levantar la interfaz gráfica."""
    lm = LicenseManager()
    if not lm.validate_license():
        root_temp = tk.Tk()
        root_temp.withdraw()
        hw_id = get_hardware_id() or "No detectado"
        messagebox.showerror(
            "Acceso Denegado - Licencia Inválida",
            "La licencia de este software no es válida o no corresponde a este equipo.\n\n"
            f"Hardware ID del equipo: {hw_id}\n\n"
            "El sistema se cerrará. Por favor, proporcione este ID al desarrollador para obtener una licencia autorizada."
        )
        root_temp.destroy()
        sys.exit(1)


# ==========================================
# APLICACIÓN PRINCIPAL (CustomTkinter)
# ==========================================
class MinimarketApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        # Inicialización de servicios
        self.db = MinimarketDB("minimarket.db")
        self.automator = Automator(self.db)
        self.whatsapp_gestor: Optional[GestorWhatsApp] = None
        self._whatsapp_monitoreo_activo: bool = False
        self._whatsapp_esta_conectado: bool = False
        self._smtp_esta_conectado: bool = False
        self._smtp_email_conectado: str = ""
        
        # Cola personalizada de destinatarios con asignación de reglas
        self.cola_destinatarios_seleccionados: List[Dict[str, Any]] = []
        self._widgets_destinatarios: List[Dict[str, Any]] = []
        self.cliente_seleccionado_id: Optional[int] = None
        self._cache_clientes_reglas: Dict[str, Dict[str, Any]] = {}
        self._cache_reglas_reglas: Dict[str, Dict[str, Any]] = {}

        # Configuración de Ventana
        self.title("Minimarket Messaging & Marketing Suite")
        self.geometry("1220x760")
        self.minsize(1020, 660)

        # Manejador de cierre de ventana para evitar procesos huérfanos
        self.protocol("WM_DELETE_WINDOW", self.al_cerrar_aplicacion)

        # Configuración de apariencia (Tema Oscuro por defecto)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        # Layout en Grid (2 columnas: menú lateral y contenedor de vistas)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        # Crear componentes
        self.crear_menu_lateral()
        self.crear_vistas()

        # Iniciar verificación de conexión inicial y chequeo periódico de WhatsApp y SMTP
        self.after(500, self._auto_verificar_conexion_inicial_whatsapp)
        self.after(800, self._auto_verificar_conexion_inicial_smtp)
        self.after(16000, self.chequeo_periodico_whatsapp)

        # Mostrar primera vista por defecto
        self.seleccionar_vista("clientes")

    def al_cerrar_aplicacion(self):
        """Cierre ordenado de la aplicación y liberación de recursos de Selenium."""
        self._whatsapp_monitoreo_activo = False
        if self.whatsapp_gestor:
            try:
                self.whatsapp_gestor.cerrar()
            except Exception:
                pass
        self.destroy()

    def obtener_gestor_whatsapp(self) -> GestorWhatsApp:
        """Obtiene o inicializa de forma diferida el gestor de WhatsApp."""
        if self.whatsapp_gestor is None or self.whatsapp_gestor.driver is None:
            self.whatsapp_gestor = GestorWhatsApp()
        return self.whatsapp_gestor

    # ------------------------------------------------------------------
    # MENÚ LATERAL (Sidebar)
    # ------------------------------------------------------------------
    def crear_menu_lateral(self):
        self.sidebar_frame = ctk.CTkFrame(self, width=230, corner_radius=0)
        self.sidebar_frame.grid(row=0, column=0, sticky="nsew")
        self.sidebar_frame.grid_rowconfigure(7, weight=1)
        self.sidebar_frame.grid_columnconfigure(0, weight=1)

        # Título / Logo
        self.logo_label = ctk.CTkLabel(
            self.sidebar_frame, 
            text="🏪 MINIMARKET", 
            font=ctk.CTkFont(size=20, weight="bold"),
            text_color=("black", "white")
        )
        self.logo_label.grid(row=0, column=0, padx=20, pady=(25, 5))
        
        self.sub_logo = ctk.CTkLabel(
            self.sidebar_frame, 
            text="Sistema de Mensajería", 
            font=ctk.CTkFont(size=12),
            text_color=("gray30", "gray70")
        )
        self.sub_logo.grid(row=1, column=0, padx=20, pady=(0, 20))

        # Botones de navegación (Alineados a la izquierda para máxima estética y orden)
        self.btn_nav_clientes = ctk.CTkButton(
            self.sidebar_frame, 
            text="👥 Gestión Clientes", 
            font=ctk.CTkFont(size=13, weight="bold"),
            height=40,
            anchor="w",
            text_color=("black", "white"),
            command=lambda: self.seleccionar_vista("clientes")
        )
        self.btn_nav_clientes.grid(row=2, column=0, padx=15, pady=5, sticky="ew")

        self.btn_nav_smtp = ctk.CTkButton(
            self.sidebar_frame, 
            text="📧 Envíos Email (SMTP)", 
            font=ctk.CTkFont(size=13, weight="bold"),
            height=40,
            anchor="w",
            text_color=("black", "white"),
            command=lambda: self.seleccionar_vista("smtp")
        )
        self.btn_nav_smtp.grid(row=3, column=0, padx=15, pady=5, sticky="ew")

        # Nuevo Botón Específico: Envíos Manuales WhatsApp
        self.btn_nav_envios_wsp = ctk.CTkButton(
            self.sidebar_frame, 
            text="💬 Envíos WhatsApp", 
            font=ctk.CTkFont(size=13, weight="bold"),
            height=40,
            anchor="w",
            text_color=("black", "white"),
            command=lambda: self.seleccionar_vista("envios_wsp")
        )
        self.btn_nav_envios_wsp.grid(row=4, column=0, padx=15, pady=5, sticky="ew")

        self.btn_nav_reglas = ctk.CTkButton(
            self.sidebar_frame, 
            text="⚡ Automatización", 
            font=ctk.CTkFont(size=13, weight="bold"),
            height=40,
            anchor="w",
            text_color=("black", "white"),
            command=lambda: self.seleccionar_vista("reglas")
        )
        self.btn_nav_reglas.grid(row=5, column=0, padx=15, pady=5, sticky="ew")

        # Botón de Conexión / Vinculación WhatsApp
        self.btn_nav_whatsapp = ctk.CTkButton(
            self.sidebar_frame, 
            text="📲 Conexión WhatsApp", 
            font=ctk.CTkFont(size=13, weight="bold"),
            height=40,
            anchor="w",
            text_color=("black", "white"),
            command=lambda: self.seleccionar_vista("whatsapp")
        )
        self.btn_nav_whatsapp.grid(row=6, column=0, padx=15, pady=5, sticky="ew")

        # Indicador Permanente de Conexión SMTP en el Sidebar (Arriba de WhatsApp)
        self.lbl_sidebar_smtp_status = ctk.CTkLabel(
            self.sidebar_frame,
            text="🔴 SMTP: Desconectado",
            font=ctk.CTkFont(size=11, weight="bold"),
            anchor="w",
            text_color="#e74c3c"
        )
        self.lbl_sidebar_smtp_status.grid(row=8, column=0, padx=20, pady=(10, 1), sticky="w")

        # Sub-etiqueta con correo de cuenta SMTP en el Sidebar
        self.lbl_sidebar_smtp_cuenta = ctk.CTkLabel(
            self.sidebar_frame,
            text="",
            font=ctk.CTkFont(size=10, weight="bold"),
            anchor="w",
            justify="left",
            wraplength=185,
            text_color=("#155724", "#4ade80")
        )
        self.lbl_sidebar_smtp_cuenta.grid(row=9, column=0, padx=20, pady=(0, 2), sticky="w")

        # Indicador Permanente de WhatsApp en el Sidebar (Alineado a la izquierda)
        self.lbl_sidebar_wsp_status = ctk.CTkLabel(
            self.sidebar_frame,
            text="🟡 WhatsApp: Comprobando...",
            font=ctk.CTkFont(size=11, weight="bold"),
            anchor="w",
            text_color="#f39c12"
        )
        self.lbl_sidebar_wsp_status.grid(row=10, column=0, padx=20, pady=(4, 1), sticky="w")

        # Sub-etiqueta con número y nombre de cuenta conectada en el Sidebar
        self.lbl_sidebar_wsp_cuenta = ctk.CTkLabel(
            self.sidebar_frame,
            text="",
            font=ctk.CTkFont(size=10, weight="bold"),
            anchor="w",
            justify="left",
            text_color=("#155724", "#4ade80")
        )
        self.lbl_sidebar_wsp_cuenta.grid(row=11, column=0, padx=20, pady=(0, 4), sticky="w")

        # Pie del menú lateral (Estado de Licencia y Modo Oscuro)
        self.license_badge = ctk.CTkLabel(
            self.sidebar_frame,
            text="🟢 Licencia Activa (Hardware OK)",
            font=ctk.CTkFont(size=11, weight="bold"),
            anchor="w",
            text_color="#2ecc71"
        )
        self.license_badge.grid(row=12, column=0, padx=20, pady=(4, 6), sticky="w")

        self.theme_switch = ctk.CTkOptionMenu(
            self.sidebar_frame,
            values=["Oscuro", "Claro"],
            command=self.cambiar_tema
        )
        self.theme_switch.set("Oscuro")
        self.theme_switch.grid(row=13, column=0, padx=15, pady=(0, 18), sticky="ew")

    def cambiar_tema(self, modo: str):
        ctk.set_appearance_mode("dark" if modo == "Oscuro" else "light")

    # ------------------------------------------------------------------
    # NAVEGACIÓN Y CONTENEDOR DE VISTAS
    # ------------------------------------------------------------------
    def crear_vistas(self):
        self.container = ctk.CTkFrame(self, corner_radius=0, fg_color="transparent")
        self.container.grid(row=0, column=1, sticky="nsew", padx=20, pady=20)
        self.container.grid_rowconfigure(0, weight=1)
        self.container.grid_columnconfigure(0, weight=1)

        # Construir cada una de las vistas
        self.vista_clientes = self.construir_vista_clientes()
        self.vista_smtp = self.construir_vista_smtp()
        self.vista_envios_wsp = self.construir_vista_envios_whatsapp()
        self.vista_reglas = self.construir_vista_reglas()
        self.vista_whatsapp = self.construir_vista_whatsapp()

    def seleccionar_vista(self, nombre: str):
        # Ocultar todas las vistas
        self.vista_clientes.grid_forget()
        self.vista_smtp.grid_forget()
        self.vista_envios_wsp.grid_forget()
        self.vista_reglas.grid_forget()
        self.vista_whatsapp.grid_forget()

        # Restablecer colores de botones inactivos
        color_inactivo = "transparent"
        self.btn_nav_clientes.configure(fg_color=color_inactivo, text_color=("black", "white"))
        self.btn_nav_smtp.configure(fg_color=color_inactivo, text_color=("black", "white"))
        self.btn_nav_envios_wsp.configure(fg_color=color_inactivo, text_color=("black", "white"))
        self.btn_nav_reglas.configure(fg_color=color_inactivo, text_color=("black", "white"))
        self.btn_nav_whatsapp.configure(fg_color=color_inactivo, text_color=("black", "white"))

        # Mostrar la vista solicitada y resaltar su botón
        if nombre == "clientes":
            self.vista_clientes.grid(row=0, column=0, sticky="nsew")
            self.btn_nav_clientes.configure(fg_color=["#3a7ebf", "#1f538d"], text_color="white")
            self.cargar_lista_clientes()
        elif nombre == "smtp":
            self.vista_smtp.grid(row=0, column=0, sticky="nsew")
            self.btn_nav_smtp.configure(fg_color=["#3a7ebf", "#1f538d"], text_color="white")
            self.actualizar_combobox_clientes()
        elif nombre == "envios_wsp":
            self.vista_envios_wsp.grid(row=0, column=0, sticky="nsew")
            self.btn_nav_envios_wsp.configure(fg_color=["#25D366", "#128C7E"], text_color="white")
            self.actualizar_combobox_clientes_whatsapp()
            self.verificar_estado_whatsapp_asincrono()
        elif nombre == "reglas":
            self.vista_reglas.grid(row=0, column=0, sticky="nsew")
            self.btn_nav_reglas.configure(fg_color=["#3a7ebf", "#1f538d"], text_color="white")
            self.actualizar_selectores_destinatarios()
            self.renderizar_cola_destinatarios()
        elif nombre == "whatsapp":
            self.vista_whatsapp.grid(row=0, column=0, sticky="nsew")
            self.btn_nav_whatsapp.configure(fg_color=["#25D366", "#128C7E"], text_color="white")
            self.verificar_estado_whatsapp_asincrono()

    # ------------------------------------------------------------------
    # VISTA 1: GESTIÓN DE CLIENTES
    # ------------------------------------------------------------------
    def construir_vista_clientes(self) -> ctk.CTkFrame:
        frame = ctk.CTkFrame(self.container, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_columnconfigure(1, weight=2)
        frame.grid_rowconfigure(0, weight=1)

        # Columna Izquierda: Formulario
        form_frame = ctk.CTkFrame(frame)
        form_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10), pady=0)
        
        # Cabecera: Título con indicador de modo al lado derecho entre paréntesis
        header_form = ctk.CTkFrame(form_frame, fg_color="transparent")
        header_form.pack(fill="x", padx=15, pady=(10, 4))

        ctk.CTkLabel(
            header_form, 
            text="Datos del Cliente", 
            font=ctk.CTkFont(size=17, weight="bold"), 
            text_color=("black", "white")
        ).pack(side="left")

        self.lbl_cliente_modo = ctk.CTkLabel(
            header_form, 
            text=" (nuevo registro)", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("gray40", "gray60")
        )
        self.lbl_cliente_modo.pack(side="left", padx=(4, 0))

        self.entry_cliente_nombre = ctk.CTkEntry(form_frame, placeholder_text="Nombre Completo", height=32)
        self.entry_cliente_nombre.pack(fill="x", padx=15, pady=3)

        self.entry_cliente_correo = ctk.CTkEntry(form_frame, placeholder_text="Correo Electrónico", height=32)
        self.entry_cliente_correo.pack(fill="x", padx=15, pady=3)

        ctk.CTkLabel(form_frame, text="Teléfono / WhatsApp:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(3, 1))
        frame_tel_cli = ctk.CTkFrame(form_frame, fg_color="transparent")
        frame_tel_cli.pack(fill="x", padx=15, pady=(0, 3))
        frame_tel_cli.grid_columnconfigure(1, weight=1)

        self.combo_cliente_pais = ctk.CTkComboBox(
            frame_tel_cli,
            values=list(PAISES_WHATSAPP.keys()),
            width=145,
            height=32,
            state="readonly"
        )
        self.combo_cliente_pais.set("🇧🇴 Bolivia (+591)")
        self.combo_cliente_pais.grid(row=0, column=0, padx=(0, 6), sticky="w")

        self.entry_cliente_telefono = ctk.CTkEntry(
            frame_tel_cli, 
            placeholder_text="Ej: 72465746", 
            height=32
        )
        self.entry_cliente_telefono.grid(row=0, column=1, sticky="ew")

        ctk.CTkLabel(form_frame, text="Estado del Cliente:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(3, 1))
        self.combo_cliente_estado = ctk.CTkComboBox(form_frame, values=["Activo", "Inactivo"], height=30, state="readonly")
        self.combo_cliente_estado.set("Activo")
        self.combo_cliente_estado.pack(fill="x", padx=15, pady=2)

        ctk.CTkLabel(form_frame, text="Categoría del Cliente:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(3, 1))
        self.combo_cliente_categoria = ctk.CTkComboBox(form_frame, values=CATEGORIAS_CLIENTE, height=30, state="readonly")
        self.combo_cliente_categoria.set("🥉 Bronce")
        self.combo_cliente_categoria.pack(fill="x", padx=15, pady=2)

        # Botones de Acción
        self.btn_guardar_cliente = ctk.CTkButton(
            form_frame, text="➕ Agregar Cliente", height=30, fg_color="#27ae60", hover_color="#219955",
            command=self.guardar_cliente
        )
        self.btn_guardar_cliente.pack(fill="x", padx=15, pady=(8, 3))

        self.btn_actualizar_cliente = ctk.CTkButton(
            form_frame, text="✏️ Actualizar Seleccionado", height=28, fg_color="#2980b9", hover_color="#1f618d",
            command=self.actualizar_cliente
        )
        self.btn_actualizar_cliente.pack(fill="x", padx=15, pady=2)

        self.btn_eliminar_cliente = ctk.CTkButton(
            form_frame, text="🗑️ Eliminar Cliente", height=28, fg_color="#c0392b", hover_color="#962d22",
            command=self.eliminar_cliente
        )
        self.btn_eliminar_cliente.pack(fill="x", padx=15, pady=2)

        self.btn_limpiar_cliente = ctk.CTkButton(
            form_frame, text="🧹 Limpiar Campos", height=28, fg_color="gray40", hover_color="gray30",
            command=self.limpiar_formulario_cliente
        )
        self.btn_limpiar_cliente.pack(fill="x", padx=15, pady=2)

        # Línea de separación estética
        sep_excel = ctk.CTkFrame(form_frame, height=2, fg_color=["gray75", "gray35"])
        sep_excel.pack(fill="x", padx=15, pady=(8, 4))

        ctk.CTkLabel(
            form_frame,
            text="📊 Importación Masiva de Clientes:",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w", padx=15, pady=(0, 2))

        # Botón 1: Descargar Plantilla Excel
        self.btn_plantilla_excel = ctk.CTkButton(
            form_frame,
            text="📥 Descargar Plantilla Excel",
            fg_color="#16a085",
            hover_color="#138d75",
            height=30,
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.generar_plantilla_excel
        )
        self.btn_plantilla_excel.pack(fill="x", padx=15, pady=(2, 3))

        # Botón 2: Cargar Clientes desde Excel
        self.btn_cargar_excel = ctk.CTkButton(
            form_frame,
            text="📂 Cargar Clientes (Excel / CSV)",
            fg_color="#8e44ad",
            hover_color="#732d91",
            height=30,
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.importar_clientes_excel
        )
        self.btn_cargar_excel.pack(fill="x", padx=15, pady=(2, 10))

        # Columna Derecha: Tabla/Lista de Clientes
        lista_frame = ctk.CTkFrame(frame)
        lista_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0), pady=0)
        lista_frame.grid_rowconfigure(1, weight=1)
        lista_frame.grid_columnconfigure(0, weight=1)

        # Cabecera de lista con Botón Refrescar y Botón Buscar Cliente
        header_frame = ctk.CTkFrame(lista_frame, fg_color="transparent")
        header_frame.grid(row=0, column=0, sticky="ew", padx=15, pady=10)
        ctk.CTkLabel(header_frame, text="Clientes Registrados", font=ctk.CTkFont(size=18, weight="bold"), text_color=("black", "white")).pack(side="left")
        
        ctk.CTkButton(header_frame, text="🔄 Refrescar", width=90, command=self.cargar_lista_clientes).pack(side="right")
        
        self.btn_buscar_cliente = ctk.CTkButton(
            header_frame, 
            text="🔍 Buscar Cliente", 
            width=125, 
            fg_color="#2980b9", 
            hover_color="#1f618d", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            command=self.abrir_modal_buscar_cliente
        )
        self.btn_buscar_cliente.pack(side="right", padx=(0, 8))

        # Scrollable Frame con tarjetas de clientes
        self.scroll_clientes = ctk.CTkScrollableFrame(lista_frame)
        self.scroll_clientes.grid(row=1, column=0, sticky="nsew", padx=15, pady=(0, 15))

        return frame

    def cargar_lista_clientes(self):
        for widget in self.scroll_clientes.winfo_children():
            widget.destroy()

        clientes = self.db.get_all_clientes()
        if not clientes:
            ctk.CTkLabel(self.scroll_clientes, text="No hay clientes registrados.", text_color=("gray40", "gray60")).pack(pady=20)
            return

        for c in clientes:
            card = ctk.CTkFrame(self.scroll_clientes)
            card.pack(fill="x", padx=5, pady=4)
            card.grid_columnconfigure(0, weight=1)

            cat = c.get("categoria") or "🥉 Bronce"
            est = c.get("estado") or "Activo"
            if est not in ["Activo", "Inactivo"]:
                est = "Activo"

            info_text = f"ID: {c['id']} | {c['nombre']}  •  {c['correo'] or 'Sin correo'}  •  Tel: {c['telefono'] or 'Sin tel'}"
            lbl = ctk.CTkLabel(card, text=info_text, font=ctk.CTkFont(size=13, weight="bold"), anchor="w", text_color=("black", "white"))
            lbl.grid(row=0, column=0, padx=10, pady=8, sticky="w")

            # Badge de Categoría
            cat_lbl = ctk.CTkLabel(card, text=f" {cat} ", font=ctk.CTkFont(size=11, weight="bold"), text_color="#3498db")
            cat_lbl.grid(row=0, column=1, padx=4, pady=8)

            # Badge de Estado [Activo] o [Inactivo]
            badge_color = "#27ae60" if est == "Activo" else "#e74c3c"
            estado_lbl = ctk.CTkLabel(card, text=f" [{est}] ", text_color=badge_color, font=ctk.CTkFont(size=12, weight="bold"))
            estado_lbl.grid(row=0, column=2, padx=4, pady=8)

            btn_sel = ctk.CTkButton(card, text="Cargar", width=65, command=lambda cli=c: self.seleccionar_cliente_para_edicion(cli))
            btn_sel.grid(row=0, column=3, padx=(5, 10), pady=6)

    def abrir_modal_buscar_cliente(self):
        """Abre una ventana emergente modal para buscar clientes en tiempo real por nombre, teléfono o correo."""
        modal = ctk.CTkToplevel(self)
        modal.title("🔍 Buscar Cliente - Minimarket")
        modal.geometry("740x540")
        modal.minsize(620, 420)
        modal.transient(self)
        modal.grab_set()

        # Centrar ventana modal respecto a la ventana principal
        modal.update_idletasks()
        try:
            x = self.winfo_x() + (self.winfo_width() - 740) // 2
            y = self.winfo_y() + (self.winfo_height() - 540) // 2
            modal.geometry(f"+{max(0, x)}+{max(0, y)}")
        except Exception:
            pass

        # Cabecera del modal
        top_frame = ctk.CTkFrame(modal, fg_color="transparent")
        top_frame.pack(fill="x", padx=20, pady=(15, 8))

        ctk.CTkLabel(
            top_frame,
            text="🔍 Búsqueda Rápida de Clientes",
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w")

        ctk.CTkLabel(
            top_frame,
            text="Escriba el nombre, teléfono o correo electrónico para filtrar al instante y cargar los datos para modificar o inactivar.",
            font=ctk.CTkFont(size=12),
            text_color=("gray40", "gray70")
        ).pack(anchor="w", pady=(2, 6))

        # Campo de Búsqueda
        entry_buscar = ctk.CTkEntry(
            modal,
            placeholder_text="🔎 Buscar por Nombre, Teléfono o Correo Electrónico...",
            height=38,
            font=ctk.CTkFont(size=13)
        )
        entry_buscar.pack(fill="x", padx=20, pady=(0, 6))
        entry_buscar.focus_set()

        # Etiqueta de conteo
        lbl_conteo = ctk.CTkLabel(
            modal,
            text="Cargando clientes...",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("gray40", "gray60")
        )
        lbl_conteo.pack(anchor="w", padx=22, pady=(0, 6))

        # Contenedor con scroll para los resultados
        scroll_modal = ctk.CTkScrollableFrame(modal)
        scroll_modal.pack(fill="both", expand=True, padx=20, pady=(0, 10))

        # Cargar todos los clientes de la base de datos
        todos_clientes = self.db.get_all_clientes()

        def cargar_en_formulario_y_cerrar(cliente_datos):
            self.seleccionar_cliente_para_edicion(cliente_datos)
            modal.destroy()

        def renderizar_resultados(query: str = ""):
            for w in scroll_modal.winfo_children():
                w.destroy()

            q = query.strip().lower()
            if not q:
                filtrados = todos_clientes
            else:
                filtrados = []
                for c in todos_clientes:
                    nom = str(c.get("nombre") or "").lower()
                    cor = str(c.get("correo") or "").lower()
                    tel = str(c.get("telefono") or "").lower()
                    cid = str(c.get("id") or "")
                    if q in nom or q in cor or q in tel or q == cid:
                        filtrados.append(c)

            lbl_conteo.configure(text=f"Resultados encontrados: {len(filtrados)} de {len(todos_clientes)} clientes registrados")

            if not filtrados:
                ctk.CTkLabel(
                    scroll_modal,
                    text="❌ No se encontraron clientes que coincidan con la búsqueda.",
                    font=ctk.CTkFont(size=13),
                    text_color=("gray40", "gray60")
                ).pack(pady=30)
                return

            for c in filtrados:
                item_frame = ctk.CTkFrame(scroll_modal)
                item_frame.pack(fill="x", padx=4, pady=4)
                item_frame.grid_columnconfigure(0, weight=1)

                cat = c.get("categoria") or "🥉 Bronce"
                est = c.get("estado") or "Activo"
                if est not in ["Activo", "Inactivo"]:
                    est = "Activo"

                info = f"ID: {c['id']} | {c['nombre']}  •  {c['correo'] or 'Sin correo'}  •  Tel: {c['telefono'] or 'Sin tel'}"
                lbl_info = ctk.CTkLabel(
                    item_frame,
                    text=info,
                    font=ctk.CTkFont(size=12, weight="bold"),
                    anchor="w",
                    text_color=("black", "white")
                )
                lbl_info.grid(row=0, column=0, padx=10, pady=8, sticky="w")

                # Badge Categoría
                lbl_cat = ctk.CTkLabel(
                    item_frame,
                    text=f" {cat} ",
                    font=ctk.CTkFont(size=11, weight="bold"),
                    text_color="#3498db"
                )
                lbl_cat.grid(row=0, column=1, padx=4, pady=8)

                # Badge Estado
                color_est = "#27ae60" if est == "Activo" else "#e74c3c"
                lbl_est = ctk.CTkLabel(
                    item_frame,
                    text=f" [{est}] ",
                    font=ctk.CTkFont(size=11, weight="bold"),
                    text_color=color_est
                )
                lbl_est.grid(row=0, column=2, padx=4, pady=8)

                # Botón Cargar
                btn_cargar = ctk.CTkButton(
                    item_frame,
                    text="Cargar ✏️",
                    width=75,
                    fg_color="#2980b9",
                    hover_color="#1f618d",
                    command=lambda cli=c: cargar_en_formulario_y_cerrar(cli)
                )
                btn_cargar.grid(row=0, column=3, padx=(6, 10), pady=6)

        entry_buscar.bind("<KeyRelease>", lambda event: renderizar_resultados(entry_buscar.get()))
        renderizar_resultados("")

        # Botón inferior para cerrar
        bot_frame = ctk.CTkFrame(modal, fg_color="transparent")
        bot_frame.pack(fill="x", padx=20, pady=(0, 12))
        ctk.CTkButton(
            bot_frame,
            text="Cerrar",
            width=100,
            fg_color="gray40",
            hover_color="gray30",
            command=modal.destroy
        ).pack(side="right")

    def seleccionar_cliente_para_edicion(self, cliente: dict):
        self.cliente_seleccionado_id = cliente["id"]
        if hasattr(self, "lbl_cliente_modo"):
            self.lbl_cliente_modo.configure(
                text=f" (✏️ editando #{cliente['id']}: {cliente['nombre']})",
                text_color="#3498db"
            )

        self.entry_cliente_nombre.delete(0, "end")
        self.entry_cliente_nombre.insert(0, cliente["nombre"])

        self.entry_cliente_correo.delete(0, "end")
        self.entry_cliente_correo.insert(0, cliente["correo"] or "")

        tel_raw = (cliente.get("telefono") or "").strip()
        pais_detectado = normalizar_pais("", tel_raw)
        prefijo = PAISES_WHATSAPP.get(pais_detectado, "+591").replace("+", "").strip()

        solo_digitos = re.sub(r'\D', '', tel_raw)
        if prefijo and solo_digitos.startswith(prefijo) and len(solo_digitos) > len(prefijo):
            num_local = solo_digitos[len(prefijo):]
        else:
            num_local = solo_digitos if solo_digitos else tel_raw

        if hasattr(self, "combo_cliente_pais"):
            self.combo_cliente_pais.set(pais_detectado)

        self.entry_cliente_telefono.delete(0, "end")
        self.entry_cliente_telefono.insert(0, num_local)

        est = cliente.get("estado") or "Activo"
        if est not in ["Activo", "Inactivo"]:
            est = "Activo"
        self.combo_cliente_estado.set(est)

        cat = cliente.get("categoria") or "🥉 Bronce"
        if cat in CATEGORIAS_CLIENTE:
            self.combo_cliente_categoria.set(cat)
        else:
            self.combo_cliente_categoria.set(normalizar_categoria(cat))

    def limpiar_formulario_cliente(self):
        self.cliente_seleccionado_id = None
        if hasattr(self, "lbl_cliente_modo"):
            self.lbl_cliente_modo.configure(
                text=" (nuevo registro)",
                text_color=("gray40", "gray60")
            )
        self.entry_cliente_nombre.delete(0, "end")
        self.entry_cliente_correo.delete(0, "end")
        if hasattr(self, "combo_cliente_pais"):
            self.combo_cliente_pais.set("🇧🇴 Bolivia (+591)")
        self.entry_cliente_telefono.delete(0, "end")
        self.combo_cliente_estado.set("Activo")
        self.combo_cliente_categoria.set("🥉 Bronce")

    def guardar_cliente(self):
        nombre = self.entry_cliente_nombre.get().strip()
        correo = self.entry_cliente_correo.get().strip()
        
        pais_sel = self.combo_cliente_pais.get() if hasattr(self, "combo_cliente_pais") else "🇧🇴 Bolivia (+591)"
        prefijo = PAISES_WHATSAPP.get(pais_sel, "+591").strip()
        num_raw = self.entry_cliente_telefono.get().strip()
        solo_dig = re.sub(r'\D', '', num_raw)

        if solo_dig:
            prefijo_num = prefijo.replace("+", "")
            if prefijo and not solo_dig.startswith(prefijo_num):
                telefono = f"{prefijo} {solo_dig}"
            elif prefijo and solo_dig.startswith(prefijo_num):
                telefono = f"+{prefijo_num} {solo_dig[len(prefijo_num):]}"
            else:
                telefono = solo_dig
        else:
            telefono = ""

        estado = normalizar_estado(self.combo_cliente_estado.get())
        categoria = normalizar_categoria(self.combo_cliente_categoria.get())

        if not nombre:
            messagebox.showwarning("Campo Requerido", "Debe ingresar al menos el Nombre del cliente.")
            return

        cid = self.db.create_cliente(nombre=nombre, correo=correo, telefono=telefono, estado=estado, categoria=categoria)
        if cid:
            messagebox.showinfo("Éxito", f"Cliente #{cid} registrado exitosamente como {categoria} [{estado}].")
            self.limpiar_formulario_cliente()
            self.cargar_lista_clientes()
            if hasattr(self, "actualizar_combobox_clientes"):
                self.actualizar_combobox_clientes()
            if hasattr(self, "actualizar_combobox_clientes_whatsapp"):
                self.actualizar_combobox_clientes_whatsapp()
            if hasattr(self, "actualizar_selectores_destinatarios"):
                self.actualizar_selectores_destinatarios()
        else:
            messagebox.showerror("Error", "No se pudo registrar el cliente en la base de datos.")

    def actualizar_cliente(self):
        if not getattr(self, "cliente_seleccionado_id", None):
            messagebox.showwarning("Seleccione Cliente", "Primero seleccione un cliente de la lista para editar (pulse 'Cargar' o use 'Buscar Cliente').")
            return

        cid = self.cliente_seleccionado_id
        nombre = self.entry_cliente_nombre.get().strip()
        correo = self.entry_cliente_correo.get().strip()
        
        pais_sel = self.combo_cliente_pais.get() if hasattr(self, "combo_cliente_pais") else "🇧🇴 Bolivia (+591)"
        prefijo = PAISES_WHATSAPP.get(pais_sel, "+591").strip()
        num_raw = self.entry_cliente_telefono.get().strip()
        solo_dig = re.sub(r'\D', '', num_raw)

        if solo_dig:
            prefijo_num = prefijo.replace("+", "")
            if prefijo and not solo_dig.startswith(prefijo_num):
                telefono = f"{prefijo} {solo_dig}"
            elif prefijo and solo_dig.startswith(prefijo_num):
                telefono = f"+{prefijo_num} {solo_dig[len(prefijo_num):]}"
            else:
                telefono = solo_dig
        else:
            telefono = ""

        estado = normalizar_estado(self.combo_cliente_estado.get())
        categoria = normalizar_categoria(self.combo_cliente_categoria.get())

        if not nombre:
            messagebox.showwarning("Campo Requerido", "El nombre del cliente no puede estar vacío.")
            return

        ok = self.db.update_cliente(cliente_id=cid, nombre=nombre, correo=correo, telefono=telefono, estado=estado, categoria=categoria)
        if ok:
            messagebox.showinfo("Éxito", f"Cliente #{cid} actualizado correctamente:\n• Categoría: {categoria}\n• Estado: {estado}")
            self.limpiar_formulario_cliente()
            self.cargar_lista_clientes()
            if hasattr(self, "actualizar_combobox_clientes"):
                self.actualizar_combobox_clientes()
            if hasattr(self, "actualizar_combobox_clientes_whatsapp"):
                self.actualizar_combobox_clientes_whatsapp()
            if hasattr(self, "actualizar_selectores_destinatarios"):
                self.actualizar_selectores_destinatarios()
        else:
            messagebox.showerror("Error", "No se pudo actualizar el cliente.")

    def eliminar_cliente(self):
        if not getattr(self, "cliente_seleccionado_id", None):
            messagebox.showwarning("Seleccione Cliente", "Seleccione un cliente para eliminar (pulse 'Cargar').")
            return

        cid = self.cliente_seleccionado_id
        if messagebox.askyesno("Confirmar Eliminación", f"¿Está seguro de eliminar al cliente #{cid}?"):
            if self.db.delete_cliente(cid):
                messagebox.showinfo("Eliminado", "Cliente eliminado correctamente.")
                self.limpiar_formulario_cliente()
                self.cargar_lista_clientes()
            else:
                messagebox.showerror("Error", "No se pudo eliminar el cliente.")

    def generar_plantilla_excel(self):
        """Genera y descarga un archivo Excel (.xlsx o .csv) con la estructura y datos de ejemplo para carga masiva."""
        ruta_archivo = filedialog.asksaveasfilename(
            title="Guardar Plantilla de Clientes Excel",
            defaultextension=".xlsx",
            initialfile="plantilla_clientes_minimarket.xlsx",
            filetypes=[("Libro de Excel (*.xlsx)", "*.xlsx"), ("Archivo CSV (*.csv)", "*.csv")]
        )
        if not ruta_archivo:
            return

        try:
            if ruta_archivo.lower().endswith(".csv"):
                with open(ruta_archivo, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    writer.writerow(["Nombre Completo", "Correo Electrónico", "País", "Teléfono", "Estado", "Categoría"])
                    writer.writerow(["Juan Pérez", "juan.perez@gmail.com", "bolivia", "72465746", "activo", "bronce"])
                    writer.writerow(["Maribel Flores", "maribel@hotmail.com", "bolivia", "67220995", "activo", "plata"])
                    writer.writerow(["Carlos Mendoza", "carlos.m@empresa.com", "peru", "912345678", "activo", "oro"])
                    writer.writerow(["Ana Gómez", "ana.gomez@gmail.com", "argentina", "1198765432", "inactivo", "platino"])
                    writer.writerow(["Roberto Silva", "roberto.vip@gmail.com", "bolivia", "76543210", "activo", "diamante/vip"])
            else:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Clientes"

                headers = ["Nombre Completo", "Correo Electrónico", "País", "Teléfono", "Estado", "Categoría"]
                ws.append(headers)

                # Estilo profesional para la fila de encabezados
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                align_center = Alignment(horizontal="center", vertical="center")

                for col_num in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = align_center

                # Filas de ejemplo limpias en minúsculas sin símbolos para compatibilidad total con Excel
                ejemplos = [
                    ["Juan Pérez", "juan.perez@gmail.com", "bolivia", "72465746", "activo", "bronce"],
                    ["Maribel Flores", "maribel@hotmail.com", "bolivia", "67220995", "activo", "plata"],
                    ["Carlos Mendoza", "carlos.m@empresa.com", "peru", "912345678", "activo", "oro"],
                    ["Ana Gómez", "ana.gomez@gmail.com", "argentina", "1198765432", "inactivo", "platino"],
                    ["Roberto Silva", "roberto.vip@gmail.com", "bolivia", "76543210", "activo", "diamante/vip"]
                ]
                for fila in ejemplos:
                    ws.append(fila)

                # Ancho de columnas óptimo
                anchos = {"A": 26, "B": 30, "C": 14, "D": 18, "E": 14, "F": 20}
                for col_letter, width in anchos.items():
                    ws.column_dimensions[col_letter].width = width

                wb.save(ruta_archivo)

            abrir = messagebox.askyesno(
                "Plantilla Generada Exitosamente",
                f"✅ ¡Plantilla de Excel guardada con éxito!\n\n"
                f"Ubicación:\n{ruta_archivo}\n\n"
                f"La plantilla está configurada en minúsculas y sin símbolos para facilitar la migración:\n"
                f"• País: bolivia | peru | argentina | chile | colombia | etc. (o dejar vacío)\n"
                f"• Estado: activo | inactivo\n"
                f"• Categoría: bronce | plata | oro | platino | diamante/vip\n\n"
                f"¿Desea abrir el archivo en Excel ahora mismo para completarlo?"
            )
            if abrir:
                try:
                    os.startfile(ruta_archivo)
                except Exception:
                    pass

        except Exception as e:
            messagebox.showerror("Error al Guardar Plantilla", f"No se pudo generar el archivo de plantilla:\n{e}")

    def importar_clientes_excel(self):
        """Carga clientes masivamente desde un archivo Excel (.xlsx, .xls) o CSV hacia la base de datos."""
        ruta_archivo = filedialog.askopenfilename(
            title="Seleccionar Archivo de Clientes (Excel o CSV)",
            filetypes=[
                ("Archivos Excel y CSV", "*.xlsx;*.xls;*.csv"),
                ("Libro de Excel (*.xlsx)", "*.xlsx"),
                ("Archivo CSV (*.csv)", "*.csv")
            ]
        )
        if not ruta_archivo:
            return

        try:
            if ruta_archivo.lower().endswith(".csv"):
                lineas = []
                for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
                    try:
                        with open(ruta_archivo, "r", encoding=enc) as f:
                            muestra = f.read(2048)
                            f.seek(0)
                            delimitador = ","
                            if ";" in muestra and muestra.count(";") > muestra.count(","):
                                delimitador = ";"
                            elif "\t" in muestra and muestra.count("\t") > muestra.count(","):
                                delimitador = "\t"
                            reader = csv.reader(f, delimiter=delimitador)
                            lineas = list(reader)
                        break
                    except Exception:
                        continue

                if not lineas:
                    messagebox.showerror("Archivo Vacío", "No se pudo leer el archivo CSV o no contiene datos.")
                    return

                cabeceras = [str(c).strip().lower() for c in lineas[0]]
                filas_datos = lineas[1:]

            else:
                import openpyxl
                wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
                ws = wb.active
                todas_filas = list(ws.iter_rows(values_only=True))
                if not todas_filas:
                    messagebox.showerror("Archivo Vacío", "La hoja de cálculo de Excel está vacía.")
                    return

                cabeceras = [str(c or "").strip().lower() for c in todas_filas[0]]
                filas_datos = todas_filas[1:]

            # Detección inteligente de columnas por palabras clave
            idx_nombre = -1
            idx_correo = -1
            idx_pais = -1
            idx_telefono = -1
            idx_estado = -1
            idx_categoria = -1

            for idx, col_name in enumerate(cabeceras):
                if any(k in col_name for k in ["categoria", "categoría", "segmento", "clasificacion", "clasificación", "nivel", "tipo"]):
                    idx_categoria = idx
                elif any(k in col_name for k in ["pais", "país", "codigo", "código", "prefijo", "cod", "nacion", "nación"]):
                    idx_pais = idx
                elif any(k in col_name for k in ["nombre", "cliente", "persona", "destinatario"]):
                    idx_nombre = idx
                elif any(k in col_name for k in ["correo", "email", "mail"]):
                    idx_correo = idx
                elif any(k in col_name for k in ["telefono", "teléfono", "celular", "tel", "movil", "móvil", "whatsapp"]):
                    idx_telefono = idx
                elif any(k in col_name for k in ["estado", "status", "condicion", "condición"]):
                    idx_estado = idx

            # Fallback por orden si no encontró coincidencia de encabezado
            if idx_nombre == -1 and len(cabeceras) >= 1: idx_nombre = 0
            if idx_correo == -1 and len(cabeceras) >= 2: idx_correo = 1
            if idx_pais == -1 and idx_telefono == -1:
                if len(cabeceras) >= 6:
                    idx_pais = 2
                    idx_telefono = 3
                    if idx_estado == -1: idx_estado = 4
                    if idx_categoria == -1: idx_categoria = 5
                else:
                    idx_telefono = 2
                    if idx_estado == -1 and len(cabeceras) >= 4: idx_estado = 3
                    if idx_categoria == -1 and len(cabeceras) >= 5: idx_categoria = 4
            elif idx_telefono == -1:
                if idx_pais != -1 and len(cabeceras) > idx_pais + 1:
                    idx_telefono = idx_pais + 1
                elif len(cabeceras) >= 3:
                    idx_telefono = 2

            # Obtener clientes actuales en base de datos para actualizar o no duplicar
            clientes_db = self.db.get_all_clientes()
            mapa_tel = {}
            mapa_nom = {}
            for cli in clientes_db:
                tel_clean = re.sub(r'\D', '', str(cli.get("telefono") or ""))
                if tel_clean:
                    mapa_tel[tel_clean] = cli
                nom_clean = (cli.get("nombre") or "").strip().lower()
                if nom_clean:
                    mapa_nom[nom_clean] = cli

            creados = 0
            actualizados = 0
            omitidos = 0

            for fila in filas_datos:
                if not fila or all(val is None or str(val).strip() == "" for val in fila):
                    continue

                nombre = str(fila[idx_nombre] or "").strip() if 0 <= idx_nombre < len(fila) else ""
                correo = str(fila[idx_correo] or "").strip() if 0 <= idx_correo < len(fila) else ""
                pais_raw = str(fila[idx_pais] or "").strip() if 0 <= idx_pais < len(fila) else ""
                tel_raw = str(fila[idx_telefono] or "").strip() if 0 <= idx_telefono < len(fila) else ""
                estado_raw = str(fila[idx_estado] or "").strip() if 0 <= idx_estado < len(fila) else "Activo"
                cat_raw = str(fila[idx_categoria] or "").strip() if 0 <= idx_categoria < len(fila) else "🥉 Bronce"

                if not nombre or nombre.lower() in ["nombre", "nombre completo", "cliente"]:
                    omitidos += 1
                    continue

                estado_norm = normalizar_estado(estado_raw)
                cat_norm = normalizar_categoria(cat_raw)

                # Normalizar país y dar formato internacional al teléfono
                pais_detectado = normalizar_pais(pais_raw, tel_raw)
                prefijo = PAISES_WHATSAPP.get(pais_detectado, "+591").strip()
                solo_digitos = re.sub(r'\D', '', tel_raw)

                if solo_digitos:
                    prefijo_num = prefijo.replace("+", "")
                    if prefijo and not solo_digitos.startswith(prefijo_num):
                        telefono_final = f"{prefijo} {solo_digitos}"
                    elif prefijo and solo_digitos.startswith(prefijo_num):
                        telefono_final = f"+{prefijo_num} {solo_digitos[len(prefijo_num):]}"
                    else:
                        telefono_final = solo_digitos
                else:
                    telefono_final = ""

                # Comprobar si ya existe por teléfono o por nombre exacto
                tel_clean = re.sub(r'\D', '', telefono_final)
                existente = None
                if tel_clean and tel_clean in mapa_tel:
                    existente = mapa_tel[tel_clean]
                elif nombre.lower() in mapa_nom:
                    existente = mapa_nom[nombre.lower()]

                if existente:
                    self.db.update_cliente(
                        cliente_id=existente["id"],
                        nombre=nombre,
                        correo=correo if correo else existente.get("correo"),
                        telefono=telefono_final if telefono_final else existente.get("telefono"),
                        estado=estado_norm,
                        categoria=cat_norm
                    )
                    actualizados += 1
                else:
                    cid = self.db.create_cliente(
                        nombre=nombre,
                        correo=correo,
                        telefono=telefono_final,
                        estado=estado_norm,
                        categoria=cat_norm
                    )
                    if cid:
                        creados += 1
                        nuevo_cli = {
                            "id": cid, 
                            "nombre": nombre, 
                            "correo": correo, 
                            "telefono": telefono_final, 
                            "estado": estado_norm,
                            "categoria": cat_norm
                        }
                        if tel_clean:
                            mapa_tel[tel_clean] = nuevo_cli
                        mapa_nom[nombre.lower()] = nuevo_cli
                    else:
                        omitidos += 1

            # Recargar y refrescar lista de clientes en la interfaz
            self.cargar_lista_clientes()
            if hasattr(self, "actualizar_combobox_clientes"):
                self.actualizar_combobox_clientes()
            if hasattr(self, "actualizar_combobox_clientes_whatsapp"):
                self.actualizar_combobox_clientes_whatsapp()
            if hasattr(self, "actualizar_selectores_destinatarios"):
                self.actualizar_selectores_destinatarios()

            messagebox.showinfo(
                "Importación Exitosa",
                f"📊 ¡Carga masiva finalizada con éxito!\n\n"
                f"• Nuevos clientes registrados: {creados}\n"
                f"• Clientes actualizados: {actualizados}\n"
                f"• Filas vacías u omitidas: {omitidos}\n\n"
                f"Total de clientes en el sistema: {len(self.db.get_all_clientes())}"
            )

        except Exception as e:
            messagebox.showerror("Error de Importación", f"Ocurrió un error al procesar el archivo:\n{e}")

    # ------------------------------------------------------------------
    # VISTA 2: ENVÍOS MANUALES (SMTP / WhatsApp)
    # ------------------------------------------------------------------
    def construir_vista_smtp(self) -> ctk.CTkFrame:
        frame = ctk.CTkFrame(self.container, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_columnconfigure(1, weight=1)
        frame.grid_rowconfigure(0, weight=1)

        # Panel Izquierdo: Configuración Servidor SMTP
        config_frame = ctk.CTkFrame(frame)
        config_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10), pady=0)

        ctk.CTkLabel(
            config_frame, 
            text="⚙️ Configuración SMTP", 
            font=ctk.CTkFont(size=18, weight="bold"), 
            text_color=("black", "white")
        ).pack(pady=(15, 8))

        # Servidor Host (Solo Lectura - Visible pero no editable)
        ctk.CTkLabel(
            config_frame, 
            text="Servidor Host: (Modo Solo Lectura)", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(5, 2))
        self.entry_smtp_host = ctk.CTkEntry(config_frame, placeholder_text="smtp.gmail.com")
        self.entry_smtp_host.insert(0, "smtp.gmail.com")
        self.entry_smtp_host.configure(state="readonly")
        self.entry_smtp_host.pack(fill="x", padx=20, pady=3)

        # Puerto (Solo Lectura - Visible pero no editable)
        ctk.CTkLabel(
            config_frame, 
            text="Puerto: (Modo Solo Lectura)", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(5, 2))
        self.entry_smtp_port = ctk.CTkEntry(config_frame, placeholder_text="587")
        self.entry_smtp_port.insert(0, "587")
        self.entry_smtp_port.configure(state="readonly")
        self.entry_smtp_port.pack(fill="x", padx=20, pady=3)

        # Correo Remitente (Guardado en BD)
        ctk.CTkLabel(
            config_frame, 
            text="Correo Remitente:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(5, 2))
        self.entry_smtp_user = ctk.CTkEntry(config_frame, placeholder_text="tu_correo@gmail.com")
        self.entry_smtp_user.pack(fill="x", padx=20, pady=3)

        # Contraseña / App Password (Guardada en BD) con botón mostrar/ocultar
        ctk.CTkLabel(
            config_frame, 
            text="Contraseña / App Password:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(5, 2))
        
        frame_pass = ctk.CTkFrame(config_frame, fg_color="transparent")
        frame_pass.pack(fill="x", padx=20, pady=3)
        frame_pass.grid_columnconfigure(0, weight=1)

        self.entry_smtp_pass = ctk.CTkEntry(frame_pass, placeholder_text="Contraseña de aplicación", show="*")
        self.entry_smtp_pass.grid(row=0, column=0, sticky="ew", padx=(0, 6))

        self.btn_toggle_smtp_pass = ctk.CTkButton(
            frame_pass,
            text="👁️",
            width=38,
            height=28,
            fg_color=["gray75", "gray30"],
            hover_color=["gray65", "gray40"],
            text_color=("black", "white"),
            command=self.toggle_ver_password_smtp
        )
        self.btn_toggle_smtp_pass.grid(row=0, column=1)

        # Switch STARTTLS
        self.switch_smtp_tls = ctk.CTkSwitch(config_frame, text="Usar STARTTLS (Recomendado)", text_color=("black", "white"))
        self.switch_smtp_tls.select()
        self.switch_smtp_tls.pack(anchor="w", padx=20, pady=(8, 6))

        # Fila compartida: Guardar Credenciales en BD y Verificar Conexión SMTP
        frame_smtp_acciones = ctk.CTkFrame(config_frame, fg_color="transparent")
        frame_smtp_acciones.pack(fill="x", padx=20, pady=(2, 10))
        frame_smtp_acciones.grid_columnconfigure(0, weight=1)
        frame_smtp_acciones.grid_columnconfigure(1, weight=1)

        self.btn_guardar_smtp = ctk.CTkButton(
            frame_smtp_acciones,
            text="💾 Guardar Credenciales",
            font=ctk.CTkFont(size=11, weight="bold"),
            height=34,
            fg_color="#27ae60",
            hover_color="#219955",
            command=self.guardar_credenciales_smtp_manual
        )
        self.btn_guardar_smtp.grid(row=0, column=0, sticky="ew", padx=(0, 4))

        self.btn_verificar_smtp = ctk.CTkButton(
            frame_smtp_acciones,
            text="🔌 Verificar Conexión",
            font=ctk.CTkFont(size=11, weight="bold"),
            height=34,
            fg_color="#2980b9",
            hover_color="#1f618d",
            command=self.verificar_conexion_smtp_manual
        )
        self.btn_verificar_smtp.grid(row=0, column=1, sticky="ew", padx=(4, 0))

        # Cargar credenciales guardadas automáticamente de la base de datos
        creds = self.db.obtener_credenciales_smtp()
        if creds.get("email"):
            self.entry_smtp_user.delete(0, "end")
            self.entry_smtp_user.insert(0, creds["email"])
            self._smtp_email_conectado = creds["email"]
            if hasattr(self, "lbl_sidebar_smtp_cuenta"):
                self.lbl_sidebar_smtp_cuenta.configure(text=f"✉️ {creds['email']}")
        if creds.get("password"):
            self.entry_smtp_pass.delete(0, "end")
            self.entry_smtp_pass.insert(0, creds["password"])
        if creds.get("tls") is not None:
            if creds["tls"]:
                self.switch_smtp_tls.select()
            else:
                self.switch_smtp_tls.deselect()

        # Panel de Estado / Registro de Envío
        ctk.CTkLabel(
            config_frame, 
            text="Historial de Actividad Reciente:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(8, 4))
        
        self.log_smtp_box = ctk.CTkTextbox(config_frame, height=135)
        self.log_smtp_box.pack(fill="both", expand=True, padx=20, pady=(0, 15))

        # Panel Derecho: Redacción y Selección de Canal
        compose_frame = ctk.CTkFrame(frame)
        compose_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0), pady=0)

        ctk.CTkLabel(
            compose_frame, 
            text="✉️ Redactar Mensaje", 
            font=ctk.CTkFont(size=18, weight="bold"), 
            text_color=("black", "white")
        ).pack(pady=(15, 8))

        # Selector de Canal: Correo SMTP o WhatsApp
        ctk.CTkLabel(compose_frame, text="Canal de Envío:", font=ctk.CTkFont(size=12, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=20, pady=(2, 2))
        self.combo_canal_manual = ctk.CTkComboBox(
            compose_frame, 
            values=["Correo SMTP", "WhatsApp"], 
            command=self.on_cambio_canal_manual,
            state="readonly"
        )
        self.combo_canal_manual.set("Correo SMTP")
        self.combo_canal_manual.pack(fill="x", padx=20, pady=2)

        # Destinatario registrado o manual con Buscador Rápido
        frame_lbl_dest = ctk.CTkFrame(compose_frame, fg_color="transparent")
        frame_lbl_dest.pack(fill="x", padx=20, pady=(4, 2))

        ctk.CTkLabel(
            frame_lbl_dest, 
            text="Seleccionar Cliente Registrado:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(side="left")

        self.btn_buscar_cliente_smtp = ctk.CTkButton(
            frame_lbl_dest,
            text="🔍 Buscar Cliente",
            width=120,
            height=26,
            fg_color="#2980b9",
            hover_color="#1f618d",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.abrir_modal_buscar_cliente_smtp
        )
        self.btn_buscar_cliente_smtp.pack(side="right")

        self.combo_destinatarios = ctk.CTkComboBox(
            compose_frame, 
            values=["(Cargar cliente...)"], 
            command=self.on_cliente_seleccionado,
            state="readonly"
        )
        self.combo_destinatarios.pack(fill="x", padx=20, pady=2)

        self.lbl_destinatario = ctk.CTkLabel(compose_frame, text="O Correo Destinatario:", font=ctk.CTkFont(size=12), text_color=("black", "white"))
        self.lbl_destinatario.pack(anchor="w", padx=20, pady=(4, 2))
        self.entry_destinatario = ctk.CTkEntry(compose_frame, placeholder_text="destino@correo.com")
        self.entry_destinatario.pack(fill="x", padx=20, pady=2)

        # SECCIÓN PLANTILLAS DE COPYWRITING PARA MINIMARKET
        ctk.CTkLabel(
            compose_frame, 
            text="Plantillas de Copywriting para Minimarket:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(6, 2))

        nombres_plantillas_email = list(PLANTILLAS_EMAIL_MINIMARKET.keys())
        self.combo_smtp_plantillas = ctk.CTkComboBox(
            compose_frame, 
            values=nombres_plantillas_email,
            state="readonly"
        )
        self.combo_smtp_plantillas.set(nombres_plantillas_email[0])
        self.combo_smtp_plantillas.pack(fill="x", padx=20, pady=2)

        frame_plantilla_btns = ctk.CTkFrame(compose_frame, fg_color="transparent")
        frame_plantilla_btns.pack(fill="x", padx=20, pady=(3, 4))
        frame_plantilla_btns.grid_columnconfigure(0, weight=1)

        btn_cargar_smtp = ctk.CTkButton(
            frame_plantilla_btns,
            text="📋 Cargar Plantilla en Asunto y Cuerpo",
            font=ctk.CTkFont(size=12, weight="bold"),
            height=28,
            fg_color=["#3a7ebf", "#1f538d"],
            command=self.cargar_plantilla_smtp_seleccionada
        )
        btn_cargar_smtp.grid(row=0, column=0, sticky="ew")

        # Inserción de variables dinámicas
        vars_email_frame = ctk.CTkFrame(compose_frame, fg_color="transparent")
        vars_email_frame.pack(fill="x", padx=20, pady=(0, 4))
        vars_email_frame.grid_columnconfigure((0, 1, 2, 3), weight=1)

        ctk.CTkButton(
            vars_email_frame, text="+ {nombre}", height=22, font=ctk.CTkFont(size=10),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_smtp("{nombre}")
        ).grid(row=0, column=0, padx=2, sticky="ew")

        ctk.CTkButton(
            vars_email_frame, text="+ {correo}", height=22, font=ctk.CTkFont(size=10),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_smtp("{correo}")
        ).grid(row=0, column=1, padx=2, sticky="ew")

        ctk.CTkButton(
            vars_email_frame, text="+ {fecha}", height=22, font=ctk.CTkFont(size=10),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_smtp("{fecha}")
        ).grid(row=0, column=2, padx=2, sticky="ew")

        ctk.CTkButton(
            vars_email_frame, text="+ {minimarket}", height=22, font=ctk.CTkFont(size=10),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_smtp("{minimarket}")
        ).grid(row=0, column=3, padx=2, sticky="ew")

        # Asunto y Cuerpo
        self.lbl_asunto = ctk.CTkLabel(compose_frame, text="Asunto (Opcional en WhatsApp):", font=ctk.CTkFont(size=12), text_color=("black", "white"))
        self.lbl_asunto.pack(anchor="w", padx=20, pady=(2, 2))
        self.entry_asunto = ctk.CTkEntry(compose_frame, placeholder_text="Ej: ¡Aviso de oferta especial!")
        self.entry_asunto.pack(fill="x", padx=20, pady=2)

        ctk.CTkLabel(compose_frame, text="Cuerpo del Mensaje:", font=ctk.CTkFont(size=12), text_color=("black", "white")).pack(anchor="w", padx=20, pady=(4, 2))
        self.text_cuerpo = ctk.CTkTextbox(compose_frame, height=130)
        self.text_cuerpo.pack(fill="both", expand=True, padx=20, pady=2)

        # Cargar plantilla por defecto
        self.cargar_plantilla_smtp_seleccionada()

        self.btn_enviar_manual = ctk.CTkButton(
            compose_frame, 
            text="🚀 Enviar Mensaje Vía SMTP", 
            font=ctk.CTkFont(size=14, weight="bold"),
            height=40,
            fg_color="#2980b9",
            hover_color="#1f618d",
            command=self.iniciar_envio_manual
        )
        self.btn_enviar_manual.pack(fill="x", padx=20, pady=(10, 15))

        return frame

    def toggle_ver_password_smtp(self):
        """Alterna la visualización de la contraseña entre oculta (*) y visible."""
        actual = self.entry_smtp_pass.cget("show")
        if actual == "*":
            self.entry_smtp_pass.configure(show="")
            self.btn_toggle_smtp_pass.configure(text="🔒")
        else:
            self.entry_smtp_pass.configure(show="*")
            self.btn_toggle_smtp_pass.configure(text="👁️")

    def guardar_credenciales_smtp_manual(self, silencioso: bool = False) -> bool:
        """Guarda permanentemente las credenciales SMTP en la base de datos."""
        user = self.entry_smtp_user.get().strip()
        pwd = self.entry_smtp_pass.get().strip()
        tls = bool(self.switch_smtp_tls.get())
        host = self.entry_smtp_host.get().strip() or "smtp.gmail.com"
        puerto_str = self.entry_smtp_port.get().strip() or "587"
        try:
            puerto = int(puerto_str)
        except ValueError:
            puerto = 587

        if not user or not pwd:
            if not silencioso:
                messagebox.showwarning("Campos Incompletos", "Por favor ingrese el Correo Remitente y la Contraseña antes de guardar.")
            return False

        exito = self.db.guardar_credenciales_smtp(email=user, password=pwd, tls=tls, host=host, puerto=puerto)
        if exito:
            self._smtp_email_conectado = user
            self._actualizar_estado_visual_smtp(self._smtp_esta_conectado, email=user)
            if not silencioso:
                self.log_smtp("✅ Credenciales guardadas en la base de datos.")
                messagebox.showinfo(
                    "Credenciales Guardadas", 
                    f"¡Excelente! El correo y contraseña han sido guardados permanentemente.\n\n"
                    f"Remitente: {user}\n"
                    "Al volver a ingresar a la aplicación, se cargarán de forma automática."
                )
            return True
        else:
            if not silencioso:
                messagebox.showerror("Error", "No se pudieron guardar las credenciales en la base de datos.")
            return False

    def verificar_conexion_smtp_manual(self):
        """Prueba la conexión y autenticación con el servidor SMTP con las credenciales ingresadas."""
        user = self.entry_smtp_user.get().strip()
        pwd = self.entry_smtp_pass.get().strip()
        tls = bool(self.switch_smtp_tls.get())
        host = self.entry_smtp_host.get().strip() or "smtp.gmail.com"
        puerto_str = self.entry_smtp_port.get().strip() or "587"

        if not user or not pwd:
            messagebox.showwarning(
                "Credenciales Requeridas",
                "Por favor ingrese el Correo Remitente y la Contraseña antes de verificar la conexión."
            )
            return

        try:
            puerto = int(puerto_str)
        except ValueError:
            messagebox.showerror("Puerto Inválido", "El puerto debe ser un número entero (ej. 587 o 465).")
            return

        self.btn_verificar_smtp.configure(state="disabled", text="⏳ Verificando...")
        self.log_smtp(f"Iniciando verificación de conexión SMTP con {user} ({host}:{puerto})...")

        def tarea():
            smtp = SMTPManager(email=user, password=pwd, host=host, puerto=puerto, use_tls=tls)
            exito, mensaje = smtp.verificar_conexion()

            def fin():
                self.btn_verificar_smtp.configure(state="normal", text="🔌 Verificar Conexión")
                if exito:
                    self._actualizar_estado_visual_smtp(True, email=user)
                    self.log_smtp(f"✅ Conexión SMTP verificada con éxito: {user}")
                    messagebox.showinfo(
                        "Conexión SMTP Exitosa",
                        f"¡Conexión verificada exitosamente!\n\n"
                        f"• Servidor: {host}:{puerto}\n"
                        f"• Correo Remitente: {user}\n"
                        f"• Seguridad: {'STARTTLS' if tls else 'SSL/Estándar'}\n"
                        f"• Estado: Conectado y autenticado correctamente.\n\n"
                        "Listo para enviar correos electrónicos a sus clientes."
                    )
                else:
                    self._actualizar_estado_visual_smtp(False, email=user)
                    self.log_smtp(f"❌ Error de verificación SMTP: {mensaje}")
                    messagebox.showerror(
                        "Fallo de Conexión SMTP",
                        f"No se pudo conectar o autenticar con el servidor SMTP:\n\n"
                        f"{mensaje}\n\n"
                        "Sugerencias:\n"
                        "1. Si usa Gmail, verifique haber generado una 'Contraseña de Aplicación' de 16 letras.\n"
                        "2. Compruebe que no haya espacios en blanco antes o después del correo o contraseña.\n"
                        "3. Revise su conexión a Internet."
                    )

            self.after(0, fin)

        threading.Thread(target=tarea, daemon=True).start()

    def _auto_verificar_conexion_inicial_smtp(self):
        """Verifica en segundo plano el estado de la conexión SMTP al iniciar la aplicación."""
        creds = self.db.obtener_credenciales_smtp()
        email = (creds.get("email") or "").strip()
        pwd = (creds.get("password") or "").strip()
        host = (creds.get("host") or "smtp.gmail.com").strip()
        puerto = creds.get("puerto") or 587
        tls = creds.get("tls", True)

        if not email or not pwd:
            self._actualizar_estado_visual_smtp(False, email="")
            return

        if hasattr(self, "lbl_sidebar_smtp_status"):
            self.lbl_sidebar_smtp_status.configure(
                text="🟡 SMTP: Comprobando...",
                text_color="#f39c12"
            )
        if hasattr(self, "lbl_sidebar_smtp_cuenta"):
            self.lbl_sidebar_smtp_cuenta.configure(
                text=f"✉️ {email}",
                text_color=("gray40", "gray70")
            )

        def tarea():
            try:
                smtp = SMTPManager(email=email, password=pwd, host=host, puerto=puerto, use_tls=tls)
                exito, _ = smtp.verificar_conexion()
            except Exception:
                exito = False
            self.after(0, lambda: self._actualizar_estado_visual_smtp(exito, email=email))

        threading.Thread(target=tarea, daemon=True).start()

    def _actualizar_estado_visual_smtp(self, esta_conectado: bool, email: str = ""):
        """Actualiza el indicador permanente de conexión SMTP en el menú lateral."""
        self._smtp_esta_conectado = bool(esta_conectado)
        if email:
            self._smtp_email_conectado = email

        email_a_mostrar = email or getattr(self, "_smtp_email_conectado", "")

        if hasattr(self, "lbl_sidebar_smtp_status"):
            if esta_conectado:
                self.lbl_sidebar_smtp_status.configure(
                    text="🟢 SMTP: Conectado",
                    text_color="#2ecc71"
                )
            else:
                self.lbl_sidebar_smtp_status.configure(
                    text="🔴 SMTP: Desconectado",
                    text_color="#e74c3c"
                )

        if hasattr(self, "lbl_sidebar_smtp_cuenta"):
            if esta_conectado and email_a_mostrar:
                self.lbl_sidebar_smtp_cuenta.configure(
                    text=f"✉️ {email_a_mostrar}",
                    text_color=("#155724", "#4ade80")
                )
            elif not esta_conectado and email_a_mostrar:
                self.lbl_sidebar_smtp_cuenta.configure(
                    text=f"✉️ {email_a_mostrar}",
                    text_color=("gray40", "gray70")
                )
            else:
                self.lbl_sidebar_smtp_cuenta.configure(
                    text="",
                    text_color=("gray40", "gray70")
                )

    def cargar_plantilla_smtp_seleccionada(self):
        """Carga la plantilla seleccionada en el asunto y cuerpo con copywriting para Minimarket."""
        plantilla_nombre = getattr(self, "combo_smtp_plantillas", None)
        if not plantilla_nombre:
            return
        nombre_sel = plantilla_nombre.get()

        canal = self.combo_canal_manual.get() if hasattr(self, "combo_canal_manual") else "Correo SMTP"

        nom = "Cliente"
        correo = ""
        if hasattr(self, "combo_destinatarios"):
            dest_val = self.combo_destinatarios.get()
            if "(" in dest_val and ")" in dest_val:
                correo = dest_val.split("(")[-1].replace(")", "").strip()
                nom_part = dest_val.split("-")[-1].split("(")[0].strip()
                if nom_part and "No hay" not in nom_part:
                    nom = nom_part
        if not correo and hasattr(self, "entry_destinatario"):
            correo = self.entry_destinatario.get().strip()

        fecha_actual = datetime.now().strftime("%d/%m/%Y")

        if canal == "WhatsApp" and nombre_sel in PLANTILLAS_MINIMARKET:
            cuerpo_raw = PLANTILLAS_MINIMARKET[nombre_sel]
            cuerpo = (
                cuerpo_raw
                .replace("{nombre}", nom)
                .replace("{correo}", correo)
                .replace("{fecha}", fecha_actual)
                .replace("{minimarket}", "Minimarket")
            )
            self.text_cuerpo.delete("1.0", "end")
            self.text_cuerpo.insert("1.0", cuerpo)
        elif nombre_sel in PLANTILLAS_EMAIL_MINIMARKET:
            item = PLANTILLAS_EMAIL_MINIMARKET[nombre_sel]
            asunto_raw = item.get("asunto", "")
            cuerpo_raw = item.get("cuerpo", "")

            asunto = (
                asunto_raw
                .replace("{nombre}", nom)
                .replace("{correo}", correo)
                .replace("{fecha}", fecha_actual)
                .replace("{minimarket}", "Minimarket")
            )
            cuerpo = (
                cuerpo_raw
                .replace("{nombre}", nom)
                .replace("{correo}", correo)
                .replace("{fecha}", fecha_actual)
                .replace("{minimarket}", "Minimarket")
            )

            self.entry_asunto.delete(0, "end")
            self.entry_asunto.insert(0, asunto)

            self.text_cuerpo.delete("1.0", "end")
            self.text_cuerpo.insert("1.0", cuerpo)

    def insertar_variable_smtp(self, var: str):
        """Inserta la variable seleccionada en el cuerpo o en el asunto si este tiene foco."""
        try:
            widget_foco = self.focus_get()
            if widget_foco == self.entry_asunto._entry:
                self.entry_asunto.insert("insert", var)
                return
        except Exception:
            pass
        self.text_cuerpo.insert("insert", var)

    def on_cambio_canal_manual(self, canal: str):
        self.actualizar_combobox_clientes()
        if canal == "WhatsApp":
            self.lbl_destinatario.configure(text="O Teléfono Destinatario (ej. 51987654321):")
            self.entry_destinatario.configure(placeholder_text="51987654321")
            self.btn_enviar_manual.configure(text="🚀 Enviar Mensaje Vía WhatsApp", fg_color="#25D366", hover_color="#1ea952")
            if hasattr(self, "combo_smtp_plantillas"):
                opciones_wsp = list(PLANTILLAS_MINIMARKET.keys())
                self.combo_smtp_plantillas.configure(values=opciones_wsp)
                self.combo_smtp_plantillas.set(opciones_wsp[0])
                self.cargar_plantilla_smtp_seleccionada()
        else:
            self.lbl_destinatario.configure(text="O Correo Destinatario:")
            self.entry_destinatario.configure(placeholder_text="destino@correo.com")
            self.btn_enviar_manual.configure(text="🚀 Enviar Mensaje Vía SMTP", fg_color="#2980b9", hover_color="#1f618d")
            if hasattr(self, "combo_smtp_plantillas"):
                opciones_email = list(PLANTILLAS_EMAIL_MINIMARKET.keys())
                self.combo_smtp_plantillas.configure(values=opciones_email)
                self.combo_smtp_plantillas.set(opciones_email[0])
                self.cargar_plantilla_smtp_seleccionada()

    def actualizar_combobox_clientes(self):
        canal = getattr(self, "combo_canal_manual", None)
        es_wsp = canal and canal.get() == "WhatsApp"
        # REGLA CRÍTICA: Filtrar únicamente clientes Activos
        clientes = [c for c in self.db.get_all_clientes() if (c.get("estado") or "").strip().lower() == "activo"]
        
        item_buscar = "🔍 Buscar Cliente..."
        if es_wsp:
            opciones_cli = [f"{c['id']} - {c['nombre']} ({c['telefono'] or 'Sin teléfono'})" for c in clientes if c.get("telefono")]
            opciones = [item_buscar] + opciones_cli if opciones_cli else ["(No hay clientes activos con teléfono)"]
        else:
            opciones_cli = [f"{c['id']} - {c['nombre']} ({c['correo']})" for c in clientes if c.get("correo")]
            opciones = [item_buscar] + opciones_cli if opciones_cli else ["(No hay clientes activos con correo)"]
                
        self.combo_destinatarios.configure(values=opciones)
        if len(opciones) > 1:
            self.combo_destinatarios.set(opciones[1])
            self.on_cliente_seleccionado(opciones[1])
        else:
            self.combo_destinatarios.set(opciones[0])

    def on_cliente_seleccionado(self, valor: str):
        if "🔍 Buscar Cliente" in valor:
            self.abrir_modal_buscar_cliente_smtp()
            return
        if "(" in valor and ")" in valor:
            dato = valor.split("(")[-1].replace(")", "").strip()
            if "Sin " not in dato and "No hay" not in dato:
                self.entry_destinatario.delete(0, "end")
                self.entry_destinatario.insert(0, dato)
                self.cargar_plantilla_smtp_seleccionada()

    def abrir_modal_buscar_cliente_smtp(self):
        """Abre ventana modal para buscar clientes activos y cargarlos como destinatario en el módulo SMTP."""
        canal = getattr(self, "combo_canal_manual", None)
        es_wsp = canal and canal.get() == "WhatsApp"

        modal = ctk.CTkToplevel(self)
        titulo_canal = "WhatsApp" if es_wsp else "Correo Electrónico (SMTP)"
        modal.title(f"🔍 Buscar Destinatario - {titulo_canal}")
        modal.geometry("750x540")
        modal.minsize(630, 420)
        modal.transient(self)
        modal.grab_set()

        try:
            x = self.winfo_x() + (self.winfo_width() - 750) // 2
            y = self.winfo_y() + (self.winfo_height() - 540) // 2
            modal.geometry(f"+{max(0, x)}+{max(0, y)}")
        except Exception:
            pass

        # Cabecera
        top_frame = ctk.CTkFrame(modal, fg_color="transparent")
        top_frame.pack(fill="x", padx=20, pady=(15, 8))

        ctk.CTkLabel(
            top_frame,
            text=f"🔍 Seleccionar Destinatario para {titulo_canal}",
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w")

        ctk.CTkLabel(
            top_frame,
            text="Filtre por nombre, correo electrónico o teléfono para seleccionar al cliente rápidamente.",
            font=ctk.CTkFont(size=12),
            text_color=("gray40", "gray70")
        ).pack(anchor="w", pady=(2, 6))

        # Buscador en vivo
        entry_buscar = ctk.CTkEntry(
            modal,
            placeholder_text="🔎 Escriba Nombre, Correo o Teléfono...",
            height=38,
            font=ctk.CTkFont(size=13)
        )
        entry_buscar.pack(fill="x", padx=20, pady=(0, 6))
        entry_buscar.focus_set()

        lbl_conteo = ctk.CTkLabel(
            modal,
            text="Cargando clientes activos...",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("gray40", "gray60")
        )
        lbl_conteo.pack(anchor="w", padx=22, pady=(0, 6))

        scroll_modal = ctk.CTkScrollableFrame(modal)
        scroll_modal.pack(fill="both", expand=True, padx=20, pady=(0, 10))

        # Solo clientes activos
        todos_clientes = [
            c for c in self.db.get_all_clientes()
            if (c.get("estado") or "").strip().lower() == "activo"
        ]

        def seleccionar_y_cerrar(cliente_datos):
            cid = cliente_datos["id"]
            nom = cliente_datos["nombre"]
            correo = (cliente_datos.get("correo") or "").strip()
            tel = (cliente_datos.get("telefono") or "").strip()

            if es_wsp:
                dato_dest = tel
                label_dest = f"{cid} - {nom} ({tel})"
            else:
                dato_dest = correo
                label_dest = f"{cid} - {nom} ({correo})"

            self.entry_destinatario.delete(0, "end")
            self.entry_destinatario.insert(0, dato_dest)

            valores_actuales = list(self.combo_destinatarios.cget("values"))
            coincidencia = None
            for v in valores_actuales:
                if str(cid) in v and (dato_dest in v or nom in v):
                    coincidencia = v
                    break
            if coincidencia:
                self.combo_destinatarios.set(coincidencia)
            elif dato_dest:
                nuevos_valores = valores_actuales + [label_dest]
                self.combo_destinatarios.configure(values=nuevos_valores)
                self.combo_destinatarios.set(label_dest)

            self.cargar_plantilla_smtp_seleccionada()
            modal.destroy()

        def renderizar_resultados(query: str = ""):
            for w in scroll_modal.winfo_children():
                w.destroy()

            q = query.strip().lower()
            if not q:
                filtrados = todos_clientes
            else:
                filtrados = []
                for c in todos_clientes:
                    nom = str(c.get("nombre") or "").lower()
                    cor = str(c.get("correo") or "").lower()
                    tel = str(c.get("telefono") or "").lower()
                    cid = str(c.get("id") or "")
                    if q in nom or q in cor or q in tel or q == cid:
                        filtrados.append(c)

            lbl_conteo.configure(text=f"Clientes encontrados: {len(filtrados)} de {len(todos_clientes)} clientes activos")

            if not filtrados:
                ctk.CTkLabel(
                    scroll_modal,
                    text="❌ No se encontraron clientes activos con esa búsqueda.",
                    font=ctk.CTkFont(size=13),
                    text_color=("gray40", "gray60")
                ).pack(pady=30)
                return

            for c in filtrados:
                item_frame = ctk.CTkFrame(scroll_modal)
                item_frame.pack(fill="x", padx=4, pady=4)
                item_frame.grid_columnconfigure(0, weight=1)

                cat = c.get("categoria") or "🥉 Bronce"
                correo_txt = c['correo'] or 'Sin correo'
                tel_txt = c['telefono'] or 'Sin tel'
                info = f"ID: {c['id']} | {c['nombre']}  •  📧 {correo_txt}  •  📱 {tel_txt}"

                lbl_info = ctk.CTkLabel(
                    item_frame,
                    text=info,
                    font=ctk.CTkFont(size=12, weight="bold"),
                    anchor="w",
                    text_color=("black", "white")
                )
                lbl_info.grid(row=0, column=0, padx=10, pady=8, sticky="w")

                lbl_cat = ctk.CTkLabel(
                    item_frame,
                    text=f" {cat} ",
                    font=ctk.CTkFont(size=11, weight="bold"),
                    text_color="#3498db"
                )
                lbl_cat.grid(row=0, column=1, padx=4, pady=8)

                btn_sel = ctk.CTkButton(
                    item_frame,
                    text="Seleccionar ✉️" if not es_wsp else "Seleccionar 💬",
                    width=100,
                    fg_color="#2980b9",
                    hover_color="#1f618d",
                    command=lambda cli=c: seleccionar_y_cerrar(cli)
                )
                btn_sel.grid(row=0, column=2, padx=(6, 10), pady=6)

        entry_buscar.bind("<KeyRelease>", lambda event: renderizar_resultados(entry_buscar.get()))
        renderizar_resultados("")

        bot_frame = ctk.CTkFrame(modal, fg_color="transparent")
        bot_frame.pack(fill="x", padx=20, pady=(0, 12))
        ctk.CTkButton(
            bot_frame,
            text="Cerrar",
            width=100,
            fg_color="gray40",
            hover_color="gray30",
            command=modal.destroy
        ).pack(side="right")

    def log_smtp(self, mensaje: str):
        self.log_smtp_box.insert("end", f"• {mensaje}\n")
        self.log_smtp_box.see("end")

    def obtener_smtp_manager(self) -> Optional[SMTPManager]:
        host = self.entry_smtp_host.get().strip() if hasattr(self, "entry_smtp_host") else "smtp.gmail.com"
        puerto_str = self.entry_smtp_port.get().strip() if hasattr(self, "entry_smtp_port") else "587"
        user = self.entry_smtp_user.get().strip() if hasattr(self, "entry_smtp_user") else ""
        pwd = self.entry_smtp_pass.get().strip() if hasattr(self, "entry_smtp_pass") else ""
        use_tls = bool(self.switch_smtp_tls.get()) if hasattr(self, "switch_smtp_tls") else True

        if not user or not pwd:
            creds = self.db.obtener_credenciales_smtp()
            if creds.get("email") and creds.get("password"):
                user = creds["email"]
                pwd = creds["password"]
                host = creds.get("host", "smtp.gmail.com") or "smtp.gmail.com"
                puerto_str = str(creds.get("puerto", 587) or 587)
                use_tls = bool(creds.get("tls", True))
            else:
                return None

        try:
            puerto = int(puerto_str)
        except ValueError:
            return None

        return SMTPManager(email=user, password=pwd, host=host, puerto=puerto, use_tls=use_tls)

    def iniciar_envio_manual(self):
        canal = self.combo_canal_manual.get() if hasattr(self, "combo_canal_manual") else "Correo SMTP"
        destinatario = self.entry_destinatario.get().strip()
        asunto = self.entry_asunto.get().strip()
        cuerpo = self.text_cuerpo.get("1.0", "end").strip()

        if not destinatario or not cuerpo:
            messagebox.showwarning("Campos Incompletos", "Por favor ingrese el destinatario y el cuerpo del mensaje.")
            return

        # REGLA CRÍTICA: Impedir envío si el destinatario corresponde a un cliente inactivo
        for c in self.db.get_all_clientes():
            if (c.get("estado") or "").strip().lower() == "inactivo":
                target_clean = re.sub(r'\D', '', destinatario)
                tel_c = re.sub(r'\D', '', str(c.get("telefono") or ""))
                correo_c = (c.get("correo") or "").strip().lower()
                if (correo_c and correo_c == destinatario.lower()) or (tel_c and target_clean and (tel_c in target_clean or target_clean.endswith(tel_c))):
                    messagebox.showwarning(
                        "Cliente Inactivo",
                        f"⛔ El cliente '{c.get('nombre')}' está registrado con estado 'Inactivo'.\n\n"
                        "El sistema restringe el envío de mensajes a clientes inactivos.\n"
                        "Para enviarle mensajes, actívelo primero desde el módulo Gestión Clientes."
                    )
                    return

        # Auto-guardar credenciales en la base de datos de manera silenciosa
        self.guardar_credenciales_smtp_manual(silencioso=True)

        # Caso Envío WhatsApp
        if canal == "WhatsApp":
            numero_whatsapp, numero_display = self.formatear_numero_con_pais(destinatario)
            if not numero_whatsapp or len(numero_whatsapp) < 8:
                messagebox.showwarning("Número Inválido", "Por favor ingrese un número telefónico válido.")
                return

            self.btn_enviar_manual.configure(state="disabled", text="Enviando WhatsApp...")
            self.log_smtp(f"Iniciando despacho WhatsApp a {numero_display}...")

            def tarea_wsp():
                try:
                    gestor = self.obtener_gestor_whatsapp()
                    exito, estado = gestor.enviar_mensaje_whatsapp(numero=numero_whatsapp, mensaje=cuerpo)
                except Exception as err:
                    exito = False
                    estado = f"Fallo: {err}"

                # Buscar id cliente si existe
                cliente_id = None
                for c in self.db.get_all_clientes():
                    if (c.get("telefono") or "").strip() in destinatario:
                        cliente_id = c["id"]
                        break
                if cliente_id:
                    self.db.create_historial(cliente_id=cliente_id, estado=f"Manual WhatsApp: {estado}")

                def fin_wsp():
                    self.btn_enviar_manual.configure(state="normal", text="🚀 Enviar Mensaje Vía WhatsApp")
                    if exito:
                        self.log_smtp(f"✅ WhatsApp enviado con éxito a {numero_display}")
                        messagebox.showinfo("Envío Exitoso", f"Mensaje de WhatsApp enviado a {numero_display}.")
                    else:
                        self.log_smtp(f"❌ WhatsApp Error: {estado}")
                        messagebox.showerror("Fallo en Envío", f"No se pudo enviar por WhatsApp:\n{estado}")
                self.after(0, fin_wsp)

            threading.Thread(target=tarea_wsp, daemon=True).start()
            return

        # Caso Envío Correo SMTP
        if not asunto:
            messagebox.showwarning("Asunto Requerido", "Por favor ingrese el asunto del correo electrónico.")
            return

        smtp = self.obtener_smtp_manager()
        if not smtp:
            return

        cliente_id = None
        for c in self.db.get_all_clientes():
            if c.get("correo") == destinatario:
                cliente_id = c["id"]
                break

        self.btn_enviar_manual.configure(state="disabled", text="Enviando correo...")
        self.log_smtp(f"Iniciando envío SMTP a {destinatario}...")

        def tarea_smtp():
            exito, estado = smtp.enviar_mensaje(destinatario=destinatario, asunto=asunto, cuerpo=cuerpo)
            if cliente_id:
                self.db.create_historial(cliente_id=cliente_id, estado=f"Manual SMTP: {estado}")

            def actualizar_gui():
                self.btn_enviar_manual.configure(state="normal", text="🚀 Enviar Mensaje Vía SMTP")
                if exito:
                    self._actualizar_estado_visual_smtp(True, email=smtp.email)
                    self.log_smtp(f"✅ Correo enviado con éxito a {destinatario}")
                    messagebox.showinfo("Envío Exitoso", f"El correo fue enviado exitosamente a {destinatario}.")
                else:
                    if "autenticación" in estado.lower() or "conectar" in estado.lower():
                        self._actualizar_estado_visual_smtp(False, email=smtp.email)
                    self.log_smtp(f"❌ Correo Error: {estado}")
                    messagebox.showerror("Fallo en Envío", f"No se pudo enviar el mensaje:\n{estado}")

            self.after(0, actualizar_gui)

        threading.Thread(target=tarea_smtp, daemon=True).start()

    # ------------------------------------------------------------------
    # VISTA ESPECÍFICA: ENVÍOS MANUALES WHATSAPP (Con Plantillas Minimarket)
    # ------------------------------------------------------------------
    def construir_vista_envios_whatsapp(self) -> ctk.CTkFrame:
        frame = ctk.CTkFrame(self.container, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_columnconfigure(1, weight=1)
        frame.grid_rowconfigure(0, weight=1)

        # Panel Izquierdo: Selección de Destinatario y Plantillas de Minimarket
        left_frame = ctk.CTkFrame(frame)
        left_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10), pady=0)

        ctk.CTkLabel(
            left_frame, 
            text="📲 Envíos Manuales WhatsApp", 
            font=ctk.CTkFont(size=18, weight="bold"), 
            text_color=("black", "white")
        ).pack(pady=(15, 4))

        # Badge indicador de estado de conexión
        self.lbl_envios_wsp_badge = ctk.CTkLabel(
            left_frame,
            text="🔴 WhatsApp Desconectado (Verifique en 'Conexión WhatsApp')",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color="#e74c3c"
        )
        self.lbl_envios_wsp_badge.pack(pady=(0, 10))

        # Sección 1: Destinatario
        frame_lbl_wsp = ctk.CTkFrame(left_frame, fg_color="transparent")
        frame_lbl_wsp.pack(fill="x", padx=20, pady=(6, 2))

        ctk.CTkLabel(
            frame_lbl_wsp, 
            text="1. Seleccionar Cliente Registrado:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(side="left")

        self.btn_buscar_cliente_wsp = ctk.CTkButton(
            frame_lbl_wsp,
            text="🔍 Buscar Cliente",
            width=120,
            height=26,
            fg_color="#25D366",
            hover_color="#1ea952",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.abrir_modal_buscar_cliente_wsp
        )
        self.btn_buscar_cliente_wsp.pack(side="right")

        self.combo_wsp_clientes = ctk.CTkComboBox(
            left_frame, 
            values=["(Cargando clientes...)"], 
            command=self.on_cliente_wsp_seleccionado,
            state="readonly"
        )
        self.combo_wsp_clientes.pack(fill="x", padx=20, pady=3)

        ctk.CTkLabel(
            left_frame, 
            text="O Número Telefónico Directo:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(6, 2))

        # Contenedor de País Predeterminado (Bolivia +591) + Entrada de Teléfono
        frame_tel_input = ctk.CTkFrame(left_frame, fg_color="transparent")
        frame_tel_input.pack(fill="x", padx=20, pady=3)
        frame_tel_input.grid_columnconfigure(1, weight=1)

        self.combo_wsp_pais = ctk.CTkComboBox(
            frame_tel_input,
            values=list(PAISES_WHATSAPP.keys()),
            width=155,
            command=self.on_pais_wsp_cambiado,
            state="readonly"
        )
        self.combo_wsp_pais.set("🇧🇴 Bolivia (+591)")
        self.combo_wsp_pais.grid(row=0, column=0, padx=(0, 6), sticky="w")

        self.entry_wsp_telefono = ctk.CTkEntry(
            frame_tel_input, 
            placeholder_text="Ej: 67220995"
        )
        self.entry_wsp_telefono.grid(row=0, column=1, sticky="ew")
        self.entry_wsp_telefono.bind("<KeyRelease>", self.actualizar_preview_whatsapp)

        # Separador visual
        ctk.CTkFrame(left_frame, height=2, fg_color=["gray80", "gray30"]).pack(fill="x", padx=20, pady=10)

        # Sección 2: Plantillas de Minimarket
        ctk.CTkLabel(
            left_frame, 
            text="2. Plantillas de Negocio para Minimarket:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(4, 2))

        nombres_plantillas = list(PLANTILLAS_MINIMARKET.keys())
        self.combo_wsp_plantillas = ctk.CTkComboBox(
            left_frame, 
            values=nombres_plantillas,
            state="readonly"
        )
        self.combo_wsp_plantillas.set(nombres_plantillas[0])
        self.combo_wsp_plantillas.pack(fill="x", padx=20, pady=3)

        btn_cargar = ctk.CTkButton(
            left_frame,
            text="📋 Cargar Plantilla en el Redactor",
            font=ctk.CTkFont(size=12, weight="bold"),
            height=32,
            fg_color=["#3a7ebf", "#1f538d"],
            command=self.cargar_plantilla_whatsapp_seleccionada
        )
        btn_cargar.pack(fill="x", padx=20, pady=(6, 10))

        # Sección 3: Inserción de Variables
        ctk.CTkLabel(
            left_frame, 
            text="Variables Dinámicas (clic para insertar):", 
            font=ctk.CTkFont(size=11, weight="bold"), 
            text_color=("gray40", "gray70")
        ).pack(anchor="w", padx=20, pady=(2, 3))

        vars_frame = ctk.CTkFrame(left_frame, fg_color="transparent")
        vars_frame.pack(fill="x", padx=20, pady=(0, 10))
        vars_frame.grid_columnconfigure((0, 1), weight=1)

        ctk.CTkButton(
            vars_frame, text="+ {nombre}", height=26, font=ctk.CTkFont(size=11),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_whatsapp("{nombre}")
        ).grid(row=0, column=0, padx=2, pady=2, sticky="ew")

        ctk.CTkButton(
            vars_frame, text="+ {telefono}", height=26, font=ctk.CTkFont(size=11),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_whatsapp("{telefono}")
        ).grid(row=0, column=1, padx=2, pady=2, sticky="ew")

        ctk.CTkButton(
            vars_frame, text="+ {fecha}", height=26, font=ctk.CTkFont(size=11),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_whatsapp("{fecha}")
        ).grid(row=1, column=0, padx=2, pady=2, sticky="ew")

        ctk.CTkButton(
            vars_frame, text="+ {minimarket}", height=26, font=ctk.CTkFont(size=11),
            fg_color=["gray80", "gray30"], hover_color=["gray70", "gray40"], text_color=("black", "white"),
            command=lambda: self.insertar_variable_whatsapp("{minimarket}")
        ).grid(row=1, column=1, padx=2, pady=2, sticky="ew")

        # Panel Derecho: Redactor, Vista Previa WhatsApp y Envío
        right_frame = ctk.CTkFrame(frame)
        right_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0), pady=0)

        ctk.CTkLabel(
            right_frame, 
            text="Redactar Mensaje WhatsApp", 
            font=ctk.CTkFont(size=18, weight="bold"), 
            text_color=("black", "white")
        ).pack(pady=(15, 4))

        # Editor de texto
        self.txt_wsp_mensaje = ctk.CTkTextbox(right_frame, height=160, font=ctk.CTkFont(size=12))
        self.txt_wsp_mensaje.pack(fill="x", padx=20, pady=(0, 6))
        self.txt_wsp_mensaje.bind("<KeyRelease>", self.actualizar_preview_whatsapp)
        self.txt_wsp_mensaje.insert("1.0", PLANTILLAS_MINIMARKET[nombres_plantillas[0]])

        # Vista Previa (Burbuja estilo WhatsApp)
        ctk.CTkLabel(
            right_frame, 
            text="📱 Vista Previa del Cliente:", 
            font=ctk.CTkFont(size=12, weight="bold"), 
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(2, 2))

        self.bubble_container = ctk.CTkFrame(
            right_frame,
            fg_color=["#dcf8c6", "#0b4d3c"],
            corner_radius=12,
            border_width=1,
            border_color=["#bfe7a3", "#073b2d"]
        )
        self.bubble_container.pack(fill="x", padx=20, pady=(2, 8))

        self.lbl_wsp_preview = ctk.CTkLabel(
            self.bubble_container,
            text="",
            font=ctk.CTkFont(size=11),
            justify="left",
            wraplength=440,
            text_color=["#111b21", "#e9edef"]
        )
        self.lbl_wsp_preview.pack(fill="both", expand=True, padx=12, pady=8)

        # Botón Principal de Envío
        self.btn_enviar_wsp_manual = ctk.CTkButton(
            right_frame,
            text="🚀  Enviar Mensaje por WhatsApp",
            font=ctk.CTkFont(size=14, weight="bold"),
            height=44,
            fg_color="#25D366",
            hover_color="#1ea952",
            text_color="white",
            command=self.ejecutar_envio_whatsapp_manual
        )
        self.btn_enviar_wsp_manual.pack(fill="x", padx=20, pady=(4, 8))

        # Log de Envíos de WhatsApp
        ctk.CTkLabel(
            right_frame,
            text="Registro de Envíos de la Sesión:",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w", padx=20, pady=(0, 2))

        self.txt_wsp_log = ctk.CTkTextbox(right_frame, height=80, font=ctk.CTkFont(size=11))
        self.txt_wsp_log.pack(fill="both", expand=True, padx=20, pady=(0, 15))

        # Actualizar vista previa inicial
        self.actualizar_preview_whatsapp()

        return frame

    def actualizar_combobox_clientes_whatsapp(self):
        """Llena el combobox de clientes activos en la vista de envíos WhatsApp."""
        clientes = self.db.get_all_clientes()
        opciones = ["(Ingreso manual de teléfono)", "🔍 Buscar Cliente..."]
        self._cache_clientes_wsp = {}

        for c in clientes:
            # REGLA CRÍTICA: Solo clientes Activos
            if (c.get("estado") or "").strip().lower() != "activo":
                continue
            tel = (c.get("telefono") or "").strip()
            nom = (c.get("nombre") or "").strip()
            cid = c.get("id")
            cat = c.get("categoria") or "🥉 Bronce"
            if tel:
                label = f"#{cid} {nom} [{cat}] - {tel}"
                opciones.append(label)
                self._cache_clientes_wsp[label] = c

        self.combo_wsp_clientes.configure(values=opciones)
        if len(opciones) > 2:
            self.combo_wsp_clientes.set(opciones[2])
            self.on_cliente_wsp_seleccionado(opciones[2])
        else:
            self.combo_wsp_clientes.set(opciones[0])

    def formatear_numero_con_pais(self, numero_raw: str) -> Tuple[str, str]:
        """
        Formatea el número asegurando el prefijo internacional del país seleccionado.
        Predeterminado para Bolivia (+591).
        Retorna (numero_limpio_para_whatsapp, numero_display).
        Ejemplo: '67220995' con Bolivia -> ('59167220995', '+591 67220995')
        """
        raw = str(numero_raw or "").strip()
        solo_digitos = re.sub(r'\D', '', raw)

        pais_sel = getattr(self, "combo_wsp_pais", None)
        pais_texto = pais_sel.get() if pais_sel else "🇧🇴 Bolivia (+591)"
        prefijo_defecto = PAISES_WHATSAPP.get(pais_texto, "+591").replace("+", "").strip()

        if not solo_digitos:
            return "", ""

        # Si el usuario ya escribió con el prefijo seleccionado (ej: 59167220995)
        if prefijo_defecto and solo_digitos.startswith(prefijo_defecto) and len(solo_digitos) > len(prefijo_defecto) + 5:
            num_final = solo_digitos
        elif prefijo_defecto and len(solo_digitos) <= 10:
            # Número local (ej. 8 dígitos de Bolivia como 67220995) -> Anteponer prefijo
            num_final = f"{prefijo_defecto}{solo_digitos}"
        else:
            num_final = solo_digitos

        if prefijo_defecto and num_final.startswith(prefijo_defecto):
            local_num = num_final[len(prefijo_defecto):]
            display = f"+{prefijo_defecto} {local_num}"
        else:
            display = f"+{num_final}" if num_final else ""

        return num_final, display

    def on_pais_wsp_cambiado(self, seleccion: str):
        """Callback al cambiar de país en el combo."""
        self.actualizar_preview_whatsapp()

    def on_cliente_wsp_seleccionado(self, seleccion: str):
        if "🔍 Buscar Cliente" in seleccion:
            self.abrir_modal_buscar_cliente_wsp()
            return

        if hasattr(self, "_cache_clientes_wsp") and seleccion in self._cache_clientes_wsp:
            c = self._cache_clientes_wsp[seleccion]
            tel_raw = (c.get("telefono") or "").strip()
            digitos = re.sub(r'\D', '', tel_raw)

            # Detectar país automáticamente con el helper normalizar_pais
            pais_norm = normalizar_pais("", tel_raw)
            if hasattr(self, "combo_wsp_pais"):
                self.combo_wsp_pais.set(pais_norm)
            prefijo = PAISES_WHATSAPP.get(pais_norm, "+591").replace("+", "").strip()

            if prefijo and digitos.startswith(prefijo) and len(digitos) > len(prefijo):
                num_local = digitos[len(prefijo):]
            else:
                num_local = digitos if digitos else tel_raw

            self.entry_wsp_telefono.delete(0, "end")
            self.entry_wsp_telefono.insert(0, num_local)

        elif seleccion == "(Ingreso manual de teléfono)":
            self.entry_wsp_telefono.delete(0, "end")

        self.actualizar_preview_whatsapp()

    def abrir_modal_buscar_cliente_wsp(self):
        """Abre ventana modal para buscar clientes activos y cargarlos como destinatario en el módulo WhatsApp."""
        modal = ctk.CTkToplevel(self)
        modal.title("🔍 Buscar Destinatario - WhatsApp")
        modal.geometry("750x540")
        modal.minsize(630, 420)
        modal.transient(self)
        modal.grab_set()

        try:
            x = self.winfo_x() + (self.winfo_width() - 750) // 2
            y = self.winfo_y() + (self.winfo_height() - 540) // 2
            modal.geometry(f"+{max(0, x)}+{max(0, y)}")
        except Exception:
            pass

        # Cabecera
        top_frame = ctk.CTkFrame(modal, fg_color="transparent")
        top_frame.pack(fill="x", padx=20, pady=(15, 8))

        ctk.CTkLabel(
            top_frame,
            text="🔍 Seleccionar Destinatario para WhatsApp",
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w")

        ctk.CTkLabel(
            top_frame,
            text="Filtre por nombre, teléfono o correo electrónico para cargar al cliente de inmediato.",
            font=ctk.CTkFont(size=12),
            text_color=("gray40", "gray70")
        ).pack(anchor="w", pady=(2, 6))

        # Buscador en vivo
        entry_buscar = ctk.CTkEntry(
            modal,
            placeholder_text="🔎 Escriba Nombre, Teléfono o Correo...",
            height=38,
            font=ctk.CTkFont(size=13)
        )
        entry_buscar.pack(fill="x", padx=20, pady=(0, 6))
        entry_buscar.focus_set()

        lbl_conteo = ctk.CTkLabel(
            modal,
            text="Cargando clientes activos...",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("gray40", "gray60")
        )
        lbl_conteo.pack(anchor="w", padx=22, pady=(0, 6))

        scroll_modal = ctk.CTkScrollableFrame(modal)
        scroll_modal.pack(fill="both", expand=True, padx=20, pady=(0, 10))

        # Clientes activos
        todos_clientes = [
            c for c in self.db.get_all_clientes()
            if (c.get("estado") or "").strip().lower() == "activo"
        ]

        def seleccionar_y_cerrar(cliente_datos):
            cid = cliente_datos["id"]
            nom = cliente_datos["nombre"]
            tel_raw = (cliente_datos.get("telefono") or "").strip()
            cat = cliente_datos.get("categoria") or "🥉 Bronce"
            label = f"#{cid} {nom} [{cat}] - {tel_raw}"

            if not hasattr(self, "_cache_clientes_wsp"):
                self._cache_clientes_wsp = {}
            self._cache_clientes_wsp[label] = cliente_datos

            valores_actuales = list(self.combo_wsp_clientes.cget("values"))
            if label not in valores_actuales:
                valores_actuales.append(label)
                self.combo_wsp_clientes.configure(values=valores_actuales)

            self.combo_wsp_clientes.set(label)
            self.on_cliente_wsp_seleccionado(label)
            modal.destroy()

        def renderizar_resultados(query: str = ""):
            for w in scroll_modal.winfo_children():
                w.destroy()

            q = query.strip().lower()
            if not q:
                filtrados = todos_clientes
            else:
                filtrados = []
                for c in todos_clientes:
                    nom = str(c.get("nombre") or "").lower()
                    cor = str(c.get("correo") or "").lower()
                    tel = str(c.get("telefono") or "").lower()
                    cid = str(c.get("id") or "")
                    if q in nom or q in cor or q in tel or q == cid:
                        filtrados.append(c)

            lbl_conteo.configure(text=f"Clientes encontrados: {len(filtrados)} de {len(todos_clientes)} clientes activos")

            if not filtrados:
                ctk.CTkLabel(
                    scroll_modal,
                    text="❌ No se encontraron clientes activos con esa búsqueda.",
                    font=ctk.CTkFont(size=13),
                    text_color=("gray40", "gray60")
                ).pack(pady=30)
                return

            for c in filtrados:
                item_frame = ctk.CTkFrame(scroll_modal)
                item_frame.pack(fill="x", padx=4, pady=4)
                item_frame.grid_columnconfigure(0, weight=1)

                cat = c.get("categoria") or "🥉 Bronce"
                correo_txt = c['correo'] or 'Sin correo'
                tel_txt = c['telefono'] or 'Sin tel'
                info = f"ID: {c['id']} | {c['nombre']}  •  📱 {tel_txt}  •  📧 {correo_txt}"

                lbl_info = ctk.CTkLabel(
                    item_frame,
                    text=info,
                    font=ctk.CTkFont(size=12, weight="bold"),
                    anchor="w",
                    text_color=("black", "white")
                )
                lbl_info.grid(row=0, column=0, padx=10, pady=8, sticky="w")

                lbl_cat = ctk.CTkLabel(
                    item_frame,
                    text=f" {cat} ",
                    font=ctk.CTkFont(size=11, weight="bold"),
                    text_color="#3498db"
                )
                lbl_cat.grid(row=0, column=1, padx=4, pady=8)

                btn_sel = ctk.CTkButton(
                    item_frame,
                    text="Seleccionar 💬",
                    width=100,
                    fg_color="#25D366",
                    hover_color="#1ea952",
                    command=lambda cli=c: seleccionar_y_cerrar(cli)
                )
                btn_sel.grid(row=0, column=2, padx=(6, 10), pady=6)

        entry_buscar.bind("<KeyRelease>", lambda event: renderizar_resultados(entry_buscar.get()))
        renderizar_resultados("")

        bot_frame = ctk.CTkFrame(modal, fg_color="transparent")
        bot_frame.pack(fill="x", padx=20, pady=(0, 12))
        ctk.CTkButton(
            bot_frame,
            text="Cerrar",
            width=100,
            fg_color="gray40",
            hover_color="gray30",
            command=modal.destroy
        ).pack(side="right")

    def cargar_plantilla_whatsapp_seleccionada(self):
        plantilla_nombre = self.combo_wsp_plantillas.get()
        if plantilla_nombre in PLANTILLAS_MINIMARKET:
            texto = PLANTILLAS_MINIMARKET[plantilla_nombre]
            self.txt_wsp_mensaje.delete("1.0", "end")
            self.txt_wsp_mensaje.insert("1.0", texto)
            self.actualizar_preview_whatsapp()

    def insertar_variable_whatsapp(self, var: str):
        self.txt_wsp_mensaje.insert("insert", var)
        self.actualizar_preview_whatsapp()

    def actualizar_preview_whatsapp(self, event=None):
        if not hasattr(self, "txt_wsp_mensaje"):
            return
        raw_text = self.txt_wsp_mensaje.get("1.0", "end-1c")
        nom = "Cliente"
        if hasattr(self, "combo_wsp_clientes"):
            sel = self.combo_wsp_clientes.get()
            if hasattr(self, "_cache_clientes_wsp") and sel in self._cache_clientes_wsp:
                nom = self._cache_clientes_wsp[sel].get("nombre", "Cliente")

        raw_tel = self.entry_wsp_telefono.get().strip() if hasattr(self, "entry_wsp_telefono") else ""
        _, tel_display = self.formatear_numero_con_pais(raw_tel)
        tel_final = tel_display or "+591 67220995"
        fecha_actual = datetime.now().strftime("%d/%m/%Y")

        preview_text = (
            raw_text
            .replace("{nombre}", nom)
            .replace("{telefono}", tel_final)
            .replace("{fecha}", fecha_actual)
            .replace("{minimarket}", "Minimarket")
        )
        if hasattr(self, "lbl_wsp_preview"):
            self.lbl_wsp_preview.configure(text=preview_text)

    def log_envios_wsp(self, mensaje: str):
        hora = datetime.now().strftime("%H:%M:%S")
        self.txt_wsp_log.insert("end", f"[{hora}] {mensaje}\n")
        self.txt_wsp_log.see("end")

    def ejecutar_envio_whatsapp_manual(self):
        raw_num = self.entry_wsp_telefono.get().strip()
        numero_whatsapp, numero_display = self.formatear_numero_con_pais(raw_num)
        cuerpo_base = self.txt_wsp_mensaje.get("1.0", "end-1c").strip()

        if not numero_whatsapp or len(numero_whatsapp) < 8:
            messagebox.showwarning(
                "Número Inválido", 
                "Por favor ingrese un número de teléfono válido.\n(Para Bolivia se auto-completará con el prefijo +591)."
            )
            return

        if not cuerpo_base:
            messagebox.showwarning("Mensaje Vacío", "Por favor escriba un mensaje o cargue una plantilla.")
            return

        # REGLA CRÍTICA: Impedir envío si el número corresponde a un cliente inactivo
        for c in self.db.get_all_clientes():
            if (c.get("estado") or "").strip().lower() == "inactivo":
                tel_c = re.sub(r'\D', '', str(c.get("telefono") or ""))
                if tel_c and (tel_c in numero_whatsapp or numero_whatsapp.endswith(tel_c)):
                    messagebox.showwarning(
                        "Cliente Inactivo", 
                        f"⛔ El número ingresado corresponde al cliente '{c.get('nombre')}' que está con estado 'Inactivo'.\n\n"
                        "El sistema tiene restringido el envío de mensajes a clientes inactivos.\n"
                        "Para enviarle mensajes, actívelo primero en el módulo Gestión Clientes."
                    )
                    return

        nom = "Cliente"
        sel = self.combo_wsp_clientes.get()
        cliente_id = None
        if hasattr(self, "_cache_clientes_wsp") and sel in self._cache_clientes_wsp:
            cliente_obj = self._cache_clientes_wsp[sel]
            nom = cliente_obj.get("nombre", "Cliente")
            cliente_id = cliente_obj.get("id")

        fecha_actual = datetime.now().strftime("%d/%m/%Y")
        mensaje_final = (
            cuerpo_base
            .replace("{nombre}", nom)
            .replace("{telefono}", numero_display)
            .replace("{fecha}", fecha_actual)
            .replace("{minimarket}", "Minimarket")
        )

        self.btn_enviar_wsp_manual.configure(state="disabled", text="⏳ Enviando por WhatsApp...")
        self.log_envios_wsp(f"Iniciando despacho a {numero_display} ({nom})...")

        def tarea():
            try:
                gestor = self.obtener_gestor_whatsapp()
                exito, estado = gestor.enviar_mensaje_whatsapp(numero=numero_whatsapp, mensaje=mensaje_final)
            except Exception as err:
                exito = False
                estado = f"Fallo al comunicar con WhatsApp: {err}"

            if cliente_id:
                self.db.create_historial(cliente_id=cliente_id, estado=f"WhatsApp Manual: {estado}")

            def on_finalizado():
                self.btn_enviar_wsp_manual.configure(state="normal", text="🚀  Enviar Mensaje por WhatsApp")
                if exito:
                    self.log_envios_wsp(f"✅ Enviado con éxito a {numero_display} ({nom})")
                    messagebox.showinfo("Envío Exitoso", f"El mensaje fue enviado correctamente a:\n{nom} ({numero_display})")
                else:
                    self.log_envios_wsp(f"❌ Error al enviar a {numero_display}: {estado}")
                    messagebox.showerror("Fallo en Envío", f"No se pudo enviar el mensaje a {numero_display}:\n{estado}")

            self.after(0, on_finalizado)

        threading.Thread(target=tarea, daemon=True).start()

    # ------------------------------------------------------------------
    # VISTA 3: REGLAS DE AUTOMATIZACIÓN (Con Cola de Envíos y Tiempos Anti-Baneo)
    # ------------------------------------------------------------------
    def construir_vista_reglas(self) -> ctk.CTkFrame:
        frame = ctk.CTkFrame(self.container, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_columnconfigure(1, weight=2)
        frame.grid_rowconfigure(0, weight=1)

        # Perfiles de Tiempos Anti-Baneo recomendados por el sistema
        self.perfiles_antiban: Dict[str, Optional[Tuple[int, int]]] = {
            "🛡️ Humano Seguro (15-30s) [Recomendado]": (15, 30),
            "🐢 Ultra Seguro (30-60s) [Cuentas Nuevas / Masivo]": (30, 60),
            "⚡ Moderado (8-15s) [Contactos Frecuentes]": (8, 15),
            "⚙️ Personalizado (Definir manualmente)": None
        }

        # Columna Izquierda: Formulario de Regla
        form_frame = ctk.CTkFrame(frame)
        form_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10), pady=0)

        ctk.CTkLabel(form_frame, text="Nueva Regla de Automatización", font=ctk.CTkFont(size=18, weight="bold"), text_color=("black", "white")).pack(pady=(15, 8))

        self.entry_regla_nombre = ctk.CTkEntry(form_frame, placeholder_text="Nombre (ej. Promo Fin de Semana)")
        self.entry_regla_nombre.pack(fill="x", padx=20, pady=4)

        # ComboBox para elegir si el envío se hará por 'Correo SMTP' o 'WhatsApp'
        ctk.CTkLabel(form_frame, text="Canal de Envío (Tipo):", font=ctk.CTkFont(size=12), text_color=("black", "white")).pack(anchor="w", padx=20, pady=(4, 2))
        self.combo_regla_tipo = ctk.CTkComboBox(form_frame, values=["WhatsApp", "Correo SMTP"], state="readonly")
        self.combo_regla_tipo.set("WhatsApp")
        self.combo_regla_tipo.pack(fill="x", padx=20, pady=4)

        ctk.CTkLabel(form_frame, text="Condición (ej. Activo, Inactivo, VIP, todos):", font=ctk.CTkFont(size=12), text_color=("black", "white")).pack(anchor="w", padx=20, pady=(4, 2))
        self.entry_regla_condicion = ctk.CTkEntry(form_frame, placeholder_text="Condición que debe cumplir el cliente")
        self.entry_regla_condicion.insert(0, "Activo")
        self.entry_regla_condicion.pack(fill="x", padx=20, pady=4)

        ctk.CTkLabel(form_frame, text="Mensaje (Usa {nombre}, {correo}, {telefono}, {estado}):", font=ctk.CTkFont(size=12), text_color=("black", "white")).pack(anchor="w", padx=20, pady=(4, 2))
        self.text_regla_mensaje = ctk.CTkTextbox(form_frame, height=110)
        self.text_regla_mensaje.insert("1.0", "Hola {nombre}, tenemos una oferta especial para ti!")
        self.text_regla_mensaje.pack(fill="x", padx=20, pady=4)

        self.switch_regla_activa = ctk.CTkSwitch(form_frame, text="Regla Activa", text_color=("black", "white"))
        self.switch_regla_activa.select()
        self.switch_regla_activa.pack(anchor="w", padx=20, pady=8)

        self.btn_guardar_regla = ctk.CTkButton(
            form_frame, text="💾 Guardar Regla", fg_color="#27ae60", hover_color="#219955",
            command=self.guardar_regla
        )
        self.btn_guardar_regla.pack(fill="x", padx=20, pady=(8, 6))

        self.btn_abrir_admin_reglas = ctk.CTkButton(
            form_frame,
            text="⚙️ Administrar Reglas (Editar / Eliminar)",
            fg_color="#34495e",
            hover_color="#2c3e50",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.abrir_ventana_admin_reglas
        )
        self.btn_abrir_admin_reglas.pack(fill="x", padx=20, pady=(0, 15))

        # Columna Derecha: Lista de Reglas, Configuración Anti-Baneo y Monitor de Cola
        lista_frame = ctk.CTkFrame(frame)
        lista_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0), pady=0)
        lista_frame.grid_columnconfigure(0, weight=1)

        # 1. Barra superior con botones de Acción
        header_frame = ctk.CTkFrame(lista_frame, fg_color="transparent")
        header_frame.pack(fill="x", padx=15, pady=(12, 6))

        ctk.CTkLabel(
            header_frame, 
            text="⚡ Automatización y Despacho Anti-Baneo", 
            font=ctk.CTkFont(size=18, weight="bold"), 
            text_color=("black", "white")
        ).pack(side="left")

        self.btn_detener_auto = ctk.CTkButton(
            header_frame, 
            text="⏹️ Detener Cola", 
            width=110,
            fg_color="#c0392b", 
            hover_color="#962d22", 
            font=ctk.CTkFont(size=12, weight="bold"),
            state="disabled",
            command=self.detener_cola_automatica
        )
        self.btn_detener_auto.pack(side="right", padx=(6, 0))

        self.btn_ejecutar_auto = ctk.CTkButton(
            header_frame, 
            text="🚀 Iniciar Cola de Envíos", 
            width=170,
            fg_color="#e67e22", 
            hover_color="#d35400", 
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.iniciar_ejecucion_automatica
        )
        self.btn_ejecutar_auto.pack(side="right")

        # 2. Tarjeta de Configuración de Tiempos Anti-Baneo
        card_tiempos = ctk.CTkFrame(lista_frame, fg_color=["gray90", "gray20"], corner_radius=8)
        card_tiempos.pack(fill="x", padx=15, pady=(0, 6))
        card_tiempos.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(
            card_tiempos,
            text="🛡️ Configuración de Intervalos Humanos Anti-Baneo (WhatsApp / Email):",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w", padx=12, pady=(8, 4))

        controles_frame = ctk.CTkFrame(card_tiempos, fg_color="transparent")
        controles_frame.pack(fill="x", padx=12, pady=(0, 4))
        controles_frame.grid_columnconfigure((1, 3, 5), weight=1)

        ctk.CTkLabel(controles_frame, text="Perfil:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).grid(row=0, column=0, padx=(0, 4), sticky="w")
        self.combo_perfil_antiban = ctk.CTkComboBox(
            controles_frame,
            values=list(self.perfiles_antiban.keys()),
            width=230,
            command=self.on_cambio_perfil_antiban,
            state="readonly"
        )
        self.combo_perfil_antiban.grid(row=0, column=1, padx=(0, 10), sticky="ew")

        ctk.CTkLabel(controles_frame, text="Min (seg):", font=ctk.CTkFont(size=11), text_color=("black", "white")).grid(row=0, column=2, padx=(0, 4), sticky="w")
        self.entry_delay_min = ctk.CTkEntry(controles_frame, width=50)
        self.entry_delay_min.grid(row=0, column=3, padx=(0, 8), sticky="w")

        ctk.CTkLabel(controles_frame, text="Max (seg):", font=ctk.CTkFont(size=11), text_color=("black", "white")).grid(row=0, column=4, padx=(0, 4), sticky="w")
        self.entry_delay_max = ctk.CTkEntry(controles_frame, width=50)
        self.entry_delay_max.grid(row=0, column=5, padx=(0, 8), sticky="w")

        self.btn_guardar_tiempos = ctk.CTkButton(
            controles_frame,
            text="💾 Guardar Tiempos",
            width=115,
            height=28,
            fg_color="#27ae60",
            hover_color="#219955",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.guardar_tiempos_antiban
        )
        self.btn_guardar_tiempos.grid(row=0, column=6, sticky="e")

        # Cargar tiempos guardados de la base de datos
        tiempos_guardados = self.db.obtener_tiempos_automatizacion()
        self.combo_perfil_antiban.set(tiempos_guardados["perfil"])
        self.entry_delay_min.insert(0, str(tiempos_guardados["min_seg"]))
        self.entry_delay_max.insert(0, str(tiempos_guardados["max_seg"]))

        # Banner informativo de recomendación del sistema
        rec_label = ctk.CTkLabel(
            card_tiempos,
            text="💡 Recomendación del Sistema: WhatsApp penaliza o banea cuentas que envían ráfagas automáticas instantáneas. Un intervalo de 15 a 30 segundos con variación aleatoria simula la velocidad natural de un ser humano, protegiendo tu número contra detecciones de spam.",
            font=ctk.CTkFont(size=10),
            text_color=("gray40", "gray70"),
            wraplength=520,
            justify="left"
        )
        rec_label.pack(anchor="w", padx=12, pady=(0, 8))

        # 3. Sección: Destinatarios a Enviar (Cola Personalizada con Asignación de Reglas)
        card_destinatarios = ctk.CTkFrame(lista_frame, fg_color=["gray90", "gray20"], corner_radius=8)
        card_destinatarios.pack(fill="x", padx=15, pady=(0, 6))

        # Encabezado y contador
        dest_header = ctk.CTkFrame(card_destinatarios, fg_color="transparent")
        dest_header.pack(fill="x", padx=12, pady=(6, 2))

        ctk.CTkLabel(
            dest_header,
            text="👥 Destinatarios a Enviar (Asignación de Regla por Cliente):",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=("black", "white")
        ).pack(side="left")

        self.lbl_contador_dest = ctk.CTkLabel(
            dest_header,
            text="Destinatarios en cola: 0",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color="#e67e22"
        )
        self.lbl_contador_dest.pack(side="right")

        # Selector interactivo de Canal de Despacho (WhatsApp / Correo SMTP)
        frame_canal_despacho = ctk.CTkFrame(card_destinatarios, fg_color=["gray82", "gray28"], corner_radius=6)
        frame_canal_despacho.pack(fill="x", padx=12, pady=(2, 6))

        ctk.CTkLabel(
            frame_canal_despacho,
            text="📡 Canal de Envío:",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("black", "white")
        ).pack(side="left", padx=(10, 6), pady=5)

        self.seg_canal_auto = ctk.CTkSegmentedButton(
            frame_canal_despacho,
            values=["📲 WhatsApp", "✉️ Correo SMTP"],
            command=self.al_cambiar_canal_despacho_auto,
            font=ctk.CTkFont(size=11, weight="bold"),
            selected_color="#2980b9",
            selected_hover_color="#1f618d"
        )
        self.seg_canal_auto.set("📲 WhatsApp")
        self.seg_canal_auto.pack(side="left", padx=4, pady=5)

        self.lbl_info_canal_auto = ctk.CTkLabel(
            frame_canal_despacho,
            text="• Envíos masivos se realizarán por WhatsApp Web con pausas anti-bloqueo.",
            font=ctk.CTkFont(size=10),
            text_color=("#155724", "#4ade80")
        )
        self.lbl_info_canal_auto.pack(side="left", padx=(10, 8), pady=5)

        # Fila de controles para agregar clientes y asignar regla
        controles_dest = ctk.CTkFrame(card_destinatarios, fg_color="transparent")
        controles_dest.pack(fill="x", padx=12, pady=(2, 4))
        controles_dest.grid_columnconfigure((0, 1), weight=1)

        # Combo Clientes
        frame_sel_cli = ctk.CTkFrame(controles_dest, fg_color="transparent")
        frame_sel_cli.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ctk.CTkLabel(frame_sel_cli, text="Seleccionar Cliente:", font=ctk.CTkFont(size=11), text_color=("black", "white")).pack(anchor="w")
        self.combo_dest_cliente = ctk.CTkComboBox(frame_sel_cli, values=["(Sin clientes)"], state="readonly")
        self.combo_dest_cliente.pack(fill="x", pady=(2, 0))

        # Combo Reglas
        frame_sel_reg = ctk.CTkFrame(controles_dest, fg_color="transparent")
        frame_sel_reg.grid(row=0, column=1, sticky="ew", padx=(6, 6))
        ctk.CTkLabel(frame_sel_reg, text="Regla a Enviar a este Cliente:", font=ctk.CTkFont(size=11), text_color=("black", "white")).pack(anchor="w")
        self.combo_dest_regla = ctk.CTkComboBox(frame_sel_reg, values=["(Sin reglas)"], state="readonly")
        self.combo_dest_regla.pack(fill="x", pady=(2, 0))

        # Botones de Acción para la Cola
        frame_btn_dest = ctk.CTkFrame(controles_dest, fg_color="transparent")
        frame_btn_dest.grid(row=0, column=2, sticky="e", padx=(6, 0))
        ctk.CTkLabel(frame_btn_dest, text="Acciones:", font=ctk.CTkFont(size=11), text_color=("black", "white")).pack(anchor="w")

        btn_box = ctk.CTkFrame(frame_btn_dest, fg_color="transparent")
        btn_box.pack(pady=(2, 0))

        self.btn_agregar_dest = ctk.CTkButton(
            btn_box,
            text="➕ Agregar",
            width=80,
            height=28,
            fg_color="#27ae60",
            hover_color="#219955",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.agregar_cliente_a_cola
        )
        self.btn_agregar_dest.pack(side="left", padx=(0, 4))

        self.btn_masivo_dest = ctk.CTkButton(
            btn_box,
            text="⚡ Masivo",
            width=75,
            height=28,
            fg_color="#3498db",
            hover_color="#2980b9",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.agregar_clientes_masivo_por_condicion
        )
        self.btn_masivo_dest.pack(side="left", padx=(0, 4))

        self.btn_vaciar_dest = ctk.CTkButton(
            btn_box,
            text="🧹 Vaciar",
            width=70,
            height=28,
            fg_color="#7f8c8d",
            hover_color="#636e72",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.vaciar_cola_destinatarios
        )
        self.btn_vaciar_dest.pack(side="left")

        # Lista deslizable de destinatarios agregados
        self.scroll_destinatarios_cola = ctk.CTkScrollableFrame(card_destinatarios, height=120)
        self.scroll_destinatarios_cola.pack(fill="both", expand=True, padx=12, pady=(4, 8))

        # 4. Monitor de Cola de Envíos en Tiempo Real
        card_cola = ctk.CTkFrame(lista_frame, fg_color=["gray90", "gray20"], corner_radius=8)
        card_cola.pack(fill="both", expand=True, padx=15, pady=(0, 12))

        ctk.CTkLabel(
            card_cola,
            text="📊 Monitor de Cola de Envíos en Tiempo Real:",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w", padx=12, pady=(6, 2))

        # Estado dinámico de la cola
        self.lbl_cola_estado = ctk.CTkLabel(
            card_cola,
            text="⚪ Cola inactiva (Presione 'Iniciar Cola de Envíos' para despachar mensajes)",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("gray40", "gray70"),
            anchor="w"
        )
        self.lbl_cola_estado.pack(fill="x", padx=12, pady=(0, 4))

        # Barra de progreso
        self.progress_cola = ctk.CTkProgressBar(card_cola, height=12, corner_radius=6)
        self.progress_cola.set(0)
        self.progress_cola.pack(fill="x", padx=12, pady=(0, 4))

        # Métricas de cola
        self.lbl_cola_metricas = ctk.CTkLabel(
            card_cola,
            text="Progreso: 0/0  |  ✅ Enviados: 0  |  ❌ Fallidos: 0  |  ⏳ En cola: 0",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=("black", "white"),
            anchor="w"
        )
        self.lbl_cola_metricas.pack(fill="x", padx=12, pady=(0, 4))

        # Registro en vivo de la cola
        self.txt_cola_log = ctk.CTkTextbox(card_cola, height=85, font=ctk.CTkFont(size=11))
        self.txt_cola_log.pack(fill="both", expand=True, padx=12, pady=(0, 8))

        # Inicializar selectores y cola de destinatarios
        self.actualizar_selectores_destinatarios()
        self.renderizar_cola_destinatarios()

        return frame

    def on_cambio_perfil_antiban(self, perfil: str):
        """Ajusta automáticamente los segundos mínimos y máximos según el perfil seleccionado."""
        if hasattr(self, "perfiles_antiban") and perfil in self.perfiles_antiban:
            rango = self.perfiles_antiban[perfil]
            if rango is not None:
                min_v, max_v = rango
                self.entry_delay_min.delete(0, "end")
                self.entry_delay_min.insert(0, str(min_v))
                self.entry_delay_max.delete(0, "end")
                self.entry_delay_max.insert(0, str(max_v))

    def guardar_tiempos_antiban(self):
        """Guarda la configuración de intervalos de tiempo anti-baneo en la base de datos."""
        perfil = self.combo_perfil_antiban.get()
        try:
            min_v = int(self.entry_delay_min.get().strip())
            max_v = int(self.entry_delay_max.get().strip())
        except ValueError:
            messagebox.showerror("Tiempos Inválidos", "Los tiempos mínimo y máximo deben ser números enteros en segundos.")
            return

        if min_v < 1 or max_v < 1:
            messagebox.showerror("Tiempos Inválidos", "Los tiempos deben ser mayores o iguales a 1 segundo.")
            return

        if min_v > max_v:
            min_v, max_v = max_v, min_v
            self.entry_delay_min.delete(0, "end")
            self.entry_delay_min.insert(0, str(min_v))
            self.entry_delay_max.delete(0, "end")
            self.entry_delay_max.insert(0, str(max_v))

        exito = self.db.guardar_tiempos_automatizacion(perfil=perfil, min_seg=min_v, max_seg=max_v)
        if exito:
            messagebox.showinfo(
                "Tiempos Guardados",
                f"Configuración de intervalos guardada exitosamente:\n\n"
                f"Perfil: {perfil}\n"
                f"Rango de espera aleatoria: {min_v} a {max_v} segundos.\n\n"
                "Se aplicará a todas las automatizaciones en curso y futuras."
            )
        else:
            messagebox.showerror("Error", "No se pudo guardar la configuración de tiempos en la base de datos.")

    def actualizar_selectores_destinatarios(self):
        """Actualiza los desplegables de selección de cliente y regla en la vista de automatización."""
        clientes = self.db.get_all_clientes()
        self._cache_clientes_reglas = {}
        valores_clientes = []
        for c in clientes:
            # REGLA CRÍTICA: Solo clientes Activos en automatización
            if (c.get("estado") or "").strip().lower() != "activo":
                continue
            cat = c.get("categoria") or "🥉 Bronce"
            etiqueta = f"#{c['id']} {c['nombre']} [{cat}] ({c.get('telefono','')})"
            self._cache_clientes_reglas[etiqueta] = c
            valores_clientes.append(etiqueta)

        if hasattr(self, "combo_dest_cliente"):
            if valores_clientes:
                self.combo_dest_cliente.configure(values=valores_clientes)
                self.combo_dest_cliente.set(valores_clientes[0])
            else:
                self.combo_dest_cliente.configure(values=["(Sin clientes activos registrados)"])
                self.combo_dest_cliente.set("(Sin clientes activos registrados)")

        reglas = self.db.get_all_reglas()
        self._cache_reglas_reglas = {}
        valores_reglas = []
        for r in reglas:
            icono = "📲" if "WHATSAPP" in str(r.get("tipo", "")).upper() else "✉️"
            estado_tag = "" if r.get("activa") else " [Inactiva]"
            etiqueta = f"#{r['id']} {r['nombre_regla']} [{icono} {r['tipo']}]{estado_tag}"
            self._cache_reglas_reglas[etiqueta] = r
            valores_reglas.append(etiqueta)

        if hasattr(self, "combo_dest_regla"):
            if valores_reglas:
                self.combo_dest_regla.configure(values=valores_reglas)
                self.combo_dest_regla.set(valores_reglas[0])
            else:
                self.combo_dest_regla.configure(values=["(Sin reglas creadas)"])
                self.combo_dest_regla.set("(Sin reglas creadas)")

    def al_cambiar_canal_despacho_auto(self, valor_sel: str):
        """Manejador al cambiar el selector de canal (WhatsApp / Correo SMTP) en Automatización."""
        es_wsp = "WHATSAPP" in valor_sel.upper()
        canal = "WhatsApp" if es_wsp else "Correo SMTP"

        if hasattr(self, "lbl_info_canal_auto"):
            if es_wsp:
                self.lbl_info_canal_auto.configure(
                    text="• Envíos masivos se realizarán por WhatsApp Web con pausas anti-bloqueo.",
                    text_color=("#155724", "#4ade80")
                )
            else:
                self.lbl_info_canal_auto.configure(
                    text="• Envíos masivos se realizarán por Correo SMTP con pausas humanas anti-spam.",
                    text_color=("#2980b9", "#60a5fa")
                )

        # Si hay clientes en la cola, actualizar los clientes en estado pendiente al nuevo canal
        if self.cola_destinatarios_seleccionados:
            cambiados = 0
            for item in self.cola_destinatarios_seleccionados:
                if item.get("estado_envio") in ("pendiente", None):
                    item["canal"] = canal
                    item["es_whatsapp"] = es_wsp
                    cambiados += 1
            if cambiados > 0:
                self.renderizar_cola_destinatarios()

    def agregar_cliente_a_cola(self):
        """Agrega el cliente seleccionado con la regla elegida al canal configurado en la cola."""
        sel_cli = self.combo_dest_cliente.get()
        sel_reg = self.combo_dest_regla.get()

        if sel_cli not in self._cache_clientes_reglas:
            messagebox.showwarning("Selección Inválida", "Por favor seleccione un cliente válido de la lista.")
            return

        if sel_reg not in self._cache_reglas_reglas:
            messagebox.showwarning("Selección Inválida", "Por favor seleccione una regla válida para asignarle al cliente.")
            return

        cliente = self._cache_clientes_reglas[sel_cli]
        regla = self._cache_reglas_reglas[sel_reg]

        # REGLA CRÍTICA: Prohibir agregar clientes inactivos
        if (cliente.get("estado") or "").strip().lower() == "inactivo":
            messagebox.showwarning(
                "Cliente Inactivo",
                f"⛔ El cliente '{cliente.get('nombre')}' está registrado con estado 'Inactivo'.\n\n"
                "El sistema tiene restringido el envío o carga de clientes inactivos en automatización."
            )
            return

        canal_sel = self.seg_canal_auto.get() if hasattr(self, "seg_canal_auto") else "WhatsApp"
        es_wsp = "WHATSAPP" in canal_sel.upper()
        canal_nombre = "WhatsApp" if es_wsp else "Correo SMTP"

        if not es_wsp and not (cliente.get("correo") or "").strip():
            messagebox.showwarning(
                "Sin Correo Registrado",
                f"⚠️ El cliente '{cliente.get('nombre')}' no tiene una dirección de correo registrada en el sistema.\n\n"
                "Puede agregarlo, pero el despacho por Correo SMTP fallará a menos que se le registre un correo."
            )
        elif es_wsp and not (cliente.get("telefono") or "").strip():
            messagebox.showwarning(
                "Sin Teléfono Registrado",
                f"⚠️ El cliente '{cliente.get('nombre')}' no tiene un número telefónico registrado en el sistema.\n\n"
                "Puede agregarlo, pero el despacho por WhatsApp fallará."
            )

        # Si el cliente ya existe en cola, actualizar regla y canal asignado
        encontrado = False
        for item in self.cola_destinatarios_seleccionados:
            if item["cliente"]["id"] == cliente["id"]:
                item["regla"] = regla
                item["canal"] = canal_nombre
                item["es_whatsapp"] = es_wsp
                item["hora_envio"] = None
                item["estado_envio"] = "pendiente"
                encontrado = True
                break

        if not encontrado:
            self.cola_destinatarios_seleccionados.append({
                "cliente": cliente,
                "regla": regla,
                "canal": canal_nombre,
                "es_whatsapp": es_wsp,
                "hora_envio": None,
                "estado_envio": "pendiente"
            })

        self.renderizar_cola_destinatarios()

    def agregar_clientes_masivo_por_condicion(self):
        """Agrega masivamente a todos los clientes que cumplan un filtro asignándoles la regla y canal seleccionados."""
        sel_reg = self.combo_dest_regla.get()
        if sel_reg not in self._cache_reglas_reglas:
            messagebox.showwarning("Regla Requerida", "Por favor seleccione primero la regla que desea asignar a los clientes.")
            return

        regla = self._cache_reglas_reglas[sel_reg]
        condicion_regla = (regla.get("condicion") or "").strip().lower()

        canal_sel = self.seg_canal_auto.get() if hasattr(self, "seg_canal_auto") else "WhatsApp"
        es_wsp = "WHATSAPP" in canal_sel.upper()
        canal_nombre = "WhatsApp" if es_wsp else "Correo SMTP"

        todos = messagebox.askyesno(
            f"Asignación Masiva vía {canal_nombre}",
            f"¿Desea agregar masivamente a los clientes para envío automatizado por {canal_nombre} con la regla '{regla.get('nombre_regla')}'?\n\n"
            f"• Presione 'Sí' para agregar a TODOS los clientes activos.\n"
            f"• Presione 'No' para agregar únicamente los que cumplan la condición '{regla.get('condicion')}'."
        )

        clientes = self.db.get_all_clientes()
        if not clientes:
            messagebox.showinfo("Sin Clientes", "No hay clientes registrados en la base de datos.")
            return

        agregados = 0
        sin_dato = 0
        for c in clientes:
            estado_c = (c.get("estado") or "").strip().lower()
            if estado_c == "inactivo":
                continue

            aplica = False
            if todos:
                aplica = True
            else:
                cat_c = (c.get("categoria") or "").strip().lower()
                if condicion_regla in ["todos", "*", ""]:
                    aplica = True
                elif condicion_regla in estado_c or estado_c == condicion_regla or condicion_regla in cat_c or cat_c == condicion_regla:
                    aplica = True

            if aplica:
                if not es_wsp and not (c.get("correo") or "").strip():
                    sin_dato += 1
                elif es_wsp and not (c.get("telefono") or "").strip():
                    sin_dato += 1

                ya_existe = False
                for item in self.cola_destinatarios_seleccionados:
                    if item["cliente"]["id"] == c["id"]:
                        item["regla"] = regla
                        item["canal"] = canal_nombre
                        item["es_whatsapp"] = es_wsp
                        item["hora_envio"] = None
                        item["estado_envio"] = "pendiente"
                        ya_existe = True
                        break
                if not ya_existe:
                    self.cola_destinatarios_seleccionados.append({
                        "cliente": c,
                        "regla": regla,
                        "canal": canal_nombre,
                        "es_whatsapp": es_wsp,
                        "hora_envio": None,
                        "estado_envio": "pendiente"
                    })
                    agregados += 1

        self.renderizar_cola_destinatarios()
        dato_tipo = "correo electrónico" if not es_wsp else "número de teléfono"
        msg_extra = f"\n\n(Aviso: {sin_dato} clientes no tienen {dato_tipo} registrado)" if sin_dato > 0 else ""
        messagebox.showinfo(
            "Asignación Masiva Exitosa", 
            f"¡Excelente! Se han preparado {len(self.cola_destinatarios_seleccionados)} clientes en la cola para envío por {canal_nombre}.{msg_extra}"
        )

    def quitar_cliente_de_cola(self, indice: int):
        """Quita un cliente específico de la cola de envíos."""
        if 0 <= indice < len(self.cola_destinatarios_seleccionados):
            self.cola_destinatarios_seleccionados.pop(indice)
            self.renderizar_cola_destinatarios()

    def vaciar_cola_destinatarios(self):
        """Vacía todos los clientes preparados en la cola."""
        if not self.cola_destinatarios_seleccionados:
            return
        if messagebox.askyesno("Vaciar Cola", "¿Está seguro de quitar todos los destinatarios preparados de la cola?"):
            self.cola_destinatarios_seleccionados.clear()
            self.renderizar_cola_destinatarios()

    def _obtener_estilo_estado_destinatario(self, estado: Optional[str], hora: Optional[str] = None):
        """
        Retorna la tupla (texto, text_color, fg_color_badge) para mostrar la hora de envío
        a la izquierda del botón de eliminación en cada tarjeta de destinatario.
        """
        if estado == "enviando":
            return "📤 Enviando...", "#2980b9", ["#d6eaf8", "#1b2a4a"]
        elif estado == "enviado":
            texto_h = hora if hora else "--:--:--"
            return f"🕒 {texto_h}  ✅", "#27ae60", ["#d4efdf", "#143625"]
        elif estado == "fallido":
            texto_h = hora if hora else "--:--:--"
            return f"🕒 {texto_h}  ❌", "#e74c3c", ["#fadbd8", "#3a1a1a"]
        else:
            return "⏳ Pendiente", ("gray45", "gray65"), ["gray78", "gray20"]

    def actualizar_estado_item_destinatario(self, idx: int, estado: str, hora: Optional[str] = None):
        """Actualiza el estado y la hora de envío en la tarjeta del destinatario en tiempo real sin recargar la lista."""
        if 0 <= idx < len(self.cola_destinatarios_seleccionados):
            item = self.cola_destinatarios_seleccionados[idx]
            item["estado_envio"] = estado
            if hora:
                item["hora_envio"] = hora

            if hasattr(self, "_widgets_destinatarios") and idx < len(self._widgets_destinatarios):
                w_info = self._widgets_destinatarios[idx]
                badge = w_info.get("badge")
                lbl = w_info.get("lbl")
                texto, fg, bg = self._obtener_estilo_estado_destinatario(estado, item.get("hora_envio"))
                try:
                    if badge and badge.winfo_exists() and lbl and lbl.winfo_exists():
                        badge.configure(fg_color=bg)
                        lbl.configure(
                            text=texto, 
                            text_color=fg, 
                            font=ctk.CTkFont(size=11, weight="bold" if estado in ("enviado", "fallido", "enviando") else "normal")
                        )
                except Exception:
                    pass

    def renderizar_cola_destinatarios(self):
        """Dibuja en la interfaz la lista de clientes preparados con su regla asignada y hora de envío."""
        if not hasattr(self, "scroll_destinatarios_cola"):
            return

        for w in self.scroll_destinatarios_cola.winfo_children():
            w.destroy()

        self._widgets_destinatarios = []
        tot = len(self.cola_destinatarios_seleccionados)
        if hasattr(self, "lbl_contador_dest"):
            self.lbl_contador_dest.configure(text=f"Destinatarios en cola: {tot}")

        if tot == 0:
            ctk.CTkLabel(
                self.scroll_destinatarios_cola,
                text="ℹ️ No hay clientes en la cola.\nSeleccione un cliente y una regla arriba y pulse '➕ Agregar', o use '⚡ Masivo'.",
                text_color=("gray40", "gray60")
            ).pack(pady=20)
            return

        for idx, item in enumerate(self.cola_destinatarios_seleccionados):
            cli = item["cliente"]
            reg = item["regla"]
            est = item.get("estado_envio", "pendiente")
            hora = item.get("hora_envio")

            card = ctk.CTkFrame(self.scroll_destinatarios_cola, fg_color=["gray85", "gray25"], corner_radius=6)
            card.pack(fill="x", padx=4, pady=3)
            card.grid_columnconfigure(1, weight=1)

            canal_item = item.get("canal")
            if not canal_item:
                tipo_reg = (reg.get("tipo") or "").upper()
                canal_item = "WhatsApp" if "WHATSAPP" in tipo_reg else "Correo SMTP"

            es_wsp_item = "WHATSAPP" in canal_item.upper()
            icono = "📲" if es_wsp_item else "✉️"
            color_icono = "#25D366" if es_wsp_item else "#3498db"

            lbl_icono = ctk.CTkLabel(card, text=icono, font=ctk.CTkFont(size=16))
            lbl_icono.grid(row=0, column=0, rowspan=2, padx=(8, 6), pady=4)

            if es_wsp_item:
                contacto_txt = f"Tel: {cli.get('telefono','-')}"
            else:
                contacto_txt = f"Correo: {cli.get('correo','-')}"

            info_cli = f"{cli.get('nombre', 'Cliente')} | {contacto_txt} | {cli.get('estado','')}"
            lbl_cli = ctk.CTkLabel(card, text=info_cli, font=ctk.CTkFont(size=12, weight="bold"), anchor="w", text_color=("black", "white"))
            lbl_cli.grid(row=0, column=1, sticky="w", padx=(0, 4), pady=(4, 0))

            info_regla = f"Canal: {canal_item}  •  Regla: #{reg.get('id')} {reg.get('nombre_regla')}"
            lbl_reg = ctk.CTkLabel(card, text=info_regla, font=ctk.CTkFont(size=10, weight="bold"), text_color=color_icono, anchor="w")
            lbl_reg.grid(row=1, column=1, sticky="w", padx=(0, 4), pady=(0, 4))

            # Badge con hora exacta de envío a la izquierda del botón rojo ❌
            texto_badge, fg_badge, bg_badge = self._obtener_estilo_estado_destinatario(est, hora)
            badge_hora = ctk.CTkFrame(card, fg_color=bg_badge, corner_radius=6)
            badge_hora.grid(row=0, column=2, rowspan=2, padx=(6, 4), pady=4)

            lbl_hora = ctk.CTkLabel(
                badge_hora,
                text=texto_badge,
                font=ctk.CTkFont(size=11, weight="bold" if est in ("enviado", "fallido", "enviando") else "normal"),
                text_color=fg_badge
            )
            lbl_hora.pack(padx=8, pady=3)

            btn_quitar = ctk.CTkButton(
                card,
                text="❌",
                width=28,
                height=24,
                fg_color="#c0392b",
                hover_color="#962d22",
                font=ctk.CTkFont(size=10),
                command=lambda i=idx: self.quitar_cliente_de_cola(i)
            )
            if getattr(self, "automator", None) and getattr(self.automator, "cola_en_ejecucion", False):
                btn_quitar.configure(state="disabled")
            btn_quitar.grid(row=0, column=3, rowspan=2, padx=(2, 8), pady=4)

            self._widgets_destinatarios.append({
                "badge": badge_hora,
                "lbl": lbl_hora,
                "btn_quitar": btn_quitar
            })

    def guardar_regla(self):
        nombre = self.entry_regla_nombre.get().strip()
        tipo = self.combo_regla_tipo.get().strip()
        condicion = self.entry_regla_condicion.get().strip()
        mensaje = self.text_regla_mensaje.get("1.0", "end").strip()
        activa = bool(self.switch_regla_activa.get())

        if not nombre or not condicion or not mensaje:
            messagebox.showwarning("Campos Requeridos", "Debe completar el nombre, condición y mensaje de la regla.")
            return

        rid = self.db.create_regla(nombre_regla=nombre, tipo=tipo, condicion=condicion, mensaje=mensaje, activa=activa)
        if rid:
            messagebox.showinfo("Éxito", f"Regla #{rid} guardada correctamente.")
            self.entry_regla_nombre.delete(0, "end")
            self.actualizar_selectores_destinatarios()
        else:
            messagebox.showerror("Error", "No se pudo guardar la regla.")

    def abrir_ventana_admin_reglas(self):
        """Abre una ventana flotante para administrar, editar, activar/desactivar y eliminar reglas."""
        modal = ctk.CTkToplevel(self)
        modal.title("⚙️ Administración y Edición de Reglas de Automatización")
        modal.geometry("900x620")
        modal.minsize(820, 540)
        modal.transient(self)
        modal.after(100, modal.lift)

        # Centrar ventana modal
        modal.update_idletasks()
        x = self.winfo_x() + max(0, (self.winfo_width() - 900) // 2)
        y = self.winfo_y() + max(0, (self.winfo_height() - 620) // 2)
        modal.geometry(f"900x620+{x}+{y}")

        # Encabezado
        header = ctk.CTkFrame(modal, fg_color="transparent")
        header.pack(fill="x", padx=20, pady=(15, 10))

        ctk.CTkLabel(
            header,
            text="⚙️ Administración de Reglas de Automatización",
            font=ctk.CTkFont(size=20, weight="bold"),
            text_color=("black", "white")
        ).pack(anchor="w")

        ctk.CTkLabel(
            header,
            text="Modifica el contenido de tus reglas, activa/desactiva su disponibilidad o elimínalas sin saturar la pantalla principal.",
            font=ctk.CTkFont(size=12),
            text_color=("gray30", "gray70")
        ).pack(anchor="w")

        # Contenedor principal de 2 columnas
        body = ctk.CTkFrame(modal, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=20, pady=(0, 15))
        body.grid_columnconfigure(0, weight=3)
        body.grid_columnconfigure(1, weight=2)
        body.grid_rowconfigure(0, weight=1)

        # Columna Izquierda: Lista de Reglas
        left_frame = ctk.CTkFrame(body)
        left_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        left_frame.grid_rowconfigure(1, weight=1)
        left_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            left_frame,
            text="📋 Reglas Registradas en la Base de Datos:",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color=("black", "white")
        ).grid(row=0, column=0, sticky="w", padx=15, pady=(12, 6))

        scroll_modal_reglas = ctk.CTkScrollableFrame(left_frame)
        scroll_modal_reglas.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 12))

        # Columna Derecha: Editor de Regla
        right_frame = ctk.CTkFrame(body)
        right_frame.grid(row=0, column=1, sticky="nsew", padx=(10, 0))

        ctk.CTkLabel(
            right_frame,
            text="✏️ Editor de Regla",
            font=ctk.CTkFont(size=15, weight="bold"),
            text_color=("black", "white")
        ).pack(padx=15, pady=(12, 4))

        lbl_editor_estado = ctk.CTkLabel(
            right_frame,
            text="👈 Seleccione una regla y presione '✏️ Editar'",
            font=ctk.CTkFont(size=11),
            text_color=("gray40", "gray60")
        )
        lbl_editor_estado.pack(padx=15, pady=(0, 8))

        # Campos de edición
        ctk.CTkLabel(right_frame, text="Nombre de la Regla:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(4, 2))
        entry_edit_nombre = ctk.CTkEntry(right_frame, placeholder_text="Nombre de regla")
        entry_edit_nombre.pack(fill="x", padx=15, pady=(0, 6))

        ctk.CTkLabel(right_frame, text="Canal de Envío:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(4, 2))
        combo_edit_tipo = ctk.CTkComboBox(right_frame, values=["WhatsApp", "Correo SMTP"], state="readonly")
        combo_edit_tipo.pack(fill="x", padx=15, pady=(0, 6))

        ctk.CTkLabel(right_frame, text="Condición (ej. Activo, Inactivo, todos):", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(4, 2))
        entry_edit_condicion = ctk.CTkEntry(right_frame, placeholder_text="Condición requerida")
        entry_edit_condicion.pack(fill="x", padx=15, pady=(0, 6))

        ctk.CTkLabel(right_frame, text="Mensaje Plantilla:", font=ctk.CTkFont(size=11, weight="bold"), text_color=("black", "white")).pack(anchor="w", padx=15, pady=(4, 2))
        text_edit_mensaje = ctk.CTkTextbox(right_frame, height=100)
        text_edit_mensaje.pack(fill="both", expand=True, padx=15, pady=(0, 6))

        switch_edit_activa = ctk.CTkSwitch(right_frame, text="Regla Activa", text_color=("black", "white"))
        switch_edit_activa.select()
        switch_edit_activa.pack(anchor="w", padx=15, pady=(4, 10))

        # Botones de acción del editor
        btn_box_edit = ctk.CTkFrame(right_frame, fg_color="transparent")
        btn_box_edit.pack(fill="x", padx=15, pady=(0, 12))

        regla_en_edicion = {"id": None}

        def limpiar_editor():
            regla_en_edicion["id"] = None
            entry_edit_nombre.delete(0, "end")
            combo_edit_tipo.set("WhatsApp")
            entry_edit_condicion.delete(0, "end")
            text_edit_mensaje.delete("1.0", "end")
            switch_edit_activa.select()
            btn_guardar_cambios.configure(state="disabled")
            lbl_editor_estado.configure(
                text="👈 Seleccione una regla y presione '✏️ Editar'",
                text_color=("gray40", "gray60")
            )

        def cargar_en_editor(r: Dict[str, Any]):
            regla_en_edicion["id"] = r["id"]
            entry_edit_nombre.delete(0, "end")
            entry_edit_nombre.insert(0, r.get("nombre_regla", ""))
            
            tipo_val = "Correo SMTP" if "SMTP" in str(r.get("tipo", "")).upper() else "WhatsApp"
            combo_edit_tipo.set(tipo_val)
            
            entry_edit_condicion.delete(0, "end")
            entry_edit_condicion.insert(0, r.get("condicion", ""))
            
            text_edit_mensaje.delete("1.0", "end")
            text_edit_mensaje.insert("1.0", r.get("mensaje", ""))
            
            if r.get("activa"):
                switch_edit_activa.select()
            else:
                switch_edit_activa.deselect()
                
            btn_guardar_cambios.configure(state="normal")
            lbl_editor_estado.configure(
                text=f"Modificando Regla #{r['id']}: {r.get('nombre_regla')}",
                text_color="#2980b9"
            )

        def guardar_cambios_regla():
            rid = regla_en_edicion.get("id")
            if not rid:
                return
            nom = entry_edit_nombre.get().strip()
            tip = combo_edit_tipo.get().strip()
            cond = entry_edit_condicion.get().strip()
            msg = text_edit_mensaje.get("1.0", "end-1c").strip()
            act = bool(switch_edit_activa.get())

            if not nom or not cond or not msg:
                messagebox.showwarning("Campos Requeridos", "Debe completar nombre, condición y mensaje.", parent=modal)
                return

            exito = self.db.update_regla(
                regla_id=rid,
                nombre_regla=nom,
                tipo=tip,
                condicion=cond,
                mensaje=msg,
                activa=act
            )
            if exito:
                messagebox.showinfo("Éxito", f"Regla #{rid} actualizada exitosamente.", parent=modal)
                limpiar_editor()
                recargar_lista_modal()
                self.actualizar_selectores_destinatarios()
                self.renderizar_cola_destinatarios()
            else:
                messagebox.showerror("Error", "No se pudo actualizar la regla.", parent=modal)

        btn_guardar_cambios = ctk.CTkButton(
            btn_box_edit,
            text="💾 Guardar Cambios",
            fg_color="#27ae60",
            hover_color="#219955",
            font=ctk.CTkFont(size=11, weight="bold"),
            state="disabled",
            command=guardar_cambios_regla
        )
        btn_guardar_cambios.pack(side="left", fill="x", expand=True, padx=(0, 4))

        btn_cancelar_edit = ctk.CTkButton(
            btn_box_edit,
            text="✖️ Cancelar",
            fg_color="#7f8c8d",
            hover_color="#636e72",
            font=ctk.CTkFont(size=11),
            command=limpiar_editor
        )
        btn_cancelar_edit.pack(side="right", fill="x", expand=True, padx=(4, 0))

        def recargar_lista_modal():
            for w in scroll_modal_reglas.winfo_children():
                w.destroy()

            reglas = self.db.get_all_reglas()
            if not reglas:
                ctk.CTkLabel(
                    scroll_modal_reglas,
                    text="No hay reglas registradas.\nPuede crear nuevas reglas desde el formulario principal.",
                    text_color=("gray40", "gray60")
                ).pack(pady=30)
                return

            for r in reglas:
                card = ctk.CTkFrame(scroll_modal_reglas, fg_color=["gray85", "gray25"], corner_radius=6)
                card.pack(fill="x", padx=4, pady=4)
                card.grid_columnconfigure(0, weight=1)

                tipo_icono = "📲" if "WHATSAPP" in str(r.get('tipo', '')).upper() else "✉️"
                titulo = f"#{r['id']} {r['nombre_regla']} [{tipo_icono} {r['tipo']}]"
                sub = f"Condición: '{r['condicion']}' | Estado: {'🟢 Activa' if r['activa'] else '🔴 Inactiva'}"

                ctk.CTkLabel(
                    card, text=titulo, font=ctk.CTkFont(size=12, weight="bold"), anchor="w", text_color=("black", "white")
                ).grid(row=0, column=0, padx=10, pady=(6, 0), sticky="w")

                ctk.CTkLabel(
                    card, text=sub, font=ctk.CTkFont(size=10), text_color=("gray30", "gray70"), anchor="w"
                ).grid(row=1, column=0, padx=10, pady=(0, 2), sticky="w")

                preview_msg = (r.get("mensaje") or "").replace("\n", " ")
                if len(preview_msg) > 65:
                    preview_msg = preview_msg[:62] + "..."
                ctk.CTkLabel(
                    card, text=f'💬 "{preview_msg}"', font=ctk.CTkFont(size=10, slant="italic"), text_color=("gray40", "gray60"), anchor="w"
                ).grid(row=2, column=0, padx=10, pady=(0, 6), sticky="w")

                b_frame = ctk.CTkFrame(card, fg_color="transparent")
                b_frame.grid(row=0, column=1, rowspan=3, padx=(4, 8), pady=4)

                btn_edit = ctk.CTkButton(
                    b_frame,
                    text="✏️ Editar",
                    width=65,
                    height=24,
                    fg_color="#3498db",
                    hover_color="#2980b9",
                    font=ctk.CTkFont(size=10, weight="bold"),
                    command=lambda regla=r: cargar_en_editor(regla)
                )
                btn_edit.pack(side="left", padx=2)

                btn_tog = ctk.CTkButton(
                    b_frame,
                    text="Desactivar" if r["activa"] else "Activar",
                    width=70,
                    height=24,
                    fg_color="#f39c12" if r["activa"] else "#27ae60",
                    hover_color="#d68910" if r["activa"] else "#219955",
                    font=ctk.CTkFont(size=10),
                    command=lambda regla=r: toggle_regla_modal(regla)
                )
                btn_tog.pack(side="left", padx=2)

                btn_del = ctk.CTkButton(
                    b_frame,
                    text="🗑️",
                    width=28,
                    height=24,
                    fg_color="#c0392b",
                    hover_color="#962d22",
                    font=ctk.CTkFont(size=10),
                    command=lambda rid=r["id"]: eliminar_regla_modal(rid)
                )
                btn_del.pack(side="left", padx=2)

        def toggle_regla_modal(r: Dict[str, Any]):
            nuevo_est = not bool(r.get("activa"))
            self.db.update_regla(regla_id=r["id"], activa=nuevo_est)
            recargar_lista_modal()
            self.actualizar_selectores_destinatarios()
            self.renderizar_cola_destinatarios()

        def eliminar_regla_modal(rid: int):
            if messagebox.askyesno("Eliminar Regla", f"¿Está seguro de eliminar la regla #{rid} permanentemente?", parent=modal):
                self.db.delete_regla(rid)
                recargar_lista_modal()
                limpiar_editor()
                self.actualizar_selectores_destinatarios()
                self.renderizar_cola_destinatarios()

        recargar_lista_modal()

    def iniciar_ejecucion_automatica(self):
        """Inicia el despacho de la cola de automatización con intervalos humanos anti-baneo."""
        # Validar que haya clientes en la cola
        if not self.cola_destinatarios_seleccionados:
            messagebox.showwarning(
                "Cola Vacía",
                "No hay clientes en la cola de envíos.\n\n"
                "Por favor agregue clientes a la sección 'Destinatarios a Enviar' y asígneles una regla antes de iniciar."
            )
            return

        # Validar tiempos configurados
        try:
            d_min = int(self.entry_delay_min.get().strip())
            d_max = int(self.entry_delay_max.get().strip())
        except ValueError:
            messagebox.showerror("Tiempos Inválidos", "Los tiempos mínimo y máximo deben ser números enteros en segundos.")
            return

        if d_min < 1 or d_max < 1:
            messagebox.showerror("Tiempos Inválidos", "Los tiempos deben ser mayores o iguales a 1 segundo.")
            return

        if d_min > d_max:
            d_min, d_max = d_max, d_min

        # Advertencia de seguridad si el intervalo es muy bajo (< 5s)
        if d_min < 5:
            continuar = messagebox.askyesno(
                "Advertencia de Seguridad Anti-Baneo",
                "⚠️ Ha configurado un intervalo menor a 5 segundos.\n\n"
                "Enviar mensajes a esta velocidad puede ser clasificado como spam y provocar el baneo de su cuenta de WhatsApp.\n\n"
                "¿Desea continuar de todas formas?"
            )
            if not continuar:
                return

        # Auto-guardar configuración de tiempos antes de iniciar
        self.guardar_tiempos_antiban()

        # Preparar cola de tareas con mensajes personalizados para cada cliente según su regla asignada
        cola_preparada: List[Dict[str, Any]] = []
        fecha_actual = datetime.now().strftime("%d/%m/%Y")
        for idx_item, item in enumerate(self.cola_destinatarios_seleccionados):
            cli = item["cliente"]
            reg = item["regla"]
            nombre = cli.get("nombre", "Cliente")
            correo = (cli.get("correo") or "").strip()
            telefono = (cli.get("telefono") or "").strip()
            estado_cli = cli.get("estado", "")
            nombre_regla = reg.get("nombre_regla", "Regla")
            canal_item = item.get("canal")
            if not canal_item:
                if hasattr(self, "seg_canal_auto"):
                    canal_item = "WhatsApp" if "WHATSAPP" in self.seg_canal_auto.get().upper() else "Correo SMTP"
                else:
                    tipo_canal = (reg.get("tipo") or "").strip().upper()
                    canal_item = "WhatsApp" if ("WHATSAPP" in tipo_canal or "WSP" in tipo_canal) else "SMTP"

            es_whatsapp = "WHATSAPP" in canal_item.upper()
            nombre_canal = "WhatsApp" if es_whatsapp else "SMTP"

            mensaje_plantilla = reg.get("mensaje", "")
            cuerpo = (
                mensaje_plantilla
                .replace("{nombre}", nombre)
                .replace("{correo}", correo)
                .replace("{telefono}", telefono)
                .replace("{estado}", estado_cli)
                .replace("{fecha}", fecha_actual)
                .replace("{minimarket}", "Minimarket")
            )

            cola_preparada.append({
                "indice_cola": idx_item,
                "cliente": cli,
                "nombre": nombre,
                "correo": correo,
                "telefono": telefono,
                "regla": reg,
                "nombre_regla": nombre_regla,
                "es_whatsapp": es_whatsapp,
                "canal": nombre_canal,
                "cuerpo": cuerpo
            })

        # Validar si hay mensajes de WhatsApp y verificar conexión antes de arrancar
        hay_whatsapp = any(item.get("es_whatsapp") for item in cola_preparada)
        hay_smtp = any(not item.get("es_whatsapp") for item in cola_preparada)

        gestor_wsp = None
        if hay_whatsapp:
            self.btn_ejecutar_auto.configure(state="disabled", text="⏳ Verificando WhatsApp...")
            self.lbl_cola_estado.configure(text="🔍 Comprobando conexión activa de WhatsApp Web...", text_color="#3a7ebf")
            self.update_idletasks()
            try:
                gestor_wsp = self.obtener_gestor_whatsapp()
                conectado = gestor_wsp.verificar_sesion_activa(forzar_navegacion=True, timeout_espera=6)
                datos_cuenta = gestor_wsp.obtener_datos_cuenta_conectada() if conectado else {}
                self._actualizar_estado_visual_whatsapp(conectado, datos_cuenta)
                if not conectado:
                    self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                    self.btn_detener_auto.configure(state="disabled")
                    self.lbl_cola_estado.configure(
                        text="⚠️ Despacho bloqueado: WhatsApp está desconectado.",
                        text_color="#e74c3c"
                    )
                    ir_a_conexion = messagebox.askyesno(
                        "WhatsApp Desconectado",
                        "⚠️ No se puede despachar la cola porque contiene envíos por WhatsApp y WhatsApp se encuentra DESCONECTADO.\n\n"
                        "¿Desea ir a la pantalla de 'Conexión WhatsApp' para conectarlo?\n\n"
                        "(Nota: Si desea enviar sus mensajes por Correo SMTP sin WhatsApp, seleccione '✉️ Correo SMTP' en el selector de Canal de Envío)."
                    )
                    if ir_a_conexion:
                        self.seleccionar_vista("whatsapp")
                    return
            except Exception as e_chk:
                self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                self.btn_detener_auto.configure(state="disabled")
                self.lbl_cola_estado.configure(text=f"❌ Error al comprobar WhatsApp: {e_chk}", text_color="#e74c3c")
                messagebox.showerror("Error de Conexión", f"No se pudo verificar la sesión de WhatsApp:\n{e_chk}")
                return

        # Validar si hay mensajes de Correo SMTP y verificar conexión antes de arrancar
        smtp = None
        if hay_smtp:
            self.btn_ejecutar_auto.configure(state="disabled", text="⏳ Verificando SMTP...")
            self.lbl_cola_estado.configure(text="🔍 Comprobando credenciales y conexión SMTP...", text_color="#3a7ebf")
            self.update_idletasks()

            smtp = self.obtener_smtp_manager()
            if not smtp:
                self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                self.btn_detener_auto.configure(state="disabled")
                self.lbl_cola_estado.configure(
                    text="⚠️ Despacho bloqueado: Faltan credenciales SMTP.",
                    text_color="#e74c3c"
                )
                messagebox.showwarning(
                    "Credenciales SMTP Requeridas",
                    "⚠️ La cola de automatización contiene envíos por Correo SMTP, pero no hay credenciales configuradas.\n\n"
                    "Por favor vaya a la pestaña 'Envíos Email (SMTP)', complete su correo remitente y contraseña, y guárdelas."
                )
                return

            exito_smtp, msg_smtp = smtp.verificar_conexion()
            self._actualizar_estado_visual_smtp(exito_smtp, email=smtp.email)
            if not exito_smtp:
                self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                self.btn_detener_auto.configure(state="disabled")
                self.lbl_cola_estado.configure(
                    text=f"⚠️ Despacho bloqueado: Error SMTP ({msg_smtp}).",
                    text_color="#e74c3c"
                )
                messagebox.showerror(
                    "Error de Conexión SMTP",
                    f"No se pudo conectar con el servidor SMTP para despachar los correos:\n\n{msg_smtp}\n\n"
                    "Por favor revise sus credenciales en 'Envíos Email (SMTP)' antes de reintentar."
                )
                return

        # Inicializar visualmente todos los items de la cola en pendiente antes de arrancar
        for idx_i, item_i in enumerate(self.cola_destinatarios_seleccionados):
            item_i["estado_envio"] = "pendiente"
            item_i["hora_envio"] = None
            self.actualizar_estado_item_destinatario(idx_i, estado="pendiente")

        # Función para proteger los botones de la cola durante el despacho
        def set_estado_botones_cola(habilitado: bool):
            st = "normal" if habilitado else "disabled"
            if hasattr(self, "seg_canal_auto"):
                self.seg_canal_auto.configure(state=st)
            if hasattr(self, "btn_agregar_dest"):
                self.btn_agregar_dest.configure(state=st)
            if hasattr(self, "btn_masivo_dest"):
                self.btn_masivo_dest.configure(state=st)
            if hasattr(self, "btn_vaciar_dest"):
                self.btn_vaciar_dest.configure(state=st)
            if hasattr(self, "_widgets_destinatarios"):
                for w in self._widgets_destinatarios:
                    btn_q = w.get("btn_quitar")
                    if btn_q and btn_q.winfo_exists():
                        btn_q.configure(state=st)

        set_estado_botones_cola(False)

        smtp = self.obtener_smtp_manager()
        self.btn_ejecutar_auto.configure(state="disabled", text="⏳ Despachando...")
        self.btn_detener_auto.configure(state="normal")
        self.progress_cola.set(0)
        self.lbl_cola_estado.configure(text="🚀 Preparando cola de envíos...", text_color="#3a7ebf")
        self.txt_cola_log.delete("1.0", "end")

        def callback_ui(evento_data: Dict[str, Any]):
            def actualizar():
                tipo_ev = evento_data.get("evento")
                hora = datetime.now().strftime("%H:%M:%S")

                if tipo_ev == "iniciado":
                    tot = evento_data["total"]
                    self.lbl_cola_estado.configure(
                        text=f"🟢 Cola activa: {tot} mensajes programados (Intervalo: {evento_data['d_min']}s a {evento_data['d_max']}s)",
                        text_color="#27ae60"
                    )
                    self.lbl_cola_metricas.configure(
                        text=f"Progreso: 0/{tot}  |  ✅ Enviados: 0  |  ❌ Fallidos: 0  |  ⏳ En cola: {tot}"
                    )
                    self.txt_cola_log.insert("end", f"[{hora}] 🚀 Cola iniciada: {tot} mensajes en total con pausas de {evento_data['d_min']}s a {evento_data['d_max']}s.\n")
                    self.txt_cola_log.see("end")

                elif tipo_ev == "enviando":
                    idx = evento_data["indice"]
                    tot = evento_data["total"]
                    cli = evento_data["cliente"]
                    canal = evento_data["canal"]
                    idx_item = evento_data.get("indice_item", idx - 1)
                    self.actualizar_estado_item_destinatario(idx_item, estado="enviando")
                    self.lbl_cola_estado.configure(
                        text=f"📤 [{idx}/{tot}] Enviando por {canal} a {cli}...",
                        text_color="#2980b9"
                    )
                    self.txt_cola_log.insert("end", f"[{hora}] [{idx}/{tot}] Despachando mensaje a {cli} vía {canal}...\n")
                    self.txt_cola_log.see("end")

                elif tipo_ev == "item_completado":
                    idx = evento_data["indice"]
                    tot = evento_data["total"]
                    cli = evento_data["cliente"]
                    exito = evento_data["exito"]
                    est = evento_data["estado"]
                    env = evento_data["enviados"]
                    fal = evento_data["fallidos"]
                    hora_envio = evento_data.get("hora_envio") or hora
                    idx_item = evento_data.get("indice_item", idx - 1)
                    self.actualizar_estado_item_destinatario(
                        idx_item,
                        estado="enviado" if exito else "fallido",
                        hora=hora_envio
                    )
                    rest = tot - idx

                    self.progress_cola.set(idx / tot)
                    self.lbl_cola_metricas.configure(
                        text=f"Progreso: {idx}/{tot}  |  ✅ Enviados: {env}  |  ❌ Fallidos: {fal}  |  ⏳ En cola: {rest}"
                    )
                    icono = "✅" if exito else "❌"
                    self.txt_cola_log.insert("end", f"[{hora_envio}] {icono} {cli}: {est}\n")
                    self.txt_cola_log.see("end")

                elif tipo_ev == "esperando":
                    restante = evento_data["segundos_restantes"]
                    pausa_total = evento_data["pausa_total"]
                    sig = evento_data["siguiente_cliente"]
                    self.lbl_cola_estado.configure(
                        text=f"⏳ Pausa Humana Anti-Baneo: Próximo mensaje a '{sig}' en {restante}s (Pausa: {pausa_total}s)...",
                        text_color="#f39c12"
                    )

                elif tipo_ev == "cancelado":
                    set_estado_botones_cola(True)
                    self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                    self.btn_detener_auto.configure(state="disabled")
                    self.lbl_cola_estado.configure(
                        text="🛑 Cola detenida por el usuario.",
                        text_color="#e74c3c"
                    )
                    self.txt_cola_log.insert("end", f"[{hora}] 🛑 Cola de envíos detenida por el usuario.\n")
                    self.txt_cola_log.see("end")
                    messagebox.showwarning(
                        "Cola Detenida",
                        f"La cola de envíos ha sido detenida.\n\nEnviados: {evento_data.get('enviados', 0)}\nFallidos: {evento_data.get('fallidos', 0)}"
                    )

                elif tipo_ev == "finalizado":
                    set_estado_botones_cola(True)
                    self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                    self.btn_detener_auto.configure(state="disabled")
                    self.progress_cola.set(1.0)
                    self.lbl_cola_estado.configure(
                        text="🎉 ¡Cola de automatizaciones completada con éxito!",
                        text_color="#27ae60"
                    )
                    tot = evento_data.get("total", 0)
                    env = evento_data.get("enviados", 0)
                    fal = evento_data.get("fallidos", 0)
                    self.lbl_cola_metricas.configure(
                        text=f"Progreso: {tot}/{tot}  |  ✅ Enviados: {env}  |  ❌ Fallidos: {fal}  |  ⏳ En cola: 0"
                    )
                    self.txt_cola_log.insert("end", f"[{hora}] 🎉 Cola finalizada. Total: {tot} | Enviados: {env} | Fallidos: {fal}.\n")
                    self.txt_cola_log.see("end")
                    messagebox.showinfo(
                        "Automatización Finalizada",
                        f"¡Cola de mensajes procesada exitosamente!\n\n"
                        f"Total Mensajes: {tot}\n"
                        f"Enviados con éxito: {env}\n"
                        f"Fallidos: {fal}"
                    )

                elif tipo_ev == "vacio":
                    set_estado_botones_cola(True)
                    self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                    self.btn_detener_auto.configure(state="disabled")
                    self.lbl_cola_estado.configure(
                        text="ℹ️ No hay clientes que cumplan las condiciones de las reglas activas.",
                        text_color=("gray30", "gray70")
                    )
                    self.txt_cola_log.insert("end", f"[{hora}] ℹ️ No hay clientes destinatarios que cumplan las condiciones.\n")
                    self.txt_cola_log.see("end")
                    messagebox.showinfo("Sin Destinatarios", "No hay clientes registrados que cumplan las condiciones de las reglas activas.")

            self.after(0, actualizar)

        def tarea_auto():
            try:
                self.automator.evaluar_y_ejecutar(
                    smtp=smtp, 
                    whatsapp=gestor_wsp or self.whatsapp_gestor,
                    delay_min=d_min,
                    delay_max=d_max,
                    callback_progreso=callback_ui,
                    cola_personalizada=cola_preparada
                )
            except Exception as e:
                def on_error(err=str(e)):
                    set_estado_botones_cola(True)
                    self.btn_ejecutar_auto.configure(state="normal", text="🚀 Iniciar Cola de Envíos")
                    self.btn_detener_auto.configure(state="disabled")
                    self.lbl_cola_estado.configure(text=f"❌ Error en la ejecución: {err}", text_color="#e74c3c")
                    messagebox.showerror("Error en Automatización", f"Ocurrió un error inesperado al despachar la cola:\n{err}")
                self.after(0, on_error)

        threading.Thread(target=tarea_auto, daemon=True).start()

    def detener_cola_automatica(self):
        """Detiene de inmediato la cola de automatización en curso."""
        confirmar = messagebox.askyesno(
            "Detener Cola de Envíos",
            "¿Está seguro de que desea detener la cola de envíos en curso?\nLos mensajes que ya fueron enviados no se revertirán."
        )
        if confirmar:
            self.automator.cancelar_cola()
            self.btn_detener_auto.configure(state="disabled")
            self.lbl_cola_estado.configure(text="⏳ Deteniendo cola de envíos...", text_color="#f39c12")

    # ------------------------------------------------------------------
    # VISTA 5: CONEXIÓN WHATSAPP (Estado en tiempo real y QR)
    # ------------------------------------------------------------------
    def construir_vista_whatsapp(self) -> ctk.CTkFrame:
        frame = ctk.CTkFrame(self.container, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(0, weight=1)

        card = ctk.CTkFrame(frame)
        card.grid(row=0, column=0, sticky="nsew", padx=30, pady=20)
        card.grid_columnconfigure(0, weight=1)

        # Encabezado
        ctk.CTkLabel(
            card, 
            text="📲 Vinculación y Conexión con WhatsApp Web", 
            font=ctk.CTkFont(size=22, weight="bold"),
            text_color=("black", "white")
        ).pack(pady=(20, 5))

        # Banner Superior de Estado en Tiempo Real (Redondo Verde / Rojo)
        self.wsp_status_banner = ctk.CTkFrame(card, fg_color=["gray90", "gray20"], corner_radius=10)
        self.wsp_status_banner.pack(fill="x", padx=40, pady=(5, 12))
        self.wsp_status_banner.grid_columnconfigure(1, weight=1)

        self.lbl_wsp_status_dot = ctk.CTkLabel(
            self.wsp_status_banner,
            text="🔴",
            font=ctk.CTkFont(size=26)
        )
        self.lbl_wsp_status_dot.grid(row=0, column=0, rowspan=2, padx=(15, 10), pady=10)

        self.lbl_wsp_status_text = ctk.CTkLabel(
            self.wsp_status_banner,
            text="WhatsApp Desconectado",
            font=ctk.CTkFont(size=16, weight="bold"),
            text_color="#e74c3c",
            anchor="w"
        )
        self.lbl_wsp_status_text.grid(row=0, column=1, sticky="w", pady=(8, 0))

        self.lbl_wsp_status_desc = ctk.CTkLabel(
            self.wsp_status_banner,
            text="Escanee el código QR para conectar WhatsApp al sistema del minimarket.",
            font=ctk.CTkFont(size=12),
            text_color=("gray30", "gray70"),
            anchor="w"
        )
        self.lbl_wsp_status_desc.grid(row=1, column=1, sticky="w", pady=(0, 8))

        # Botón para comprobar estado manualmente
        btn_recheck = ctk.CTkButton(
            self.wsp_status_banner,
            text="🔄 Verificar Estado",
            width=135,
            height=32,
            fg_color=["#3a7ebf", "#1f538d"],
            command=self.verificar_estado_whatsapp_manual
        )
        btn_recheck.grid(row=0, column=2, rowspan=2, padx=15, pady=10)

        # Contenedor dinámico 1: Frame CUANDO ESTÁ CONECTADO (🟢)
        self.frame_wsp_conectado = ctk.CTkFrame(card, fg_color="transparent")
        
        # Tarjeta estilizada de éxito
        card_exito = ctk.CTkFrame(self.frame_wsp_conectado, fg_color=["#d4edda", "#143a21"], corner_radius=12, border_width=1, border_color="#28a745")
        card_exito.pack(fill="x", padx=40, pady=15)
        
        ctk.CTkLabel(
            card_exito,
            text="🟢 ¡WHATSAPP CONECTADO Y VINCULADO!",
            font=ctk.CTkFont(size=18, weight="bold"),
            text_color=("#155724", "#2ecc71")
        ).pack(pady=(16, 6))

        ctk.CTkLabel(
            card_exito,
            text="Tu cuenta de WhatsApp ya se encuentra vinculada y sincronizada en segundo plano.\nEl código QR se oculta automáticamente porque no es necesario volver a escanearlo.\nPuedes enviar mensajes manuales con plantillas o ejecutar automatizaciones.",
            font=ctk.CTkFont(size=13),
            text_color=("#155724", "#d4edda"),
            justify="center"
        ).pack(padx=20, pady=(0, 12))

        # Cuadro de Información de la Cuenta Conectada (Número y Nombre)
        self.box_cuenta_wsp = ctk.CTkFrame(
            card_exito, 
            fg_color=["#c3e6cb", "#0f2d18"], 
            corner_radius=10, 
            border_width=1, 
            border_color=["#85cf9d", "#1e6b36"]
        )
        self.box_cuenta_wsp.pack(fill="x", padx=30, pady=(0, 16))
        self.box_cuenta_wsp.grid_columnconfigure((0, 1), weight=1)

        # Columna 1: Número Conectado
        f_num = ctk.CTkFrame(self.box_cuenta_wsp, fg_color="transparent")
        f_num.grid(row=0, column=0, padx=20, pady=12, sticky="w")
        ctk.CTkLabel(
            f_num, 
            text="📱 Número Conectado:", 
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=("#155724", "#86efac")
        ).pack(anchor="w")
        self.lbl_wsp_cuenta_numero = ctk.CTkLabel(
            f_num,
            text="Detectando...",
            font=ctk.CTkFont(size=16, weight="bold"),
            text_color=("#0b2e13", "#ffffff")
        )
        self.lbl_wsp_cuenta_numero.pack(anchor="w", pady=(2, 0))

        # Columna 2: Nombre de la Cuenta / Perfil
        f_nom = ctk.CTkFrame(self.box_cuenta_wsp, fg_color="transparent")
        f_nom.grid(row=0, column=1, padx=20, pady=12, sticky="w")
        ctk.CTkLabel(
            f_nom, 
            text="👤 Nombre de Cuenta / Perfil:", 
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=("#155724", "#86efac")
        ).pack(anchor="w")
        self.lbl_wsp_cuenta_nombre = ctk.CTkLabel(
            f_nom,
            text="Detectando...",
            font=ctk.CTkFont(size=16, weight="bold"),
            text_color=("#0b2e13", "#ffffff")
        )
        self.lbl_wsp_cuenta_nombre.pack(anchor="w", pady=(2, 0))

        btn_ir_envios = ctk.CTkButton(
            self.frame_wsp_conectado,
            text="💬 Ir al Módulo de Envíos WhatsApp",
            font=ctk.CTkFont(size=15, weight="bold"),
            height=44,
            fg_color="#25D366",
            hover_color="#1ea952",
            text_color="white",
            command=lambda: self.seleccionar_vista("envios_wsp")
        )
        btn_ir_envios.pack(fill="x", padx=80, pady=10)

        btn_desvincular = ctk.CTkButton(
            self.frame_wsp_conectado,
            text="🚪 Cerrar Sesión / Desvincular Cuenta",
            font=ctk.CTkFont(size=13),
            height=36,
            fg_color="#e74c3c",
            hover_color="#c0392b",
            command=self.desconectar_sesion_whatsapp
        )
        btn_desvincular.pack(fill="x", padx=120, pady=8)

        # Contenedor dinámico 2: Frame CUANDO ESTÁ DESCONECTADO (🔴 Muestra instrucciones y QR)
        self.frame_wsp_desconectado = ctk.CTkFrame(card, fg_color="transparent")
        
        info_frame = ctk.CTkFrame(self.frame_wsp_desconectado, fg_color=["gray90", "gray20"])
        info_frame.pack(fill="x", padx=40, pady=6)

        instrucciones = (
            "📌 Instrucciones para conectar tu WhatsApp:\n"
            "1. Haz clic en 'Generar Código QR' (el navegador cargará de forma 100% oculta en segundo plano).\n"
            "2. En tu teléfono, abre WhatsApp > Menú (o Configuración) > Dispositivos vinculados.\n"
            "3. Toca en 'Vincular un dispositivo' y escanea el código QR que se mostrará aquí abajo.\n"
            "💡 La sesión quedará guardada de forma segura para no tener que escanearlo nuevamente."
        )

        ctk.CTkLabel(
            info_frame, 
            text=instrucciones, 
            font=ctk.CTkFont(size=12),
            justify="left",
            text_color=("black", "white")
        ).pack(padx=20, pady=8, anchor="w")

        # Botón para Generar / Actualizar Código QR
        self.btn_generar_qr = ctk.CTkButton(
            self.frame_wsp_desconectado,
            text="📲  Generar Código QR",
            font=ctk.CTkFont(size=16, weight="bold"),
            height=46,
            fg_color="#25D366",
            hover_color="#1ea952",
            text_color="white",
            command=self.solicitar_qr_whatsapp
        )
        self.btn_generar_qr.pack(fill="x", padx=60, pady=(8, 4))

        # Mensaje dinámico de estado
        self.lbl_estado_whatsapp = ctk.CTkLabel(
            self.frame_wsp_desconectado,
            text="Estado: Presione 'Generar Código QR' para comenzar.",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=("gray40", "gray70")
        )
        self.lbl_estado_whatsapp.pack(pady=3)

        # Contenedor visual para el código QR: Tarjeta blanca pura con bordes definidos
        self.qr_card_outer = ctk.CTkFrame(
            self.frame_wsp_desconectado, 
            width=280, 
            height=280, 
            fg_color="#FFFFFF", 
            corner_radius=14,
            border_width=2,
            border_color=["#cbd5e1", "#334155"]
        )
        self.qr_card_outer.pack(pady=(4, 8))
        self.qr_card_outer.pack_propagate(False)

        self.lbl_qr_display = ctk.CTkLabel(
            self.qr_card_outer,
            text="El código QR aparecerá aquí\nuna vez generado.",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#333333"
        )
        self.lbl_qr_display.pack(expand=True, fill="both", padx=8, pady=8)

        # Botón para cerrar o reiniciar el navegador en segundo plano
        self.btn_cerrar_whatsapp = ctk.CTkButton(
            self.frame_wsp_desconectado,
            text="Cerrar Navegador en Segundo Plano",
            font=ctk.CTkFont(size=12),
            height=30,
            fg_color="gray40",
            hover_color="gray30",
            command=self.cerrar_navegador_whatsapp
        )
        self.btn_cerrar_whatsapp.pack(pady=(4, 10))

        # Por defecto empaquetar frame_wsp_desconectado
        self.frame_wsp_desconectado.pack(fill="both", expand=True)

        return frame

    # ------------------------------------------------------------------
    # GESTIÓN Y SINCRONIZACIÓN VISUAL DEL ESTADO DE WHATSAPP (🟢 / 🔴)
    # ------------------------------------------------------------------
    def _actualizar_estado_visual_whatsapp(self, esta_conectado: bool, datos_cuenta: Optional[Dict[str, str]] = None):
        """Actualiza todos los indicadores de la aplicación con redondos verde (🟢) o rojo (🔴) y datos de la cuenta."""
        self._whatsapp_esta_conectado = bool(esta_conectado)

        if not esta_conectado:
            self._whatsapp_datos_cuenta = {}
            tel_fmt = ""
            nom_cuenta = ""
        else:
            if datos_cuenta:
                self._whatsapp_datos_cuenta = dict(datos_cuenta)
            elif not hasattr(self, "_whatsapp_datos_cuenta") or not self._whatsapp_datos_cuenta:
                if self.whatsapp_gestor:
                    try:
                        self._whatsapp_datos_cuenta = self.whatsapp_gestor.obtener_datos_cuenta_conectada()
                    except Exception:
                        self._whatsapp_datos_cuenta = {}
                else:
                    self._whatsapp_datos_cuenta = {}

            tel_fmt = self._whatsapp_datos_cuenta.get("telefono_formateado") or (("+" + self._whatsapp_datos_cuenta.get("telefono")) if self._whatsapp_datos_cuenta.get("telefono") else "")
            nom_cuenta = (self._whatsapp_datos_cuenta.get("nombre") or "").strip()

        # 1. Indicador en el Menú Lateral (Sidebar)
        if hasattr(self, "lbl_sidebar_wsp_status"):
            if esta_conectado:
                self.lbl_sidebar_wsp_status.configure(
                    text="🟢 WhatsApp: Conectado",
                    text_color="#2ecc71"
                )
            else:
                self.lbl_sidebar_wsp_status.configure(
                    text="🔴 WhatsApp: Desconectado",
                    text_color="#e74c3c"
                )

        if hasattr(self, "lbl_sidebar_wsp_cuenta"):
            if esta_conectado:
                if tel_fmt and nom_cuenta:
                    txt_cuenta = f"📱 {tel_fmt}\n👤 {nom_cuenta}"
                elif tel_fmt:
                    txt_cuenta = f"📱 {tel_fmt}"
                elif nom_cuenta:
                    txt_cuenta = f"👤 {nom_cuenta}"
                else:
                    txt_cuenta = "📱 Línea Activa"
                self.lbl_sidebar_wsp_cuenta.configure(
                    text=txt_cuenta,
                    text_color=("#155724", "#4ade80")
                )
            else:
                self.lbl_sidebar_wsp_cuenta.configure(
                    text="",
                    text_color=("gray40", "gray70")
                )

        # 2. Indicador en la Vista de Envíos Manuales WhatsApp
        if hasattr(self, "lbl_envios_wsp_badge"):
            if esta_conectado:
                if tel_fmt:
                    info_ext = f": {tel_fmt}" + (f" ({nom_cuenta})" if nom_cuenta else "")
                    self.lbl_envios_wsp_badge.configure(
                        text=f"🟢 WhatsApp Conectado{info_ext} - Listo para Enviar",
                        text_color="#2ecc71"
                    )
                else:
                    self.lbl_envios_wsp_badge.configure(
                        text="🟢 WhatsApp Conectado y Listo para Enviar",
                        text_color="#2ecc71"
                    )
            else:
                self.lbl_envios_wsp_badge.configure(
                    text="🔴 WhatsApp Desconectado (Verifique en 'Conexión WhatsApp')",
                    text_color="#e74c3c"
                )

        # 3. Vista de Vinculación WhatsApp: Banner superior y alternancia de tarjetas
        if hasattr(self, "lbl_wsp_status_dot") and hasattr(self, "lbl_wsp_status_text"):
            if esta_conectado:
                self.lbl_wsp_status_dot.configure(text="🟢")
                titulo_banner = f"WhatsApp Conectado ({tel_fmt})" if tel_fmt else "WhatsApp Conectado"
                self.lbl_wsp_status_text.configure(
                    text=titulo_banner,
                    text_color="#2ecc71"
                )
                if hasattr(self, "lbl_wsp_status_desc"):
                    desc = f"Línea activa: {tel_fmt}" if tel_fmt else "Línea activa vinculada"
                    if nom_cuenta:
                        desc += f"  |  Titular: {nom_cuenta}"
                    desc += ". Sesión activa y sincronizada. Listo para enviar mensajes."
                    self.lbl_wsp_status_desc.configure(text=desc)

                # Actualizar cuadro de detalles en la tarjeta de éxito
                if hasattr(self, "lbl_wsp_cuenta_numero"):
                    self.lbl_wsp_cuenta_numero.configure(text=tel_fmt if tel_fmt else "(Detectado)")
                if hasattr(self, "lbl_wsp_cuenta_nombre"):
                    self.lbl_wsp_cuenta_nombre.configure(text=nom_cuenta if nom_cuenta else "(Sin nombre de perfil)")

                # Alternar marcos: Ocultar QR y mostrar éxito
                if hasattr(self, "frame_wsp_desconectado"):
                    self.frame_wsp_desconectado.pack_forget()
                if hasattr(self, "frame_wsp_conectado"):
                    self.frame_wsp_conectado.pack(fill="both", expand=True)
            else:
                self.lbl_wsp_status_dot.configure(text="🔴")
                self.lbl_wsp_status_text.configure(
                    text="WhatsApp Desconectado",
                    text_color="#e74c3c"
                )
                if hasattr(self, "lbl_wsp_status_desc"):
                    self.lbl_wsp_status_desc.configure(
                        text="Escanee el código QR para conectar WhatsApp al sistema del minimarket."
                    )
                if hasattr(self, "lbl_wsp_cuenta_numero"):
                    self.lbl_wsp_cuenta_numero.configure(text="Desconectado")
                if hasattr(self, "lbl_wsp_cuenta_nombre"):
                    self.lbl_wsp_cuenta_nombre.configure(text="Ninguna")

                # Alternar marcos: Ocultar éxito y mostrar QR
                if hasattr(self, "frame_wsp_conectado"):
                    self.frame_wsp_conectado.pack_forget()
                if hasattr(self, "frame_wsp_desconectado"):
                    if not self.frame_wsp_desconectado.winfo_ismapped():
                        self.frame_wsp_desconectado.pack(fill="both", expand=True)

    def verificar_estado_whatsapp_asincrono(self):
        """Verificación rápida en segundo plano de la sesión de WhatsApp."""
        if not self.whatsapp_gestor or not self.whatsapp_gestor.driver:
            self._actualizar_estado_visual_whatsapp(False, {})
            return

        def tarea():
            conectado = False
            datos_cuenta = {}
            try:
                conectado = self.whatsapp_gestor.verificar_sesion_activa()
                if conectado:
                    datos_cuenta = self.whatsapp_gestor.obtener_datos_cuenta_conectada()
            except Exception:
                conectado = False
            self.after(0, lambda: self._actualizar_estado_visual_whatsapp(conectado, datos_cuenta))

        threading.Thread(target=tarea, daemon=True).start()

    def ejecutar_verificacion_whatsapp(self, es_manual: bool = False):
        """
        Comprueba el estado de conexión con WhatsApp Web de manera idéntica al botón 'Verificar Estado'.
        Se utiliza tanto al arrancar el sistema como al presionar el botón de verificación manual.
        """
        if getattr(self, "_whatsapp_verificando", False):
            if es_manual:
                messagebox.showinfo("En Proceso", "Ya se está comprobando la conexión de WhatsApp. Por favor espere un momento...")
            return

        self._whatsapp_verificando = True

        if hasattr(self, "lbl_sidebar_wsp_status"):
            self.lbl_sidebar_wsp_status.configure(
                text="🟡 WhatsApp: Comprobando...",
                text_color="#f39c12"
            )
        if hasattr(self, "lbl_sidebar_wsp_cuenta"):
            self.lbl_sidebar_wsp_cuenta.configure(
                text="Verificando cuenta...",
                text_color=("gray40", "gray70")
            )
        if hasattr(self, "lbl_wsp_status_desc"):
            self.lbl_wsp_status_desc.configure(text="⏳ Verificando sesión con WhatsApp Web en segundo plano...")

        def tarea():
            conectado = False
            datos_cuenta = {}
            try:
                gestor = self.obtener_gestor_whatsapp()
                conectado = gestor.verificar_sesion_activa(forzar_navegacion=True, timeout_espera=10)
                if conectado:
                    datos_cuenta = gestor.obtener_datos_cuenta_conectada(forzar_refresco=True)
            except Exception as e:
                logging.error(f"Error verificando conexión de WhatsApp: {e}")
                conectado = False
            finally:
                self._whatsapp_verificando = False
                def fin():
                    self._actualizar_estado_visual_whatsapp(conectado, datos_cuenta)
                    if es_manual:
                        if conectado:
                            tel = datos_cuenta.get("telefono_formateado") or datos_cuenta.get("telefono") or ""
                            nom = datos_cuenta.get("nombre") or ""
                            info_msg = ""
                            if tel:
                                info_msg += f"\n📱 Número: {tel}"
                            if nom:
                                info_msg += f"\n👤 Titular: {nom}"
                            messagebox.showinfo(
                                "WhatsApp Conectado",
                                f"🟢 ¡WhatsApp Web está Conectado y Vinculado!{info_msg}\n\nTu sesión se encuentra activa y lista para realizar envíos de mensajes."
                            )
                        else:
                            messagebox.showwarning(
                                "WhatsApp Desconectado",
                                "🔴 WhatsApp se encuentra Desconectado.\n\nPor favor presione 'Generar Código QR' para escanearlo desde su teléfono."
                            )
                self.after(0, fin)

        threading.Thread(target=tarea, daemon=True).start()

    def _auto_verificar_conexion_inicial_whatsapp(self):
        """Ejecuta la verificación al iniciar el sistema con la misma lógica del botón manual."""
        self.ejecutar_verificacion_whatsapp(es_manual=False)

    def verificar_estado_whatsapp_manual(self):
        """Disparado por el botón 'Verificar Estado'."""
        self.ejecutar_verificacion_whatsapp(es_manual=True)

    def chequeo_periodico_whatsapp(self):
        """Monitoreo periódico recurrente cada 12 segundos del estado de WhatsApp."""
        if (getattr(self, "_whatsapp_monitoreo_activo", False) or 
            getattr(self, "_whatsapp_generando_qr", False) or
            getattr(self, "_whatsapp_verificando", False)):
            self.after(12000, self.chequeo_periodico_whatsapp)
            return

        def tarea():
            try:
                if self.whatsapp_gestor and self.whatsapp_gestor.driver and not getattr(self, "_whatsapp_verificando", False):
                    conectado = self.whatsapp_gestor.verificar_sesion_activa(forzar_navegacion=False, timeout_espera=2)
                    datos_cuenta = {}
                    if conectado:
                        datos_cuenta = self.whatsapp_gestor.obtener_datos_cuenta_conectada(forzar_refresco=False)
                    self.after(0, lambda: self._actualizar_estado_visual_whatsapp(conectado, datos_cuenta))
            except Exception:
                pass

        if self.whatsapp_gestor and self.whatsapp_gestor.driver:
            threading.Thread(target=tarea, daemon=True).start()

        self.after(12000, self.chequeo_periodico_whatsapp)

    def desconectar_sesion_whatsapp(self):
        """Cierra la sesión activa de WhatsApp Web y limpia los datos de sesión."""
        confirmar = messagebox.askyesno(
            "Desvincular WhatsApp",
            "¿Está seguro de que desea cerrar la sesión de WhatsApp Web?\nPara volver a enviar mensajes tendrá que escanear un nuevo código QR."
        )
        if not confirmar:
            return

        def tarea():
            if self.whatsapp_gestor:
                try:
                    self.whatsapp_gestor.cerrar_sesion_whatsapp()
                except Exception:
                    try:
                        self.whatsapp_gestor.cerrar()
                    except Exception:
                        pass
                self.whatsapp_gestor = None

            def al_desconectar():
                self._actualizar_estado_visual_whatsapp(False, {})
                self.btn_generar_qr.configure(state="normal", text="📲  Generar Código QR")
                self.lbl_estado_whatsapp.configure(text="Sesión desvinculada. Genere un nuevo QR.", text_color=("gray40", "gray70"))
                self.lbl_qr_display.configure(image=None, text="El código QR aparecerá aquí\nuna vez generado.", text_color="#333333")
                messagebox.showinfo("Sesión Cerrada", "La sesión de WhatsApp ha sido cerrada correctamente.")

            self.after(0, al_desconectar)

        threading.Thread(target=tarea, daemon=True).start()

    def _iniciar_monitoreo_sesion_whatsapp(self):
        """Monitorea en segundo plano si el usuario ya escaneó el QR y vinculó el dispositivo."""
        self._whatsapp_monitoreo_activo = True

        def monitor():
            intentos = 0
            while self._whatsapp_monitoreo_activo and intentos < 45:
                time.sleep(2)
                intentos += 1
                if not self._whatsapp_monitoreo_activo:
                    break
                if self.whatsapp_gestor and self.whatsapp_gestor.verificar_sesion_activa():
                    datos_cuenta = self.whatsapp_gestor.obtener_datos_cuenta_conectada(forzar_refresco=True)
                    def on_vinculado_exitoso(d=datos_cuenta):
                        self._whatsapp_monitoreo_activo = False
                        self._actualizar_estado_visual_whatsapp(True, d)
                        tel = d.get("telefono_formateado") or d.get("telefono") or ""
                        nom = d.get("nombre") or ""
                        txt_extra = f"\n📱 Número: {tel}" if tel else ""
                        if nom: txt_extra += f"\n👤 Titular: {nom}"
                        messagebox.showinfo("WhatsApp Vinculado", f"¡Excelente! Tu WhatsApp ha sido vinculado.{txt_extra}\nLa sesión está activa y lista para enviar.")
                    self.after(0, on_vinculado_exitoso)
                    break

        threading.Thread(target=monitor, daemon=True).start()

    def solicitar_qr_whatsapp(self):
        if str(self.btn_generar_qr.cget("state")) == "disabled":
            return

        self._whatsapp_generando_qr = True
        self._whatsapp_monitoreo_activo = False
        self.btn_generar_qr.configure(state="disabled", text="⏳ Generando QR...")
        self.lbl_estado_whatsapp.configure(
            text="Generando QR seguro, por favor espere...",
            text_color="#f39c12"
        )
        self.lbl_qr_display.configure(
            image=None,
            text="⏳ Conectando con WhatsApp Web...\nPor favor espere unos segundos.",
            text_color="#333333"
        )

        def tarea_qr():
            try:
                gestor = self.obtener_gestor_whatsapp()
                qr_b64 = gestor.obtener_qr_base64()

                if qr_b64 == "CONECTADO":
                    datos_cuenta = gestor.obtener_datos_cuenta_conectada()
                    def on_conectado(d=datos_cuenta):
                        self._whatsapp_generando_qr = False
                        self._whatsapp_monitoreo_activo = False
                        self.btn_generar_qr.configure(state="normal", text="📲  Generar Código QR")
                        self._actualizar_estado_visual_whatsapp(True, d)
                        messagebox.showinfo("WhatsApp Vinculado", "Tu sesión de WhatsApp ya se encuentra guardada y activa.")
                    self.after(0, on_conectado)

                elif qr_b64:
                    raw_b64 = qr_b64.split("base64,")[-1] if "base64," in qr_b64 else qr_b64
                    img_bytes = base64.b64decode(raw_b64)
                    pil_img = Image.open(io.BytesIO(img_bytes))

                    def on_qr_listo(img_obj=pil_img):
                        self._whatsapp_generando_qr = False
                        ctk_img = ctk.CTkImage(light_image=img_obj, dark_image=img_obj, size=(260, 260))
                        self.lbl_qr_display.configure(image=ctk_img, text="")
                        self.btn_generar_qr.configure(state="normal", text="🔄  Actualizar Código QR")
                        self.lbl_estado_whatsapp.configure(
                            text="📷 Código QR listo. Escanéalo desde WhatsApp en tu teléfono.",
                            text_color="#27ae60"
                        )
                        # Iniciar escucha automática de escaneo exitoso
                        self._iniciar_monitoreo_sesion_whatsapp()
                    self.after(0, on_qr_listo)

                else:
                    def on_qr_timeout():
                        self._whatsapp_generando_qr = False
                        self._whatsapp_monitoreo_activo = False
                        self.btn_generar_qr.configure(state="normal", text="📲  Generar Código QR")
                        self.lbl_estado_whatsapp.configure(
                            text="❌ Tiempo de espera agotado al obtener el QR.",
                            text_color="#e74c3c"
                        )
                        self.lbl_qr_display.configure(image=None, text="⚠️ No se detectó el código QR.", text_color="#c0392b")
                        messagebox.showwarning("Aviso", "No se detectó el código QR en el tiempo esperado. Verifique su conexión.")
                    self.after(0, on_qr_timeout)

            except Exception as err:
                error_str = str(err)
                def on_qr_error(msg=error_str):
                    self._whatsapp_generando_qr = False
                    self._whatsapp_monitoreo_activo = False
                    self.btn_generar_qr.configure(state="normal", text="📲  Generar Código QR")
                    self.lbl_estado_whatsapp.configure(text=f"Error: {msg}", text_color="#e74c3c")
                    self.lbl_qr_display.configure(image=None, text="❌ Error al cargar WhatsApp Web.", text_color="#c0392b")
                    messagebox.showerror("Error al Iniciar WhatsApp", f"No se pudo generar el código QR:\n{msg}")
                self.after(0, on_qr_error)

        threading.Thread(target=tarea_qr, daemon=True).start()

    def cerrar_navegador_whatsapp(self):
        self._whatsapp_monitoreo_activo = False
        if self.whatsapp_gestor:
            try:
                self.whatsapp_gestor.cerrar()
            except Exception:
                pass
            finally:
                self.whatsapp_gestor = None
                self._actualizar_estado_visual_whatsapp(False)
                self.lbl_estado_whatsapp.configure(text="Navegador cerrado correctamente.", text_color=("gray40", "gray70"))
                self.lbl_qr_display.configure(image=None, text="El código QR aparecerá aquí\nuna vez generado.", text_color="#333333")
                self.btn_generar_qr.configure(state="normal", text="📲  Generar Código QR")
                messagebox.showinfo("WhatsApp", "El navegador de WhatsApp en segundo plano se ha cerrado.")
        else:
            self._actualizar_estado_visual_whatsapp(False)
            messagebox.showinfo("WhatsApp", "El navegador ya se encontraba cerrado.")


# ==========================================
# PUNTO DE ENTRADA
# ==========================================
if __name__ == "__main__":
    validar_licencia_o_salir()
    app = MinimarketApp()
    app.mainloop()
